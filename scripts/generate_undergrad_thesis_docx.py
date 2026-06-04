from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
OUT_DOCX = ROOT / "基于STM32的水质监测及断面数据展示系统_本科毕业论文.docx"

TITLE_CN = "基于STM32的水质监测及断面数据展示系统"
TITLE_EN = "Design and Implementation of an STM32-Based Water Quality Monitoring and Section Data Display System"
STUDENT = "胡卓凡"
STUDENT_ID = "202220050409"
ADVISOR = "李俊吉"
COLLEGE = "计算机科学与技术学院"
MAJOR_CLASS = "物联网222004"
DATE_TEXT = "2026年5月"
HEADER_TEXT = "太原科技大学学士学位论文"

TASK_BOOK_TITLE = "太原科技大学毕业设计（论文）任务书"
TASK_BOOK_META = [
    ("学 生 姓 名", STUDENT, "学      号", STUDENT_ID),
    ("专 业 班 级", MAJOR_CLASS, "同  组  人", "无"),
    ("任务下发时间", "2026年2月1日", "任务完成时间", "2026年5月20日"),
]
TASK_BOOK_TOPIC = TITLE_CN
TASK_BOOK_PURPOSE = (
    "本课题面向养殖水体巡检、河湖断面监测和教学实验等应用场景，设计并实现一套基于 STM32 的水质监测及断面数据展示系统。"
    "要求完成多参数采集终端、设备数据接入、历史与快照数据组织、断面地图展示、趋势分析和系统测试等核心内容，并形成符合本科毕业论文规范的设计文档。"
)
TASK_BOOK_CONTENT = [
    "完成以 STM32F103C8 为核心的终端方案设计，接入 DS18B20 温度探头以及 pH、浊度、TDS 模拟传感器，明确主控、采样、通信和调试接口之间的连接关系。",
    "完成终端侧采样、换算、状态控制和上报流程设计，结合 ADC、DMA 与 ESP-01 通信链路实现现场数据采集与系统接入。",
    "基于 Django、MySQL 和 Redis/Channels 完成后台系统设计，建立统一数据模型、历史数据表、最新快照表、坐标缓存表和 AI 日志表，支撑实时查询与历史回溯。",
    "基于 Vue 3、Element Plus、ECharts 和高德地图实现监测概览、综合分析、断面地图、数据源设置和 AI 辅助分析等页面，形成断面数据展示系统。",
    "完成系统联调、功能测试与运行结果分析，整理关键实现过程、测试数据和论文文稿。",
]
TASK_BOOK_SUBMIT = [
    "毕业设计论文一份（含中文摘要、英文摘要、目录、正文、参考文献和附录）；",
    "基于 STM32 的水质监测及断面数据展示系统完整源代码；",
    "系统使用说明、关键接口说明及测试记录；",
    "系统运行截图、附录英文资料翻译和代表性程序代码。",
]

CN_ABSTRACT = (
    "水质监测数据具有明显的时效性与空间关联性，直接影响养殖管理、水环境巡检以及断面研判的准确性。"
    "针对人工取样频率低、现场设备与后台服务脱节、断面信息难以统一展示等实际情况，本文依据当前项目代码、"
    "硬件原型和开题材料，对基于 STM32 的水质监测及断面数据展示系统展开设计与实现研究。终端以 "
    "STM32F103C8 为控制核心，接入 DS18B20 温度探头以及 pH、浊度、TDS 模拟传感器，通过 ADC 与 DMA "
    "完成连续采样，经经验标定换算后由 ESP-01 以 HTTP 接口上报。后台系统采用 Django、MySQL 与 "
    "Redis/Channels 构建统一接入与展示链路，利用 DataTransformer 实现多源字段归一化，借助 "
    "sensor_data 历史表、sensor_data_latest 快照表和 station_locations 坐标缓存表支撑趋势分析、"
    "实时查询与地图展示。前端基于 Vue 3、Element Plus、ECharts 和高德地图实现监测概览、综合分析、"
    "断面地图、数据源配置及 AI 辅助分析等功能。结合 2026 年 3 月 5 日至 2026 年 3 月 11 日的运行日志统计，"
    "系统共完成 465 次同步任务，其中 257 次产生新增写入，代表性地理编码批次平均成功率为 98.78%。"
    "测试结果表明，该系统已经形成“终端采集、系统入库、断面展示、辅助研判”的完整业务闭环，能够满足本科毕业设计对系统性、"
    "工程性和应用性的要求。"
)

EN_ABSTRACT = (
    "Water quality monitoring data is highly time-sensitive and spatially correlated, and it directly affects aquaculture "
    "management, water-environment inspection, and section-level assessment. To address the practical problems of low "
    "manual sampling frequency, weak linkage between field devices and the backend service, and fragmented section display, this "
    "thesis studies the design and implementation of an STM32-based water quality monitoring and section data display system "
    "based on the current project code, hardware prototype, and proposal materials. The terminal uses STM32F103C8 as the "
    "main controller, connects a DS18B20 temperature sensor and analog pH, turbidity, and TDS modules, performs continuous "
    "sampling through ADC and DMA, and uploads calibrated values through an ESP-01 module over an HTTP interface. The "
    "backend service is built with Django, MySQL, and Redis/Channels. DataTransformer is used to normalize heterogeneous fields, "
    "while the sensor_data history table, sensor_data_latest snapshot table, and station_locations cache table support "
    "trend analysis, real-time query, and map visualization. The frontend is implemented with Vue 3, Element Plus, "
    "ECharts, and AMap, and provides dashboard, analysis, map, data-source setting, and AI-assisted analysis pages. "
    "According to runtime logs from March 5, 2026 to March 11, 2026, the system completed 465 synchronization tasks, "
    "257 of which created new records, and representative geocoding batches achieved an average success rate of 98.78 "
    "percent. The results indicate that the project has formed a complete workflow covering terminal acquisition, backend "
    "storage, section display, and auxiliary interpretation, and it satisfies the systemic and engineering requirements of "
    "an undergraduate graduation project."
)

KEYWORDS_CN = "STM32；水质监测；断面展示；数据可视化；Django"
KEYWORDS_EN = "STM32; water quality monitoring; section display; data visualization; Django"

REFERENCES = [
    "[1] 国家环境保护总局. 地表水环境质量标准: GB 3838-2002[S]. 北京: 中国环境科学出版社, 2002.",
    "[2] OASIS. MQTT Version 3.1.1[S]. OASIS Standard, 2014.",
    "[3] OASIS. MQTT Version 5.0[S]. OASIS Standard, 2019.",
    "[4] Django Software Foundation. Django Documentation[EB/OL].",
    "[5] Encode OSS Ltd. Django REST framework Documentation[EB/OL].",
    "[6] Vue.js Team. Vue.js Documentation[EB/OL].",
    "[7] Apache Software Foundation. Apache ECharts Handbook[EB/OL].",
    "[8] 高德开放平台. Web服务 API 文档[EB/OL].",
    "[9] Oracle Corporation. MySQL 8.0 Reference Manual[EB/OL].",
    "[10] STMicroelectronics. RM0008 Reference Manual[EB/OL].",
    "[11] Python Software Foundation. Python Documentation[EB/OL].",
    "[12] Redis Labs. Redis Documentation[EB/OL].",
    "[13] 乐鑫信息科技. ESP8266 AT 指令集[EB/OL].",
    "[14] Maxim Integrated. DS18B20 Programmable Resolution 1-Wire Digital Thermometer[EB/OL].",
]

BLOCKS = [
    ("chapter", "第一章 系统概述"),
    ("section", "1.1 课题背景与研究意义"),
    ("subsection", "1.1.1 行业背景"),
    (
        "paragraph",
        "水质参数是养殖生产、水环境巡检和断面研判的重要依据。传统人工取样虽然能够完成阶段性检测，但采样频率有限、记录形式分散、"
        "结果回传速度较慢，很难反映连续变化过程中的短时波动。当监测对象从单点扩展到多个断面或多个养殖单元时，人工方式在数据时效性、"
        "空间关联性和历史追溯能力方面都会明显受限。[1]"
    ),
    (
        "paragraph",
        "随着嵌入式采集、无线通信、时序存储和可视化技术的发展，水质监测系统已经不再停留于单点仪表读数，而是逐步演化为由终端感知、"
        "平台接入、历史存储、地图展示和辅助分析共同组成的综合性平台。对于断面化管理场景而言，只有把现场采样和平台展示放到同一条数据链路中考虑，"
        "监测结果才具有连续观测和辅助决策的价值。"
    ),
    ("subsection", "1.1.2 研究意义"),
    (
        "paragraph",
        "本课题的现实意义主要体现在三个方面。其一，借助 STM32 终端和 Web 平台协同设计，可以把原本低频、离散的人工检测过程转化为连续的数据链路，"
        "提高监测工作的时效性。其二，平台将实时列表、趋势图和断面地图放在统一界面中展示，有助于管理人员从空间分布和时间变化两个维度理解水质状态。"
        "其三，统一数据模型和标准化接口能够减少多源接入时的字段差异，降低平台扩展成本。"
    ),
    (
        "paragraph",
        "从本科毕业设计训练目标看，该课题同时覆盖嵌入式开发、网络通信、后端建模、数据库组织和前端可视化等多个知识点。论文以当前项目的代码实现为基础，"
        "通过对硬件方案、平台结构和运行结果进行系统化梳理，能够较完整地体现学生在分析、设计、实现和验证各阶段的综合能力。"
    ),
    ("section", "1.2 国内外研究现状"),
    ("subsection", "1.2.1 国外研究现状"),
    (
        "paragraph",
        "国外相关研究起步较早，水质监测系统通常与低功耗传感网络、远程通信和长期运维结合开展。研究重点除了单一参数测量外，还包括节点稳定性、"
        "漂移补偿、分布式部署和数据回传效率等问题。轻量级网络协议、云端数据服务以及地图化展示已成为常见技术组合。[2][3]"
    ),
    (
        "paragraph",
        "在平台侧，国外方案普遍重视历史数据管理与可视分析的结合。一方面需要保存足够细粒度的时序记录以支撑趋势分析，另一方面又强调面向现场使用的简洁界面，"
        "使管理人员能够快速定位风险站点、查看空间分布并开展后续处理。"
    ),
    ("subsection", "1.2.2 国内研究现状"),
    (
        "paragraph",
        "国内水质监测研究多与智慧渔业、流域治理、数字水务和环境教学实验等场景相关。近年来，基于 STM32、ESP 系列模组或其他嵌入式平台的终端方案较为常见，"
        "Web 平台通常配套用于实时展示、历史查询和统计分析。相关研究在工程落地方面推进较快，已经形成了较丰富的硬件原型与平台样例。"
    ),
    (
        "paragraph",
        "与此同时，国内许多工程实践更关注单一设备接入或单一数据源展示。当系统同时接入自建设备、公开平台和云侧接口时，字段命名、单位表达、地理信息完整性和数据刷新方式"
        "往往会出现较大差异。如何在不破坏现有代码结构的前提下实现统一建模与断面展示，已经成为平台型课题中的关键问题。"
    ),
    ("subsection", "1.2.3 本课题研究切入点"),
    (
        "paragraph",
        "结合当前项目基础，本文不把研究重点放在某一传感器精度的孤立比较上，而是围绕“终端采集、统一入库、断面展示、辅助分析”这一完整链路展开。"
        "终端部分提供真实采样入口，平台部分负责统一字段模型、快照与历史数据管理，前端部分则将断面状态、趋势变化和地图分布组织为可读界面。"
    ),
    (
        "paragraph",
        "与仅使用公开断面数据的可视化系统相比，本课题保留了自建 STM32 节点作为现场采集入口；与仅关注终端采集的方案相比，平台侧又补充了 DataTransformer、"
        "SensorSnapshot、StationLocation 和 AiInsightLog 等结构，使论文内容能够围绕实际项目形成较完整的系统分析。"
    ),
    ("section", "1.3 研究内容与技术路线"),
    (
        "paragraph",
        "依据开题报告和课题申请表，本文围绕终端设计、平台设计、页面展示和测试分析四项内容展开。第一，设计基于 STM32F103C8 的多参数水质采集终端，"
        "完成 DS18B20、pH、浊度和 TDS 模块接入，以及上报流程实现。第二，构建基于 Django、MySQL 和 Redis/Channels 的平台结构，实现多源数据统一接入、"
        "历史保存、最新快照维护和实时广播。第三，基于 Vue 3、Element Plus、ECharts 和高德地图实现监测概览、综合分析、断面地图、AI 辅助分析与配置页面。"
        "第四，结合运行日志和功能联调结果对系统进行验证。"
    ),
    (
        "table",
        {
            "caption": "表1-1 论文研究内容与项目实现基础",
            "headers": ["研究内容", "对应项目基础", "论文关注点"],
            "rows": [
                ["终端采样与上传", "程序/USER/main.c、esp01_at.c、adc.c、ds18b20.c", "采样流程、换算公式、上报机制"],
                ["平台数据接入", "backend/apps/sensors/views.py、core/data_transformer.py", "统一模型、设备入库、字段映射"],
                ["历史与快照存储", "backend/apps/sensors/models.py、core/realtime_store.py", "双表结构、去重、坐标缓存"],
                ["前端断面展示", "frontend/src/views/*、sensorStore.js", "概览、分析、地图与缓存控制"],
                ["辅助分析与日志", "backend/apps/dashboard/views.py、dashboard/models.py", "上下文组织、问答记录、可追溯性"],
            ],
        },
    ),
    (
        "paragraph",
        "从当前项目状态看，终端正式上行路径采用 HTTP 接口，平台侧已经形成 REST 查询与 WebSocket 广播并存的展示链路。为了保证论文结论与项目状态一致，"
        "文中对实现细节的描述均以现有代码和运行结果为依据，不将未纳入当前实现的功能作为既有成果进行陈述。"
    ),
    ("section", "1.4 论文结构安排"),
    (
        "paragraph",
        "全文共分为六章。第一章阐述课题背景、研究意义、研究现状以及论文的研究内容；第二章给出业务场景、功能需求和可行性分析；第三章从总体层面对系统架构、"
        "数据库与接口进行设计；第四章对终端、平台和展示分析模块展开详细设计；第五章结合现有项目代码说明系统实现；第六章给出测试环境、联调结果与运行分析。"
        "正文之后依次给出结束语、致谢、参考文献和附录。"
    ),
    ("chapter", "第二章 需求分析"),
    ("section", "2.1 业务场景与角色定位"),
    ("subsection", "2.1.1 业务场景分解"),
    (
        "paragraph",
        "本系统主要面向养殖水体巡检、河湖断面监测和教学实验三类场景。对于养殖管理者而言，重点在于持续掌握水温、酸碱度和浊度等参数的变化情况；"
        "对于断面巡检人员而言，重点在于通过地图和筛选条件快速定位目标断面并查看历史趋势；对于教学演示场景而言，系统还承担了展示感知终端、平台入库和可视化联动流程的任务。"
    ),
    (
        "paragraph",
        "虽然三个场景的关注重点并不完全相同，但它们都要求系统能够把采样终端、平台服务和页面展示组织为统一链路。也就是说，系统既要解决现场数据如何稳定进入平台的问题，"
        "也要解决平台结果如何按照断面、区域和时间维度被有效读取的问题。"
    ),
    (
        "table",
        {
            "caption": "表2-1 典型业务场景分解",
            "headers": ["应用场景", "主要对象", "核心关注点", "对应实现基础"],
            "rows": [
                ["养殖水体巡检", "养殖池或单体监测节点", "温度、pH、浊度等指标变化是否及时可见", "STM32 终端、Dashboard 页面、history 接口"],
                ["河湖断面监测", "行政区划或流域下的断面集合", "空间分布、断面筛选和历史趋势对比", "SensorSnapshot、WaterMap 页面、realtime 接口"],
                ["教学实验与演示", "终端到平台的完整链路", "采样、上报、入库和展示是否可联调", "main.c、device_ingest_gateway、overview 页面"],
            ],
        },
    ),
    ("subsection", "2.1.2 角色职责分析"),
    (
        "paragraph",
        "结合当前平台页面与接口功能，系统的主要使用角色可以划分为平台管理人员、监测业务人员和演示观察人员三类。不同角色对同一套数据的关注点不同，"
        "因此需求分析不仅要说明系统能做什么，还要说明这些功能对谁有价值。"
    ),
    (
        "table",
        {
            "caption": "表2-2 系统角色与职责分析",
            "headers": ["角色", "主要关注信息", "典型操作", "期望结果"],
            "rows": [
                ["平台管理人员", "数据源模式、同步状态、运行日志", "切换 auto/manual 模式、触发同步、查看概览", "平台运行状态清晰且可控制"],
                ["监测业务人员", "断面实时数据、历史趋势、地图分布", "筛选区域、查看曲线、点击地图标记", "能够快速定位重点断面并读取变化过程"],
                ["演示观察人员", "终端采样结果与页面联动效果", "观察串口、刷新页面、提问 AI 助手", "能够理解系统完整链路和核心功能"],
            ],
        },
    ),
    (
        "paragraph",
        "角色划分带来的直接要求是：平台不能只提供底层数据接口，而应提供与具体使用方式相匹配的页面与统计结构。例如管理人员更需要概览和模式控制，"
        "业务人员更需要趋势和地图，而演示场景则更关注链路是否完整、结果是否直观。"
    ),
    ("subsection", "2.1.3 业务流程概述"),
    (
        "paragraph",
        "从系统运行过程看，一次典型业务流程包含现场采样、设备上报、统一转换、历史入库、快照更新、页面展示和辅助分析七个环节。"
        "各环节之间并不是松散堆叠关系，而是前后依赖的连续过程。"
    ),
    (
        "code",
        {
            "caption": "图2-1 监测业务流程示意",
            "lines": [
                "传感器采样",
                "   ↓",
                "STM32 进行换算与状态控制",
                "   ↓",
                "ESP-01 通过 HTTP 上报 /sensor",
                "   ↓",
                "后端统一字段转换与历史入库",
                "   ↓",
                "更新 sensor_data_latest 快照",
                "   ↓",
                "前端实时展示、地图展示与 AI 辅助分析",
            ],
        },
    ),
    (
        "paragraph",
        "上述流程说明，需求分析不能只停留在“页面有哪些按钮”或“终端接了哪些传感器”，而需要把数据如何跨越终端、接口、数据库和前端页面的全过程纳入考虑。"
    ),
    ("section", "2.2 功能需求分析"),
    (
        "paragraph",
        "系统功能需求建立在当前项目已经实现的代码基础之上。需求分析的目标不是重新设想一套与现有工程无关的功能集合，而是把已经在项目中落地的能力按照监测场景重新梳理，"
        "使其形成逻辑清楚、层次明确的论文表述。系统核心功能需求如表2-3所示。"
    ),
    (
        "table",
        {
            "caption": "表2-3 系统核心功能需求",
            "headers": ["功能模块", "输入或触发", "处理过程", "输出结果"],
            "rows": [
                ["多参数采集", "DS18B20、pH、浊度、TDS 模块", "ADC/DMA 连续采样、温度读取、经验换算", "温度、pH、浊度、TDS 指标"],
                ["设备上报", "定时任务、WiFi 状态", "ESP-01 建连并调用 /sensor 网关", "设备历史记录与最新快照"],
                ["多源同步", "国家平台、华为云或本地数据", "字段归一化、省市补齐、去重入库", "统一水质监测模型记录"],
                ["实时监测", "页面轮询或 WebSocket 广播", "快照筛选、版本比较、统计汇总", "实时列表和概览卡片"],
                ["历史分析", "站点编号、时间范围", "读取历史表并按时间序列返回", "折线图和趋势分析数据"],
                ["断面地图", "经纬度、水质类别和站点信息", "地图渲染、重复坐标偏移、详情弹窗", "空间分布展示结果"],
                ["辅助分析", "用户问题和快照上下文", "构造上下文、调用模型、记录日志", "文字化分析结果"],
                ["系统配置", "自动或手动模式选择", "更新数据源偏好、控制允许来源", "模式切换结果"],
            ],
        },
    ),
    ("subsection", "2.2.1 采集与上报需求"),
    (
        "paragraph",
        "终端需要完成水温、pH、浊度和 TDS 四类代表性指标采集，其中温度采用数字传感器，另外三类采用模拟量输入。采样过程既要保证周期稳定，"
        "也要能够在串口中输出可读日志，便于开发阶段校验换算结果。"
    ),
    (
        "paragraph",
        "从当前固件实现看，系统已经采用 1 s 周期采样、15 s 周期上报的组织方式，并在主循环中加入心跳翻转与 WiFi 重试逻辑。由此可见，终端层的需求并不仅仅是“采到值”，"
        "还包括周期组织、异常重试和调试可观测性。"
    ),
    ("subsection", "2.2.2 统一接入与数据治理需求"),
    (
        "paragraph",
        "平台侧并非只接收 STM32 终端数据，还承担国家平台数据、华为云侧数据和数据库内部数据的统一组织任务。因此，系统必须在进入业务层之前完成字段名称、"
        "时间格式、区域字段和水质类别表达的归一化处理。"
    ),
    (
        "paragraph",
        "这一需求直接决定了平台需要设置统一转换层、历史表和快照表。若缺少数据治理环节，前端页面将不得不分别适配多种来源，既增加实现复杂度，也不利于后续维护。"
    ),
    ("subsection", "2.2.3 实时监测需求"),
    (
        "paragraph",
        "平台需要同时满足“看最新状态”和“看当前统计”两类需求。前者要求能够快速返回每个站点最近一次有效记录，后者要求能够在相同时间点上对在线数量、均值指标和水质分布进行聚合，"
        "因此系统必须具备面向实时查询的快照结构。"
    ),
    (
        "paragraph",
        "除此之外，监测类页面的刷新逻辑还应考虑无变化数据的处理方式。当前项目已经通过 data_version 和前端缓存机制减少重复下发，这说明需求层面不仅关注“能刷新”，"
        "也关注刷新是否高效、是否适合长时间监测页面持续运行。"
    ),
    ("subsection", "2.2.4 历史分析与回溯需求"),
    (
        "paragraph",
        "除了实时列表，用户还需要查看指定站点在一定时间窗口内的连续变化情况，用于判断某一指标是在短时间内突变，还是在较长时间内逐步变化。"
        "这意味着系统不能只保留最后一条记录，而应保存足够完整的时序历史。"
    ),
    (
        "paragraph",
        "历史分析需求还要求平台在接口层面区分短期历史查询与全量历史查询，前端则需要能够以折线图等形式呈现随时间变化的数据，"
        "从而为断面状态判断提供连续依据。"
    ),
    ("subsection", "2.2.5 断面地图与区域筛选需求"),
    (
        "paragraph",
        "断面展示不仅要求返回经纬度，还要求把站点名称、行政区划、水质类别和短期历史趋势组织为同一页面交互。对于坐标重叠的站点，前端还应具备基本的可点击区分能力，"
        "以保证地图页面具备实际查看价值。"
    ),
    (
        "paragraph",
        "区域筛选需求同样具有实际意义。当前系统支持按省市、流域和名称关键字进行过滤，这意味着平台必须为快照数据补齐 province、city 和 river_basin 等字段，"
        "否则地图展示和列表筛选都难以有效实现。"
    ),
    ("subsection", "2.2.6 辅助分析与配置需求"),
    (
        "paragraph",
        "当前项目已经实现 AI 辅助分析和数据源设置页面，因此系统还需要提供上下文构造、问答记录和模式切换能力。辅助分析并不替代人工判断，"
        "它的作用在于将实时快照转化为更容易理解的文字说明。"
    ),
    (
        "paragraph",
        "数据源设置需求则体现为平台运行策略的可切换性。系统需要根据 auto 和 manual 两种模式控制允许数据源，从而在公开断面展示和自建设备演示两类使用方式之间保持一致的操作入口。"
    ),
    ("section", "2.3 非功能需求分析"),
    (
        "paragraph",
        "除了完成核心业务功能之外，系统还必须满足面向长期展示和多源接入的非功能需求。对于水质监测平台而言，非功能需求直接影响页面是否稳定、接口是否可复用以及终端与平台是否能够持续联动。"
    ),
    (
        "table",
        {
            "caption": "表2-4 非功能需求与设计支撑关系",
            "headers": ["非功能属性", "需求说明", "当前项目中的设计支撑"],
            "rows": [
                ["实时性", "秒级看到最近一次有效监测结果", "快照表、概览接口、WebSocket 广播"],
                ["一致性", "多源数据以统一字段进入平台", "DataTransformer、统一模型字段"],
                ["稳定性", "网络波动和重复同步不影响系统运行", "WiFi 重试、去重写入、data_version 机制"],
                ["可扩展性", "便于接入新页面、新数据源和新指标", "app 分层、数据源模式、统一接口"],
                ["可维护性", "具备日志、调试和状态观察能力", "调试串口、运行日志、AI 问答日志"],
            ],
        },
    ),
    ("subsection", "2.3.1 实时性需求"),
    (
        "paragraph",
        "水质监测平台对实时性的要求主要体现在秒级刷新而非毫秒级控制。当前固件以 1 s 为采样周期、15 s 为上报周期，平台端以快照表支撑实时查询，"
        "能够满足一般监测场景下的展示要求。"
    ),
    (
        "paragraph",
        "实时性需求还体现在页面更新链路的完整性上。终端上报后不仅要完成数据库写入，还要能够让前端在合理时间内得到更新结果，因此快照更新和广播通知属于实时性需求的重要组成部分。"
    ),
    ("subsection", "2.3.2 数据一致性需求"),
    (
        "paragraph",
        "多源接入情况下，字段一致性直接决定前端组件是否能够复用。无论数据来自 STM32 终端、国家平台还是华为云侧，都应在进入数据库之前转换为统一字段集合，"
        "从而保证页面逻辑和统计逻辑建立在同一语义基础上。"
    ),
    ("subsection", "2.3.3 稳定性需求"),
    (
        "paragraph",
        "稳定性要求体现在两个层面。终端侧应能够在网络波动时自动重试并保持主循环运行；平台侧应能够在重复同步时避免无意义的重复写入，在前端频繁刷新时避免整批无变化数据反复下发。"
    ),
    (
        "paragraph",
        "对断面展示系统而言，稳定性还意味着地图、概览和分析页面应共享同一份当前状态来源，避免因不同接口时间不一致而产生结果分裂。当前项目通过快照表统一支撑多个页面，符合这一要求。"
    ),
    ("subsection", "2.3.4 可扩展性需求"),
    (
        "paragraph",
        "可扩展性需求主要体现为两个方面：一是平台未来可以接入新的数据来源或新的传感器指标，二是前端可以继续扩展页面而不破坏现有接口结构。"
        "当前项目已经通过 app 分层、统一模型和配置化数据源模式为后续扩展保留了结构空间。"
    ),
    ("subsection", "2.3.5 可维护性需求"),
    (
        "paragraph",
        "毕业设计项目不仅需要具备运行结果，也需要具备可讲解和可调试特征。当前终端保留调试串口和心跳灯，平台侧保留运行日志、AI 问答日志和可切换的数据源模式，"
        "这些设计都有利于后续定位问题和整理论文材料。"
    ),
    ("section", "2.4 可行性分析"),
    ("subsection", "2.4.1 技术可行性"),
    (
        "paragraph",
        "课题所使用的 STM32F103C8、DS18B20、ESP-01、Django、Vue 3、MySQL 和 Redis/Channels 均为成熟技术方案。"
        "其中 STM32F1 系列能够满足 ADC 采样、串口通信和简单调度任务需求；Django 适合快速构建模型和 REST 接口；Vue 3 与 ECharts 则便于实现监测类页面。"
        "[4][5][6][7][10][11][12][13][14]"
    ),
    ("subsection", "2.4.2 经济可行性"),
    (
        "paragraph",
        "系统所需硬件以开发板和常见传感器模块为主，平台软件主要依赖开源框架和现有运行环境，经费负担较轻。相较于直接采购成套工业监测设备，"
        "本方案更符合本科毕业设计的实验条件，也便于围绕现有项目开展二次实现。"
    ),
    ("subsection", "2.4.3 实施可行性"),
    (
        "paragraph",
        "当前仓库已经具备较完整的软件代码和固件工程，前端页面、后端接口和终端采样流程均有对应实现。开题阶段提出的核心内容已经可以在现有项目中找到实现基础，"
        "因此从实施条件看，课题具备较好的完成条件。"
    ),
    (
        "paragraph",
        "此外，项目已经形成覆盖终端固件、后端服务、前端页面和运行日志的多层材料基础，这使论文编写不需要脱离工程事实另行虚构实现过程，"
        "也使各章节之间能够保持较强的一致性。"
    ),
    ("subsection", "2.4.4 运行可行性"),
    (
        "paragraph",
        "从现有运行结果看，平台已经能够完成多轮同步、历史入库、快照更新、地图展示和辅助分析。运行日志中同步次数、增量写入情况和地理编码成功率均可统计，"
        "说明系统并非停留在静态代码层面，而是已经具备可运行、可观察和可验证的基础。"
    ),
    (
        "paragraph",
        "对于本科毕业设计而言，这种“已有真实项目基础，再围绕其结构化整理与验证”的课题组织方式具备较好的运行可行性，也更有利于答辩阶段对系统结构和实现依据进行清晰说明。"
    ),
    ("chapter", "第三章 系统总体设计"),
    ("section", "3.1 设计目标与原则"),
    (
        "paragraph",
        "系统总体设计遵循“链路完整、数据统一、展示直观、结构清晰”的基本原则。链路完整要求终端采样、平台入库和前端展示相互对应；数据统一要求不同来源在进入平台后转化为同一模型；"
        "展示直观要求页面能够服务于监测场景而非停留在技术演示层面；结构清晰则要求后端应用划分、数据表职责和接口语义保持明确。"
    ),
    (
        "paragraph",
        "基于上述原则，系统在总体设计上采用分层架构和双表存储思路。前者用于划分终端、接入、数据服务和展示分析四个层级；后者用于同时满足实时查询与历史回放两类典型需求。"
    ),
    ("section", "3.2 系统总体架构"),
    ("subsection", "3.2.1 分层结构"),
    (
        "table",
        {
            "caption": "表3-1 系统分层结构",
            "headers": ["层次", "主要组成", "核心职责"],
            "rows": [
                ["感知层", "STM32F103C8、DS18B20、pH/浊度/TDS 模块、ESP-01", "现场采样、指标换算、状态检测和上报"],
                ["接入层", "/sensor 网关、多源同步逻辑、HTTP 请求入口", "接收设备数据并统一送入平台流程"],
                ["数据服务层", "Django、MySQL、Redis/Channels、地理编码缓存", "统一建模、历史保存、快照更新、广播消息"],
                ["展示分析层", "Vue 3、Element Plus、ECharts、高德地图、AI 页面", "实时监测、趋势分析、地图展示和文字化辅助分析"],
            ],
        },
    ),
    (
        "paragraph",
        "该分层结构的特点在于把终端和平台视为同一套系统的不同层次，而不是彼此独立的两套工程。采样字段、数据库字段和页面组件围绕统一语义进行设计，"
        "从而降低后续联调中的理解成本。"
    ),
    ("subsection", "3.2.2 数据流闭环"),
    (
        "paragraph",
        "系统数据流以“采样、上传、入库、展示、反馈”为主线。终端采集到的温度、pH、浊度和 TDS 指标经 ESP-01 上传至 /sensor 网关，"
        "平台完成字段转换后写入历史表与快照表，再通过 REST 接口和 WebSocket 事件向前端提供实时数据与历史数据。对于国家平台和华为云等外部数据源，"
        "平台同样通过统一转换流程写入相同的数据结构，实现多源结果的统一呈现。"
    ),
    (
        "code",
        {
            "caption": "图3-1 系统数据流示意",
            "lines": [
                "DS18B20 / pH / 浊度 / TDS",
                "          ↓",
                "     STM32F103C8 采样与换算",
                "          ↓",
                "      ESP-01 HTTP 上报",
                "          ↓",
                "  /sensor 网关与统一转换层",
                "          ↓",
                "sensor_data + sensor_data_latest",
                "          ↓",
                "Dashboard / Analysis / WaterMap / AI",
            ],
        },
    ),
    ("section", "3.3 模块划分"),
    ("subsection", "3.3.1 终端模块划分"),
    (
        "paragraph",
        "终端由采样模块、温度模块、通信模块、状态指示模块和调试模块构成。采样模块负责 ADC 与 DMA 配置，温度模块负责 DS18B20 初始化与读取，"
        "通信模块负责 ESP-01 的 AT 指令驱动与 HTTP 请求发送，状态指示模块通过 PC13 心跳灯反映主循环运行情况，调试模块则通过 USART1 输出关键日志。"
    ),
    ("subsection", "3.3.2 平台模块划分"),
    (
        "paragraph",
        "后端按照业务职责划分为 sensors、dashboard 和 alerts 三个应用。sensors 负责核心数据模型、设备网关和历史查询；dashboard 负责概览统计、"
        "AI 辅助分析和数据源设置；alerts 保留统一告警结构。前端则围绕 Dashboard、Analysis、WaterMap、AiAssistant 和 Settings 五个页面进行组织。"
    ),
    ("section", "3.4 数据库设计"),
    ("subsection", "3.4.1 核心数据表设计"),
    (
        "table",
        {
            "caption": "表3-2 平台核心数据表",
            "headers": ["数据表", "关键字段", "主要用途"],
            "rows": [
                ["sensor_data", "station_id、recorded_at、temperature、ph 等", "保存历史数据，支撑趋势分析和全量回溯"],
                ["sensor_data_latest", "station_id、recorded_at、province、city 等", "保存站点最新快照，支撑实时查询"],
                ["station_locations", "station_id、station_name、longitude、latitude", "缓存断面坐标，减少重复地理编码"],
                ["alerts", "station_id、alert_type、alert_level、message", "提供统一告警数据结构"],
                ["dashboard_datasourcepreference", "mode、updated_at", "保存当前数据源模式"],
                ["ai_insight_logs", "question、answer、model、duration_ms", "记录 AI 辅助分析过程"],
            ],
        },
    ),
    ("subsection", "3.4.2 历史表与快照表协同设计"),
    (
        "paragraph",
        "sensor_data 与 sensor_data_latest 的双表结构是总体设计中的关键部分。历史表保留时间维度上的完整记录，主要服务于 history 和 all_history 接口；"
        "快照表按 station_id 唯一保存最新状态，主要服务于 overview 和 realtime 接口。通过职责分离，平台既能够支持趋势回溯，又能保持实时查询的较高效率。"
    ),
    (
        "paragraph",
        "在此基础上，station_locations 用于管理坐标缓存，AiInsightLog 用于保留辅助分析记录，DataSourcePreference 则负责维护平台的数据源模式。"
        "这些表共同构成了当前项目的软件数据基础。"
    ),
    ("section", "3.5 接口与通信设计"),
    ("subsection", "3.5.1 设备接入与同步接口"),
    (
        "table",
        {
            "caption": "表3-3 平台主要接口设计",
            "headers": ["接口路径", "方法", "功能说明"],
            "rows": [
                ["/sensor 或 /sensor/", "GET/POST", "兼容 STM32 终端上报入口"],
                ["/api/v1/sensors/device-ingest/", "GET/POST", "设备入库兼容入口"],
                ["/api/v1/sensors/data/realtime/", "GET", "查询最新快照并支持筛选"],
                ["/api/v1/sensors/data/history/", "GET", "查询指定站点的时间窗口历史"],
                ["/api/v1/sensors/data/all_history/", "GET", "查询站点全量历史数据"],
                ["/api/v1/sensors/data/sync_realtime/", "POST", "触发多源同步与统一入库"],
                ["/api/v1/dashboard/overview/", "GET", "返回概览统计、预览列表和规则提示"],
                ["/api/v1/dashboard/ai-insight/", "POST", "基于快照数据生成辅助分析结果"],
            ],
        },
    ),
    ("subsection", "3.5.2 页面访问接口与实时通道"),
    (
        "paragraph",
        "除 REST 接口外，系统还通过 /ws/realtime/ 提供 WebSocket 广播通道，用于在设备入库或同步完成后通知前端页面刷新。"
        "这种设计避免了纯轮询方式在监测场景下产生的高频重复请求，同时保持了实现结构的清晰性。"
    ),
    ("section", "3.6 部署与运行方案"),
    (
        "paragraph",
        "在当前项目中，终端运行于 STM32 固件工程，后端采用 Django 工程组织，前端采用 Vite 构建的 Vue 3 工程组织，数据库使用 MySQL。"
        "系统运行时由终端或同步任务产生数据，后端完成统一入库和快照维护，前端负责展示与交互。由于各模块职责边界明确，系统便于分阶段调试和功能验证。"
    ),
    ("chapter", "第四章 系统详细设计"),
    ("section", "4.1 终端硬件与固件设计"),
    ("subsection", "4.1.1 硬件组成与接口关系"),
    (
        "paragraph",
        "从固件工程可以确认，终端主控芯片为 STM32F103C8，温度采样采用 DS18B20，模拟量采样由 ADC1 与 DMA 完成，ESP-01 通过串口与主控连接。"
        "调试串口和心跳指示灯在代码中均有独立初始化流程，便于开发阶段观察设备状态。终端硬件接口分配如表4-1所示。"
    ),
    (
        "table",
        {
            "caption": "表4-1 终端主要硬件与接口分配",
            "headers": ["模块", "连接引脚或接口", "作用说明"],
            "rows": [
                ["STM32F103C8", "核心控制器", "完成采样调度、串口通信和运行状态控制"],
                ["DS18B20 温度探头", "PB1", "读取水温数字量"],
                ["浊度传感器", "PA6 / ADC Channel 6", "采集浊度模拟电压"],
                ["pH 传感器", "PA7 / ADC Channel 7", "采集 pH 模块模拟电压"],
                ["TDS 传感器", "PA5 / ADC Channel 5", "采集 TDS 模块模拟电压"],
                ["ESP-01 WiFi 模块", "USART2", "完成无线联网与 HTTP 上报"],
                ["调试串口", "USART1", "输出启动和采样日志"],
                ["心跳指示灯", "PC13", "反映主循环运行状态"],
            ],
        },
    ),
    (
        "paragraph",
        "该组合兼顾了原型实现的可获得性和系统演示的完整性。STM32F103C8 拥有足够的 GPIO、ADC 和串口资源，DS18B20 具备读取简单、结果稳定的优点，"
        "ESP-01 则适合在现有项目条件下快速打通终端与平台之间的网络链路。[10][13][14]"
    ),
    ("subsection", "4.1.2 采样换算与指标计算"),
    (
        "paragraph",
        "终端程序将 pH、浊度和 TDS 模块的 ADC 采样结果先换算为电压，再依据当前实验环境下的经验系数转换为业务指标。该方法并不以实验室级标定为目标，"
        "而是服务于系统原型阶段对多参数采样链路的验证。主要换算公式如下。"
    ),
    ("equation", "Vph = ADCph × 3.3 / 4095    （4-1）"),
    ("equation", "pH = -5.7541 × Vph + 16.654    （4-2）"),
    ("equation", "Turb = 2047.19 - 865.68 × Vturb - 200    （4-3）"),
    ("equation", "TDS = 66.71 × Vtds^3 - 127.93 × Vtds^2 + 428.7 × Vtds    （4-4）"),
    (
        "paragraph",
        "为保证页面显示结果的可读性，程序还对部分换算结果进行了简单裁剪。例如 pH 被限制在 0 至 14 范围内，浊度低于 35 时直接置零，TDS 低于 20 时亦置零。"
        "这样的处理能够在原型阶段抑制明显异常值，使平台端更容易得到稳定的展示结果。"
    ),
    ("subsection", "4.1.3 主循环与重试机制"),
    (
        "paragraph",
        "终端采用时间片主循环方式组织任务。main.c 中定义了 SENSOR_SCAN_MS、ESP_RETRY_MS、UPLOAD_PERIOD_MS、HEARTBEAT_MS 和 LOOP_DELAY_MS 等常量，"
        "主循环每次延时 20 ms 后累加运行时间，根据时间差依次触发心跳翻转、采样更新、WiFi 重试和数据上报。相较于引入实时操作系统，这种实现更适合当前功能规模，"
        "也便于在论文中清晰说明执行顺序。"
    ),
    (
        "code",
        {
            "caption": "图4-1 终端主循环流程示意",
            "lines": [
                "系统上电",
                "  ↓",
                "初始化延时、心跳灯、调试串口、ADC/DMA、DS18B20、USART2",
                "  ↓",
                "尝试执行 ESP01_BasicSetup 和 ESP01_ConnectWiFi",
                "  ↓",
                "进入循环：心跳翻转 → 周期采样 → 串口打印 → 条件满足时 HTTP 上报",
                "  ↓",
                "若发送失败，则将 wifi_ready 置零，等待下一轮重试",
            ],
        },
    ),
    ("section", "4.2 平台数据接入与标准化设计"),
    ("subsection", "4.2.1 设备入库网关设计"),
    (
        "paragraph",
        "后端在 URL 配置中同时注册了 /sensor 和 /sensor/ 两个入口，统一指向 device_ingest_gateway，用于兼容当前 STM32 终端的 HTTP 上报方式。"
        "该接口支持 GET 与 POST 两种请求形式，能够根据 query_params 或 request.data 构造统一负载。"
    ),
    (
        "paragraph",
        "在字段组织阶段，_build_device_sensor_payload 会将 temp、ph、turb、tds 等紧凑参数转换为平台内部统一字段集合，并补充站点标识、时间戳和来源类型。"
        "考虑到当前统一模型以 conductivity 作为导电性相关字段，终端上报中的 tds 在现有实现中被映射为 conductivity，以保证平台内部的数据结构一致。"
    ),
    (
        "paragraph",
        "入库时，系统先向 sensor_data 写入一条历史记录，再根据 station_id 在 sensor_data_latest 中创建或更新快照。若新到达数据的 recorded_at 早于已有快照时间，"
        "平台仅保留历史记录而不覆盖最新状态，从而保证快照表始终代表站点当前状态。"
    ),
    ("subsection", "4.2.2 多源字段归一化设计"),
    (
        "paragraph",
        "DataTransformer 是平台数据层中的关键组件。该模块针对 national、huawei 和 database 三类来源分别定义字段提取逻辑，统一输出 station_id、station_name、"
        "province、city、river_basin、water_quality、temperature、ph、dissolved_oxygen、conductivity 和 turbidity 等字段。"
    ),
    (
        "table",
        {
            "caption": "表4-2 多源数据统一映射示例",
            "headers": ["数据来源", "原始字段特征", "归一化后的关键字段"],
            "rows": [
                ["STM32 设备", "temp、ph、turb、tds", "temperature、ph、turbidity、conductivity"],
                ["国家平台", "断面名称、行政区划、水质类别等", "station_id、province、river_basin、water_quality"],
                ["华为云", "云侧水质参数与时间戳", "temperature、ph、dissolved_oxygen、conductivity"],
                ["数据库快照", "统一模型字段", "直接转换为前端展示格式"],
            ],
        },
    ),
    (
        "paragraph",
        "统一转换层的价值在于把数据差异收敛在平台内部，而不是让前端页面分别适配不同来源。同步流程在完成字段归一化后，还会基于断面名称、位置描述和行政区划信息补齐省市字段，"
        "进一步提高后续筛选和地图展示的完整性。"
    ),
    ("subsection", "4.2.3 坐标缓存与地理编码设计"),
    (
        "paragraph",
        "为了满足断面地图展示需求，平台引入了 station_locations 坐标缓存表，并通过批量地理编码流程补齐缺失坐标。对于已经存在缓存的站点，系统直接复用历史结果；"
        "对于缺少坐标信息的站点，系统调用高德地理编码服务，获得经纬度后写回缓存表，并在需要时同步到快照表中。[8]"
    ),
    (
        "paragraph",
        "这种设计一方面减少了对外部地理编码接口的重复依赖，另一方面也使地图展示能力能够随系统运行不断完善。坐标缓存表的引入，使断面地图不再依赖每次页面请求时重新解析地理位置。"
    ),
    ("section", "4.3 实时监测与断面展示设计"),
    ("subsection", "4.3.1 快照读取与版本控制"),
    (
        "paragraph",
        "realtime 接口的核心设计目标是保证返回结果与页面展示节奏匹配。接口读取 sensor_data_latest 中的最新快照，根据区域、断面名称和流域条件完成筛选后，再经 DataTransformer "
        "转换为页面可直接使用的统一格式。为了减少无变化情况下的重复传输，后端进一步计算 data_version 用于比较前后两次数据状态。"
    ),
    ("equation", "data_version = MD5(station_id | timestamp | water_quality | temperature | ph | dissolved_oxygen)    （4-5）"),
    (
        "paragraph",
        "当前端提交的 last_version 与服务端计算结果一致时，接口只返回 changed=false 和必要的元数据，不再重复发送完整列表。"
        "这种机制与前端缓存策略配合后，可显著降低监测页面在稳定时段的无效刷新。"
    ),
    ("subsection", "4.3.2 监测概览页面设计"),
    (
        "paragraph",
        "Dashboard 页面承担平台首页功能。页面顶部提供断面数量、刷新按钮和区域筛选，中部通过统计卡片展示总设备数、在线设备数、告警数量和均值指标，下部以列表形式展示断面最新状态。"
        "页面结构围绕“先看总体，再看个体”展开，符合监测场景的使用习惯。"
    ),
    (
        "paragraph",
        "为了避免重复请求造成页面闪烁，前端通过 sensorStore 保存传感器列表、概览数据、时间戳和版本号，并在 30 s 缓存窗口内复用已有结果。"
        "当设备入库或同步任务触发广播事件时，页面再根据需要主动刷新。"
    ),
    ("subsection", "4.3.3 断面地图与实时推送设计"),
    (
        "paragraph",
        "WaterMap 页面基于高德地图实现断面空间展示。页面首先过滤无坐标数据，再对完全重叠的站点坐标进行轻量偏移，使密集断面仍然保持可点击状态。"
        "用户点击地图标记后，页面显示断面详情并加载短期历史趋势。"
    ),
    (
        "paragraph",
        "后端通过 RealtimeUpdateConsumer 维护 realtime_updates 广播组，broadcast_realtime_event 在设备入库或多源同步成功后向通道层发送更新事件。"
        "前端页面在接收到事件后触发重新拉取，从而形成较轻量的实时联动链路。"
    ),
    ("section", "4.4 综合分析与 AI 辅助设计"),
    ("subsection", "4.4.1 风险排序策略设计"),
    (
        "paragraph",
        "Analysis 页面不仅展示均值指标，还通过启发式规则对风险站点进行排序。当前实现优先关注溶解氧偏低、pH 超出常用参考范围以及水质类别较差的断面。"
        "在页面实现中，风险值采用规则加权方式计算，其表达式可写为："
    ),
    ("equation", "R = 3I(DO < 5) + 2I(pH < 6.5 or pH > 8.5) + 4I(WQ ∈ {Ⅴ, 劣Ⅴ})    （4-6）"),
    (
        "paragraph",
        "该评分并不作为正式水质评价标准，而是用于在页面中快速排序和突出展示监测重点。通过这种方式，用户可以先定位风险较高的断面，再结合历史曲线与地图位置进行进一步分析。"
    ),
    ("subsection", "4.4.2 AI 上下文组织与日志记录"),
    (
        "paragraph",
        "AI 辅助分析模块建立在最新快照数据之上。后台在处理用户问题时，会提取断面总数、在线数量、水质分布和代表性风险站点构成上下文；当问题中出现某个省份名称时，"
        "系统还会额外补充该省的样本数据，以提高回答与当前监测场景的一致性。"
    ),
    (
        "paragraph",
        "为了保证辅助分析过程具备可追溯性，系统将每次问答的 question、answer、model、duration_ms、请求上下文和客户端信息写入 AiInsightLog。"
        "这样既便于后续查看模型响应情况，也保证 AI 功能始终围绕真实监测数据开展解释。"
    ),
    ("section", "4.5 数据源模式与系统配置设计"),
    (
        "paragraph",
        "项目为平台设计了 auto 和 manual 两种数据源模式。自动模式下，平台优先使用国家平台数据并保留本地设备信息；手动模式下，则更强调手动设备数据和指定来源数据的展示。"
        "这一设计使系统既能服务于实际断面展示，也能支持以 STM32 节点为中心的演示与实验。"
    ),
    (
        "paragraph",
        "当前模式由 DataSourcePreference 持久化保存，并通过 Settings 页面供用户切换。由于模式控制位于后端统一入口，前端页面在不修改业务组件的情况下即可切换展示来源，"
        "体现了配置与表现分离的设计思路。"
    ),
    ("chapter", "第五章 系统实现"),
    ("section", "5.1 开发环境与关键技术"),
    (
        "table",
        {
            "caption": "表5-1 系统开发环境与关键技术",
            "headers": ["类别", "配置或技术", "说明"],
            "rows": [
                ["终端硬件", "STM32F103C8 + DS18B20 + pH/浊度/TDS + ESP-01", "完成现场采样与无线接入"],
                ["固件环境", "C 语言 + 标准外设库", "实现 ADC、串口和主循环逻辑"],
                ["后端", "Python 3、Django 4.2.7、DRF、Channels", "实现统一建模、接口和广播"],
                ["数据库", "MySQL", "保存历史表、快照表和配置数据"],
                ["缓存与消息", "Redis/Channels 或内存通道层", "支持实时广播和运行缓存"],
                ["前端", "Vue 3、Vite、Element Plus、ECharts", "实现可视化页面与交互"],
                ["地图服务", "高德地图 JS 与地理编码接口", "实现断面定位和地图展示"],
            ],
        },
    ),
    ("section", "5.2 终端功能实现"),
    ("subsection", "5.2.1 采样与状态控制实现"),
    (
        "paragraph",
        "在 main.c 中，Sensors_Init 负责完成 ADC1_DMA_Config 和 DS18B20_Init 调用，Heartbeat_Init 负责初始化 PC13 心跳灯，DebugUart_Init 用于开启调试串口。"
        "主循环内部依次处理心跳翻转、传感器更新、WiFi 重试和上传任务，使终端能够按照确定的时间节奏运行。"
    ),
    (
        "paragraph",
        "Sensors_Update 对三个模拟量通道和一个温度通道进行集中处理，其中 ADCConvertedValue 数组承担 DMA 搬运后的采样缓存。"
        "采样结果经换算和裁剪后被写入 PH_temp、turbidity_temp、temperature_temp 和 TDS_DAT 等变量，再由 Debug_PrintSensors 输出到串口。"
    ),
    ("subsection", "5.2.2 HTTP 上报实现"),
    (
        "paragraph",
        "esp01_at.c 负责终端通信细节实现。ESP01_BasicSetup 依次发送 AT、ATE0、AT+CWMODE=1、AT+CIPMUX=0 和 AT+CWAUTOCONN=1 等指令，"
        "确保模块处于可联网状态；ESP01_ConnectWiFi 负责加入目标无线网络；ESP01_ReportSensorsUpstream 则根据当前上行模式调用 HTTP 上报函数。"
    ),
    (
        "paragraph",
        "当前固件以 HTTP GET 作为正式上行路径，请求参数包括 ph、tds、turb 和 temp 四项。若发送失败，主循环会将 wifi_ready 置零，随后重新执行联网流程。"
        "这一实现直接对应平台中的 /sensor 设备网关。"
    ),
    ("section", "5.3 后端核心功能实现"),
    ("subsection", "5.3.1 设备入库与快照更新实现"),
    (
        "paragraph",
        "backend/apps/sensors/views.py 中的 device_ingest_gateway 是终端数据进入平台的第一入口。该函数调用 _store_device_sensor_payload 完成历史入库和快照更新，"
        "随后通过 broadcast_realtime_event 向 WebSocket 组广播更新原因和站点编号。"
    ),
    (
        "paragraph",
        "快照更新逻辑体现了平台对“当前状态”与“历史记录”的区分。如果新到达记录的 recorded_at 早于现有快照时间，系统只写入历史表而不覆盖快照，从而避免较旧数据影响当前页面结果。"
    ),
    ("subsection", "5.3.2 多源同步与坐标缓存实现"),
    (
        "paragraph",
        "core/realtime_store.py 中的 sync_realtime_data 负责从国家平台或华为云侧抓取实时数据，并按统一流程完成字段转换、历史写入、快照更新和坐标补齐。"
        "函数通过 station_id 与 recorded_at 组成签名检查已有记录，从而实现增量写入。"
    ),
    (
        "paragraph",
        "对于无坐标站点，平台优先查询 StationLocation 缓存；若缓存不存在，再调用地理编码服务生成经纬度。缓存命中与批量补齐并存的机制，使地图页面既能保持较高可用性，又不会在每次同步时重复请求外部接口。"
    ),
    ("subsection", "5.3.3 概览、辅助分析与实时广播实现"),
    (
        "paragraph",
        "dashboard/views.py 中的 overview 接口直接基于 SensorSnapshot 统计均值指标、在线数量和水质分布，并返回用于首页展示的预览列表。"
        "与直接扫描历史表相比，基于快照表聚合能够更稳定地服务页面实时查询。"
    ),
    (
        "paragraph",
        "同一文件中的 ai_insight 接口根据快照内容组织上下文，并把问答过程写入 AiInsightLog。realtime_events.py 则封装了广播函数，向 realtime_updates 通道组发送统一事件格式。"
        "这样，概览统计、辅助分析和实时刷新三类功能被放在相对独立但能够协同工作的模块中实现。"
    ),
    ("section", "5.4 前端页面实现"),
    ("subsection", "5.4.1 监测概览页面实现"),
    (
        "paragraph",
        "Dashboard.vue 对应首页监测概览。页面通过概览接口加载统计卡片，通过 realtime 接口加载站点列表，并提供省市筛选、关键词检索和手动同步按钮。"
        "界面设计强调信息密度与可读性的平衡，使用户无需切换页面即可掌握当前整体状态。"
    ),
    ("subsection", "5.4.2 综合分析页面实现"),
    (
        "paragraph",
        "Analysis.vue 负责展示均值指标、趋势图和风险站点排序。页面根据快照和历史接口返回的数据构造可视化图表，并以规则化方式突出需要优先关注的站点。"
        "这一页面使平台从“看数据”进一步扩展到“读数据”。"
    ),
    ("subsection", "5.4.3 断面地图页面实现"),
    (
        "paragraph",
        "WaterMap.vue 使用高德地图组件渲染断面位置，并在弹窗中展示站点信息和短期历史趋势。对完全重叠的坐标，页面会进行轻量偏移处理，以保证密集区域内多个断面都能够被点击查看。"
    ),
    ("subsection", "5.4.4 设置与 AI 助手页面实现"),
    (
        "paragraph",
        "Settings.vue 负责数据源模式切换，AiAssistant.vue 负责展示对话式辅助分析结果。两者分别体现了系统的配置能力和解释能力，使平台在实时展示之外具备了面向场景化使用的交互深度。"
    ),
    (
        "table",
        {
            "caption": "表5-2 前端页面实现要点",
            "headers": ["页面", "主要功能", "关键实现点"],
            "rows": [
                ["Dashboard", "概览统计、实时列表、区域筛选", "结合 overview 与 realtime 接口并使用 sensorStore 缓存"],
                ["Analysis", "均值指标、趋势图、风险排序", "基于快照和历史接口组织图表"],
                ["WaterMap", "地图展示、断面详情、历史趋势", "坐标过滤、重叠偏移和详情弹窗"],
                ["AiAssistant", "问题提交、结果展示、日志联动", "调用 ai-insight 接口并展示文本结果"],
                ["Settings", "数据源模式切换", "调用配置接口更新 auto/manual 模式"],
            ],
        },
    ),
    ("section", "5.5 实现效果说明"),
    (
        "paragraph",
        "从整体实现结果看，系统已经形成从终端采样到平台展示的完整链路。终端能够稳定输出多参数数据，平台能够完成历史与快照双轨管理，"
        "前端能够在统一界面中展示实时状态、历史趋势和断面地图，辅助分析模块则将快照数据转化为文字化解释。"
    ),
    (
        "paragraph",
        "这一实现效果表明，当前项目并非单一界面或单一板卡的局部实现，而是一套围绕真实监测流程组织起来的综合系统。"
    ),
    ("chapter", "第六章 系统测试与分析"),
    ("section", "6.1 测试环境与方法"),
    (
        "paragraph",
        "系统测试围绕终端采样、设备入库、前端展示和运行日志四个层面展开。测试方法以功能联调和运行结果核对为主，通过观察串口输出、调用平台接口、"
        "打开前端页面和统计运行日志，验证系统是否达到设计目标。"
    ),
    ("section", "6.2 功能测试"),
    (
        "table",
        {
            "caption": "表6-1 主要功能测试结果",
            "headers": ["测试项", "测试方法", "预期结果", "实际结果"],
            "rows": [
                ["终端采样", "观察串口输出 ADC 值和换算值", "能够周期输出水温、pH、浊度和 TDS", "符合预期"],
                ["设备入库", "终端调用 /sensor 接口", "生成历史记录并更新快照", "符合预期"],
                ["多源同步", "触发 sync_realtime 接口", "完成统一转换并写入数据库", "符合预期"],
                ["实时概览", "访问 Dashboard 页面", "卡片和列表正常刷新", "符合预期"],
                ["历史分析", "选择断面加载历史曲线", "按时间顺序展示趋势图", "符合预期"],
                ["断面地图", "打开 WaterMap 页面并点击标记", "展示断面详情和历史趋势", "符合预期"],
                ["实时广播", "执行设备入库或同步任务", "前端接收到更新事件", "符合预期"],
                ["AI 辅助分析", "提交推荐问题或自定义问题", "返回分析结果并记录日志", "符合预期"],
            ],
        },
    ),
    (
        "paragraph",
        "功能测试结果表明，当前项目的主要业务流程均已打通。终端、后端和前端之间不存在明显的结构性断点，系统具备完成论文所要求展示任务的基础。"
    ),
    ("section", "6.3 联调结果与运行统计"),
    ("subsection", "6.3.1 采集到查询的链路联调"),
    (
        "paragraph",
        "在终端上电并完成网络连接后，串口能够持续输出采样结果，平台在收到 HTTP 请求后写入历史表和快照表，前端页面刷新后能够显示新增站点的最新状态。"
        "这一过程验证了从物理采样到页面查询的全链路连通性。"
    ),
    ("subsection", "6.3.2 运行日志统计结果"),
    (
        "table",
        {
            "caption": "表6-2 运行日志统计结果",
            "headers": ["统计指标", "结果", "说明"],
            "rows": [
                ["同步任务总次数", "465 次", "统计区间为 2026-03-05 至 2026-03-11"],
                ["产生新增写入的任务数", "257 次", "created 大于 0 的任务数"],
                ["无新增写入的任务数", "208 次", "说明重复同步不会持续膨胀历史表"],
                ["单次新增记录上限", "1000 条", "与同步接口默认 count 上限一致"],
                ["有新增任务平均写入量", "349.24 条", "仅统计 created 大于 0 的任务"],
                ["有新增任务中位数", "174 条", "反映常规同步任务规模"],
                ["代表性地理编码批次", "297/300、296/300、296/300", "平均成功率约 98.78%"],
            ],
        },
    ),
    (
        "paragraph",
        "从表6-2可以看出，平台能够在多轮同步过程中保持较稳定的写入节奏。created=0 的任务数量较多，说明平台在识别重复记录方面起到了实际作用，"
        "历史表不会因为重复同步而无序膨胀。"
    ),
    ("subsection", "6.3.3 地图定位与断面展示结果"),
    (
        "paragraph",
        "三次代表性地理编码批处理结果分别达到 297/300、296/300 和 296/300，平均成功率为 98.78%。这表明“坐标缓存 + 批量补齐”的方案能够满足当前地图展示需求，"
        "绝大多数断面都可以在页面中获得稳定的空间位置。"
    ),
    ("subsection", "6.3.4 广播与缓存控制结果"),
    (
        "paragraph",
        "realtime 接口的数据版本机制与 sensorStore 缓存逻辑配合后，能够在无变化情况下直接返回 changed=false。"
        "页面因此不需要在每次轮询时重新渲染整批数据，既降低了前端刷新负担，也保持了监测页面的稳定性。"
    ),
    ("section", "6.4 测试结果分析"),
    (
        "paragraph",
        "综合功能测试与日志统计可以认为，系统已经达到论文设计目标。终端能够形成稳定采样和上报链路，平台能够完成统一入库和双表管理，前端能够围绕断面场景完成概览、分析、地图和辅助解释。"
    ),
    (
        "paragraph",
        "从本科毕业设计评价角度看，当前项目不仅具备代码实现，也具备较完整的系统结构、明确的数据流向和可验证的运行结果，能够较好体现“感知终端 + 数据平台 + 可视展示”一体化设计的工程特征。"
    ),
    ("chapter_no_number", "结束语"),
    (
        "paragraph",
        "本文结合当前项目代码、硬件原型、开题材料和学校论文规范，对基于 STM32 的水质监测系统及断面数据展示平台进行了系统分析与整理。"
        "论文围绕课题背景、需求分析、总体设计、详细设计、系统实现和测试结果六个方面展开，较完整地呈现了终端采样、平台入库、断面展示和辅助分析之间的对应关系。"
    ),
    (
        "paragraph",
        "研究结果表明，终端侧已经完成水温、pH、浊度和 TDS 指标采集与 HTTP 上报，平台侧已经形成统一字段模型、历史与快照双表管理、坐标缓存与实时广播机制，"
        "前端侧已经实现概览、分析、地图、设置和 AI 辅助分析等页面。运行统计进一步验证了系统在多轮同步、增量写入和断面地图展示方面的有效性。"
    ),
    (
        "paragraph",
        "总体而言，本课题已经形成一套结构清晰、链路完整、具有工程实现基础的水质监测与断面数据展示方案，能够满足本科毕业设计对系统性、学术性和应用性的基本要求。"
    ),
    ("chapter_no_number", "致谢"),
    (
        "paragraph",
        "本课题从选题、实现到论文整理过程中，得到了指导教师李俊吉老师的持续指导。老师在系统结构、实现边界和论文写作方面提出了许多具体意见，使课题能够始终围绕真实项目展开，而不是脱离代码和硬件另写一套内容。"
    ),
    (
        "paragraph",
        "同时，感谢在课程学习和项目实践中给予帮助的老师与同学。无论是嵌入式调试、前后端联调还是文档整理，这些交流都为本论文的完成提供了直接支持。"
    ),
    (
        "paragraph",
        "最后，感谢家人在毕业设计阶段给予的理解与支持。正是这些帮助，使我能够较完整地完成本课题的系统实现、测试分析和论文撰写工作。"
    ),
    ("chapter_no_number", "参考文献"),
]


def _replace_nested_text(value, replacements: list[tuple[str, str]]):
    if isinstance(value, str):
        updated = value
        for old, new in replacements:
            updated = updated.replace(old, new)
        return updated
    if isinstance(value, list):
        return [_replace_nested_text(item, replacements) for item in value]
    if isinstance(value, dict):
        return {key: _replace_nested_text(item, replacements) for key, item in value.items()}
    return value


def replace_blocks(blocks: list[tuple[str, object]], replacements: list[tuple[str, str]]) -> list[tuple[str, object]]:
    return [(block_type, _replace_nested_text(payload, replacements)) for block_type, payload in blocks]


def insert_after_heading(
    blocks: list[tuple[str, object]], heading: str, additions: list[tuple[str, object]]
) -> list[tuple[str, object]]:
    result: list[tuple[str, object]] = []
    inserted = False
    for block in blocks:
        result.append(block)
        if block[1] == heading:
            result.extend(additions)
            inserted = True
    if not inserted:
        raise ValueError(f"Heading not found for insertion: {heading}")
    return result


def insert_before_heading(
    blocks: list[tuple[str, object]], heading: str, additions: list[tuple[str, object]]
) -> list[tuple[str, object]]:
    result: list[tuple[str, object]] = []
    inserted = False
    for block in blocks:
        if block[1] == heading and not inserted:
            result.extend(additions)
            inserted = True
        result.append(block)
    if not inserted:
        raise ValueError(f"Heading not found for insertion: {heading}")
    return result


BLOCKS = replace_blocks(
    BLOCKS,
    [
        ("基于 STM32 的水质监测系统及断面数据展示平台", "基于 STM32 的水质监测及断面数据展示系统"),
        ("水质监测系统及断面数据展示平台", "水质监测及断面数据展示系统"),
        ("Water Quality Monitoring System and Section Data Display Platform", "Water Quality Monitoring and Section Data Display System"),
        ("平台入库", "系统入库"),
        ("平台侧", "服务端"),
        ("数据平台 + 可视展示", "数据处理 + 可视展示"),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "1.4 论文结构安排",
    [
        ("subsection", "1.3.1 研究对象与实现范围"),
        (
            "paragraph",
            "本文的研究对象不是抽象的概念性方案，而是当前仓库中已经存在并能够运行的水质监测及断面数据展示系统。"
            "从结构上看，该系统由终端采集单元、后台服务单元和前端展示单元构成：终端负责多参数采样与上报，后台负责统一接入、存储和广播，"
            "前端负责实时监测、趋势分析和地图展示。"
        ),
        (
            "paragraph",
            "论文所讨论的实现范围与当前工程状态保持一致。终端上行链路以 HTTP 网关为正式实现路径，后台采用 Django、MySQL 和 Redis/Channels 组织核心服务，"
            "前端采用 Vue 3 组织页面和交互。凡属代码中已经落地的内容，论文据实展开分析；凡属尚未形成正式实现的内容，不作为既有结果叙述。"
        ),
        ("subsection", "1.3.2 研究目标"),
        (
            "paragraph",
            "本课题的研究目标可以概括为四点：第一，形成稳定的现场采样与上传链路；第二，形成适合多源数据统一管理的后台结构；第三，形成面向断面监测场景的展示页面与交互机制；"
            "第四，通过运行结果验证系统结构的完整性。相比单独强调某一个模块性能，本文更关注整套系统在工程组织上的合理性。"
        ),
        (
            "paragraph",
            "在论文写作上，相应目标表现为：把硬件连接关系、接口结构、数据库职责、页面组织和测试结论说明清楚，使论文内容能够与现有项目代码一一对应。"
            "这也是本科毕业论文区别于简单项目汇报的关键所在。"
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "3.6 部署与运行方案",
    [
        ("subsection", "3.5.3 页面路由与访问关系"),
        (
            "paragraph",
            "前端访问关系由 router/index.js 统一管理。当前工程采用静态导入方式注册 Dashboard、Analysis、WaterMap、AiAssistant 和 Settings 五个页面，"
            "避免懒加载在监测首页首次进入时带来的额外等待。不同页面围绕同一组后端接口组合出不同展示重点，从而构成完整的断面数据展示系统。"
        ),
        (
            "table",
            {
                "caption": "表3-4 页面路由与功能对应关系",
                "headers": ["路由路径", "页面名称", "主要展示内容", "核心依赖接口或状态"],
                "rows": [
                    ["/", "Dashboard", "概览统计、实时列表、区域筛选", "overview、realtime、sensorStore"],
                    ["/analysis", "Analysis", "均值指标、趋势图、风险排序", "realtime、history"],
                    ["/map", "WaterMap", "断面地图、详情弹窗、短期趋势", "realtime、history、高德地图"],
                    ["/ai", "AiAssistant", "对话式辅助分析与推荐问题", "overview、ai-insight、AI 存储状态"],
                    ["/settings", "Settings", "数据源模式切换与可用性提示", "getDataSourceSettings、updateDataSourceSettings"],
                ],
            },
        ),
        ("subsection", "3.5.4 数据访问路径设计"),
        (
            "paragraph",
            "从系统访问路径看，不同接口虽然面向不同页面，但其数据入口并不分散。概览和实时页面依赖快照表，历史页面依赖历史表，辅助分析则在快照摘要基础上构造上下文。"
            "这种访问设计的优势在于：页面各自独立，但其核心数据语义保持一致。"
        ),
        (
            "table",
            {
                "caption": "表3-5 关键数据访问路径",
                "headers": ["访问目标", "主要接口", "核心数据来源", "设计目的"],
                "rows": [
                    ["系统首页概览", "/api/v1/dashboard/overview/", "sensor_data_latest", "快速展示当前总体状态"],
                    ["实时断面列表", "/api/v1/sensors/data/realtime/", "sensor_data_latest", "支撑筛选、列表和地图"],
                    ["历史趋势图", "/api/v1/sensors/data/history/", "sensor_data", "支撑时间窗口分析"],
                    ["全量历史追溯", "/api/v1/sensors/data/all_history/", "sensor_data", "支撑较长时间跨度回放"],
                    ["实时广播", "/ws/realtime/", "broadcast_realtime_event", "降低页面无效轮询成本"],
                    ["辅助分析", "/api/v1/dashboard/ai-insight/", "sensor_data_latest + AiInsightLog", "形成文本化解释与记录"],
                ],
            },
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "4.2 平台数据接入与标准化设计",
    [
        ("subsection", "4.1.4 上报参数组织设计"),
        (
            "paragraph",
            "从 esp01_at.c 可以看出，当前固件把上报报文控制在较为紧凑的范围内，HTTP GET 请求中只包含 ph、tds、turb 和 temp 四个核心参数。"
            "这种设计有利于降低终端构造报文的复杂度，也便于在串口和服务端日志中直接核对请求内容。"
        ),
        (
            "paragraph",
            "由于终端报文保持了较小的字段集合，后台系统在 device_ingest_gateway 中负责补充统一字段结构，例如 recorded_at、station_id、station_name 和 data_source 等。"
            "这样一来，终端实现可以保持简洁，系统内部的数据结构一致性则由服务端负责维护。"
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "4.3 实时监测与断面展示设计",
    [
        ("subsection", "4.2.4 历史入库与快照更新规则"),
        (
            "paragraph",
            "后台系统对历史入库和快照更新采用分离处理策略。对于终端直连数据，_store_device_sensor_payload 会先执行 SensorData.objects.create 完成历史写入，"
            "再根据 station_id 获取或创建 SensorSnapshot 记录；对于外部同步数据，sync_realtime_data 会先构造 station_id 与 recorded_at 组成的签名，"
            "再决定是否写入历史表。"
        ),
        (
            "table",
            {
                "caption": "表4-3 历史入库与快照更新规则",
                "headers": ["处理场景", "历史表处理方式", "快照表处理方式", "设计目的"],
                "rows": [
                    ["终端上报新记录", "直接写入 sensor_data", "若时间更新则覆盖快照", "保留完整采样轨迹并维护当前状态"],
                    ["终端上报旧时间记录", "保留历史记录", "不覆盖现有快照", "避免旧数据影响当前展示"],
                    ["外部数据同步重复记录", "按签名跳过重复写入", "仅在需要时更新快照", "控制历史表膨胀"],
                    ["坐标补齐后回填", "不改动原历史数据", "必要时补充快照经纬度", "提高地图展示完整性"],
                ],
            },
        ),
        (
            "paragraph",
            "这种规则设计保证了历史回溯与当前展示不会互相干扰。历史表强调“记录完整”，快照表强调“状态最新”，两者职责清晰后，页面层的逻辑组织也更容易保持稳定。"
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "4.4 综合分析与 AI 辅助设计",
    [
        ("subsection", "4.3.4 页面缓存与刷新控制设计"),
        (
            "paragraph",
            "frontend/src/stores/sensorStore.js 对实时数据、概览数据和同步任务分别维护了独立的 loading 状态、请求键值和 inflight promise。"
            "其中 realtime、overview 和 sync 三类任务都采用去重执行策略，以避免相同条件下的重复请求并发进入页面。"
        ),
        (
            "paragraph",
            "在缓存策略上，系统把实时数据缓存有效期设置为 30000 ms。若页面在缓存期内再次请求且数据版本未发生变化，则直接复用已有 sensors、total 和 timestamp。"
            "这一设计使系统能够适应监测页面“长时间停留、频繁刷新”的使用方式。"
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "4.5 数据源模式与系统配置设计",
    [
        ("subsection", "4.4.3 AI 页面交互与日志结构设计"),
        (
            "paragraph",
            "AiAssistant.vue 采用“概览摘要 + 对话区域 + 推荐问题”的页面组织方式。页面顶部显示更新时间，左侧显示监测断面数和未恢复告警数，"
            "中部对话区提供推荐问题、消息列表、加载动画和清空按钮，使辅助分析功能不再停留在单次接口调用层面，而成为系统内部的持续交互模块。"
        ),
        (
            "table",
            {
                "caption": "表4-4 AI 问答日志关键字段设计",
                "headers": ["字段名", "类型或含义", "用途说明"],
                "rows": [
                    ["question", "用户问题文本", "记录提问内容"],
                    ["answer", "系统返回文本", "保留回答结果"],
                    ["model", "模型名称", "区分不同模型调用来源"],
                    ["success、error_message", "调用状态与错误信息", "判断问答是否成功"],
                    ["request_payload、response_meta", "请求上下文与响应元数据", "便于回溯上下文组织过程"],
                    ["ip_address、user_agent", "客户端信息", "保留访问来源特征"],
                    ["duration_ms、created_at", "耗时与创建时间", "用于运行分析和审计"],
                ],
            },
        ),
        (
            "paragraph",
            "日志模型的引入使 AI 功能具有明确的可追溯基础。对于本科毕业论文而言，这一点能够说明系统没有把辅助分析当作不可解释的外部黑盒，而是把它纳入了系统内部可记录、可查询的实现结构。"
        ),
    ],
)

BLOCKS = insert_after_heading(
    BLOCKS,
    "5.1 开发环境与关键技术",
    [
        (
            "paragraph",
            "本章实现内容全部对应当前工程目录中的真实代码，而不是根据设计图进行推演。固件相关代码位于程序/USER 目录，后台相关代码位于 backend 目录，"
            "前端相关代码位于 frontend 目录。通过这种目录层级组织，系统实现部分能够在论文中直接对应到具体文件。"
        ),
        ("subsection", "5.1.1 工程目录组织实现"),
        (
            "paragraph",
            "终端工程以 main.c 为主入口，adc.c、ds18b20.c 和 esp01_at.c 分别承担采样、温度读取和联网发送职责；后台工程按 apps 和 core 划分，"
            "前者承载业务模型与视图，后者承载统一转换、同步存储和数据源偏好等公共逻辑；前端工程则按 views、stores、router 和 api 划分。"
        ),
        ("subsection", "5.1.2 关键模块协同实现"),
        (
            "paragraph",
            "前后端协同的关键不在于技术栈堆叠，而在于模块之间的边界是否清楚。当前系统中，终端只负责提供核心指标值，服务端负责完成统一字段转换与持久化，"
            "页面只面向统一接口取数。这样的实现方式降低了前后端耦合，也使后续测试更容易定位问题来源。"
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "5.5 实现效果说明",
    [
        ("subsection", "5.4.5 路由与状态管理实现"),
        (
            "paragraph",
            "router/index.js 中以静态导入方式注册了五个核心页面，分别对应“/”“/analysis”“/map”“/ai”和“/settings”五条访问路径。"
            "监测首页、分析页和地图页围绕同一组实时数据形成不同视图，而 AI 和设置页面则分别承担解释与配置职责。"
        ),
        (
            "table",
            {
                "caption": "表5-3 页面路由与状态管理实现要点",
                "headers": ["页面", "路由路径", "核心状态或存储", "实现特点"],
                "rows": [
                    ["Dashboard", "/", "sensorStore.cache.overview", "首页统计与实时列表共享缓存"],
                    ["Analysis", "/analysis", "sensorStore.cache.sensors", "基于快照与历史数据生成分析视图"],
                    ["WaterMap", "/map", "实时数据 + 地图局部状态", "依据坐标和详情请求完成空间展示"],
                    ["AiAssistant", "/ai", "aiAssistantStore.messages/draft", "支持推荐问题、对话历史和取消请求"],
                    ["Settings", "/settings", "mode、availability", "保存 auto/manual 模式并反馈数据源可用性"],
                ],
            },
        ),
        ("subsection", "5.4.6 AI 页面与设置页面实现细节"),
        (
            "paragraph",
            "AiAssistant 页面在进入时会先请求 overview 接口以加载摘要数据，再结合 quickPrompts 提供推荐问题，问题提交后调用存储层 ask 方法并滚动到底部。"
            "Settings 页面则在 onMounted 阶段读取当前数据源模式，用户保存后再调用 updateDataSourceSettings 完成持久化更新。"
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "6.2 功能测试",
    [
        ("subsection", "6.1.1 测试判据"),
        (
            "paragraph",
            "本文采用的测试判据并不局限于“页面能够打开”或“接口能够返回 200”。对于该系统而言，更重要的判据是：终端采样值是否能够进入统一数据结构，"
            "历史记录和快照是否保持一致，地图与概览是否引用相同来源的数据，以及缓存和广播逻辑是否真正减少了重复刷新。"
        ),
        ("subsection", "6.1.2 测试覆盖范围"),
        (
            "paragraph",
            "基于上述判据，测试覆盖范围包括终端采样与串口输出、/sensor 网关入库、实时与历史查询接口、概览与地图页面、数据源设置页面、AI 辅助分析页面以及运行日志统计。"
            "通过覆盖这些关键路径，可以对系统链路完整性形成较为充分的验证。"
        ),
    ],
)

BLOCKS = insert_before_heading(
    BLOCKS,
    "6.4 测试结果分析",
    [
        ("subsection", "6.3.5 页面级联与配置验证"),
        (
            "paragraph",
            "在页面级联验证中，Settings 页面切换 auto/manual 模式后，概览与实时页面会依据新的数据源策略重新获取结果，页面顶部可用性标记也会同步更新。"
            "这一结果说明系统配置并非孤立页面，而是能够影响后续数据展示行为。"
        ),
        (
            "table",
            {
                "caption": "表6-3 页面级联验证结果",
                "headers": ["验证项", "验证方式", "结果说明"],
                "rows": [
                    ["模式切换生效", "在 Settings 页面保存 auto/manual 模式", "后台配置更新后页面可重新取到对应数据"],
                    ["AI 摘要联动", "刷新 AI 页面并查看概览摘要", "监测断面数与告警数能与 overview 结果保持一致"],
                    ["地图详情联动", "在 WaterMap 页面点击标记后查看趋势", "详情弹窗与历史查询接口返回一致"],
                    ["缓存复用", "在短时间内重复请求实时数据", "无变化时直接复用已有缓存结果"],
                ],
            },
        ),
    ],
)

APPENDIX_I_EN = """
Continuous water quality monitoring is essential in aquaculture, environmental inspection, and section-based river management. In many routine scenarios, managers do not need raw voltages or register values; they need stable indicators, recent trends, and timely warning signals. When monitoring still depends on manual sampling, the collection interval is usually long, the processing chain is fragmented, and the final record often arrives after the abnormal condition has already expanded. An embedded monitoring terminal connected to a networked platform can shorten this path and transform scattered measurements into a continuous operational data stream.

The system discussed in this appendix follows an integrated design idea. The embedded terminal, the communication path, the database service, and the visualization interface are treated as a complete chain rather than four isolated components. This engineering view is important because many practical problems are produced at the boundary between modules. If the terminal reports unstable field names, the server cannot normalize data efficiently. If the server stores data without distinguishing historical records from latest snapshots, the dashboard will become slow. If the map module cannot obtain reliable coordinates, a visually polished interface still fails to present useful spatial information.

At the terminal side, STM32F103C8 is used as the core controller. This microcontroller is not selected for novelty but for balance. It provides enough GPIO resources, multiple serial interfaces, ADC capability, and a stable development environment for a graduation project prototype. In the current implementation, DS18B20 is connected for temperature acquisition, while pH, turbidity, and TDS modules provide analog outputs. The design goal is not to claim laboratory-grade precision. Instead, it is to build a reproducible embedded path that can sample, convert, package, and upload several representative water quality indicators in a predictable cycle.

Analog water quality modules are easy to wire but difficult to trust without calibration. Their outputs are affected by supply voltage, cable length, probe condition, and water characteristics. For this reason, the firmware does not directly expose ADC counts to the platform. It first converts the sampled values into voltages and then applies empirical calibration formulas. Simple threshold clipping is also used to suppress obviously unreasonable results. This strategy does not replace formal calibration experiments, yet it improves the readability of the uploaded data and prevents the platform from being flooded by meaningless outliers during early-stage verification.

The terminal firmware uses a time-slice main loop instead of a real-time operating system. This decision reduces structural complexity and matches the scale of the current task. Several timers are maintained in software: one for heartbeat indication, one for periodic sensing, one for Wi-Fi retry, and one for upstream reporting. Such an arrangement is easy to inspect in source code and convenient for debugging through serial logs. In a teaching or prototype environment, a clear loop with deterministic periods is often more valuable than a heavy abstraction layer that obscures execution order.

Wireless transmission is implemented through an ESP-01 module driven by AT commands. The firmware first checks whether the module responds to basic commands, then configures station mode, disables command echo, and joins the local wireless network. After the connection becomes available, the terminal opens a TCP socket and sends an HTTP GET request carrying temperature, pH, turbidity, and TDS-related values. This choice keeps the communication chain short and observable. HTTP is not the only possible option, and it is not the final target of the long-term design, but it is very suitable for proving that the terminal and platform can already exchange real measurements through a stable and testable interface.

The communication design also keeps room for future migration. In the current code, a switch constant for MQTT mode is reserved even though HTTP remains the active path. This reflects a practical engineering trade-off. For a graduation project, finishing a full end-to-end chain with manageable complexity is more important than enabling every advanced feature at once. Once the basic path works, later iterations can add local buffering, retransmission, message confirmation, and topic-based uplink mechanisms without discarding the existing acquisition logic.

On the server side, the first important design goal is field normalization. Data may come from the STM32 terminal, the national public monitoring platform, a cloud-side source, or the local database itself. These inputs use different field names, different optional attributes, and sometimes different semantic expressions for the same concept. The DataTransformer component is therefore not just a convenience layer; it is the central place where heterogeneous payloads are converted into a stable internal model. Without this step, every API and every page would have to deal with source-specific branches, and the maintenance cost would rise rapidly.

The device ingestion gateway provides a concrete example of this idea. The terminal uploads compact parameters such as temp, ph, turb, and tds. The backend converts them into a unified payload that contains station identity, geographic hints, normalized metric names, timestamp, and source type. In the present model, the TDS channel is temporarily mapped into the unified conductivity field so that the downstream storage and visualization code can continue to use one shared structure. This is a deliberate compromise for integration. The important point is that the platform receives device data in the same conceptual shape used by other monitoring sources.

Database design is another key part of the platform. The current project separates the history table from the latest snapshot table. The history table preserves all valid records and supports trend analysis, while the snapshot table keeps only the newest record for each station and supports dashboard queries. This split is simple but highly effective. If a real-time page always scans a long history table, response latency increases as data volume grows. If the system keeps only the latest row, long-term analysis becomes impossible. By combining both tables, the platform preserves traceability and still maintains fast access to the current state.

Section-level visualization requires more than metric storage; it also requires spatial positioning. In real data sources, section coordinates are often incomplete or inconsistent. The platform addresses this issue by introducing a dedicated coordinate cache table and a geocoding workflow. Existing cached coordinates are reused first, and only unresolved sections are sent to the geocoding service. Successful results are written back into the cache and may also be copied into the latest snapshot table. This design reduces repeated external requests and makes the map view progressively more complete as the system runs.

Real-time display is not solved only by WebSocket pushing. It also depends on how repeated polling is handled when the data has not changed. The project computes a version signature from several load-bearing fields such as station identifier, timestamp, water quality class, temperature, pH, and dissolved oxygen. When the front end sends the last known version and the current snapshot is unchanged, the backend returns only a small response that marks the data as unchanged. This mechanism reduces unnecessary transmission, avoids repeated rendering, and fits the practical rhythm of monitoring pages in which most refresh cycles contain no meaningful update.

The visualization layer is organized around five pages: dashboard, analysis, map, AI assistant, and settings. The dashboard focuses on current snapshots and quick filters. The analysis page highlights trends, averages, and a lightweight risk ranking. The map page turns station records into clickable markers and connects spatial distribution with history review. The settings page controls the active data-source mode. The AI assistant page uses the latest snapshot as context and produces textual explanations or suggestions. Although each page has its own interaction style, all of them depend on the same normalized storage model and therefore remain consistent in meaning.

Risk interpretation in the current platform is intentionally conservative. The system uses rule-based ranking for low dissolved oxygen, abnormal pH, and poor water quality classes, while the AI module is positioned as an assistant rather than an automatic decision maker. This separation is important in engineering practice. Rules are transparent, reproducible, and easy to audit. AI outputs are useful for interpretation, summarization, and recommendation, but they can be unstable if context is incomplete or external model services are unavailable. Recording every question, answer, model name, and execution time helps the platform keep this auxiliary function traceable.

Another point worth noting is maintenance. A water quality terminal is not a closed mathematical object; it is a physical node deployed in a changing environment. Probes age, connectors loosen, local power quality fluctuates, and wireless signals vary over time. Because of this, maintainability should be considered part of the architecture rather than an afterthought. The current design supports serial debug output, explicit Wi-Fi retry logic, and readable configuration constants. These details may appear small when compared with sensors or databases, yet they strongly influence whether the system can be diagnosed and restored quickly after a field-side abnormal condition.

Filtering and query organization on the server side also play a practical role. Monitoring users rarely ask for all data in the same way every time. Sometimes they focus on one station, sometimes on one province, and sometimes on a river basin or a short time window. The realtime, history, and all_history interfaces therefore expose different levels of granularity. The realtime endpoint is optimized for current snapshots, while the history endpoints preserve chronological traces. This division helps each request remain simple in meaning. It also avoids forcing one endpoint to carry too many unrelated responsibilities, which would make both backend logic and front-end integration harder to control.

The map page reflects a useful design lesson for section-based monitoring systems. A map is not valuable only because it contains coordinates; it becomes useful when position, time, and indicator values are presented together. The current implementation does more than place markers. It filters out records without coordinates, offsets overlapping markers so that dense points remain clickable, and connects each popup to a short history chart. This means the map is not merely decorative. It functions as a compact spatial entry for diagnosis: users can move from regional distribution to station detail without switching mental context or manually searching another page.

The AI assistant should also be understood in a limited but meaningful way. In this project, the assistant does not invent its own monitoring dataset. It receives structured context assembled from the snapshot table and uses that context to explain risk points, summarize system state, or answer operation-oriented questions. This constraint is necessary. Once a model is allowed to answer without controlled context, the explanation may become fluent but detached from the actual monitoring record. The platform therefore logs request payloads, model names, response status, and duration, which makes the AI layer observable and keeps it aligned with the real dataset instead of turning it into an opaque black box.

Finally, the system has educational value beyond the immediate monitoring task. It demonstrates how an undergraduate project can move from separate technical exercises to a complete engineering workflow. The embedded part shows acquisition and communication, the backend part shows normalization and persistence, and the front-end part shows interaction and decision support. When these pieces are connected, students can observe that engineering quality is determined not only by whether each module runs independently, but also by whether the boundaries between modules are clearly defined. In this sense, the project serves both as a water quality platform and as a practical case of end-to-end IoT system integration.

Runtime observations from the existing project provide useful evidence for the design. The synchronization log records hundreds of scheduled runs, and only a part of them generate new records, which indicates that duplicate insertion is effectively controlled. Representative geocoding batches also show a high success rate, demonstrating the value of the coordinate cache plus batch completion strategy. These observations do not prove that the system is finished, but they do show that the current architecture is not merely theoretical. It has already supported repeated data synchronization, incremental storage, section positioning, and front-end presentation in a stable manner.

From an engineering perspective, the most valuable achievement of this system is not a single algorithm or a single interface. Its real value lies in turning embedded acquisition, server-side normalization, structured storage, map-based section display, and interactive analysis into one coherent workflow. The existing implementation already demonstrates a complete and teachable path from physical sensing to visual decision support, which is exactly what a comprehensive undergraduate project should emphasize.
""".strip().split("\n\n")

APPENDIX_I_CN = """
连续水质监测在智慧养殖、环境巡检和断面化管理中具有基础意义。在很多日常场景下，管理人员并不需要直接看到原始电压或寄存器值，他们真正关心的是稳定的监测指标、近期变化趋势以及异常预警信号。如果监测仍然主要依赖人工取样，那么采样间隔通常较长，处理链路也较为分散，最终记录往往在异常已经扩散后才到达管理端。将嵌入式监测终端与网络平台连接起来，可以显著缩短这一链路，把零散测量转化为连续的数据流。

本附录讨论的系统遵循一体化设计思路。嵌入式终端、通信链路、数据库服务和可视化界面被视为一条完整的工作链，而不是四个彼此独立的模块。这种工程视角非常重要，因为实际问题往往出现在模块边界处。如果终端上传的字段名不稳定，服务端就难以高效归一化；如果服务端没有区分历史记录和最新快照，监测首页就会逐渐变慢；如果地图模块拿不到可靠坐标，即使界面再美观，也无法提供真正有用的空间信息。

在终端侧，系统选择 STM32F103C8 作为核心控制器。这一选择并不是为了追求新颖，而是为了追求平衡。该芯片拥有足够的 GPIO 资源、串口接口、ADC 能力以及较成熟的开发环境，适合毕业设计原型实现。当前方案中，DS18B20 用于采集温度，pH、浊度和 TDS 模块提供模拟量输出。设计目标并不是宣称达到实验室级精度，而是建立一条可复现的嵌入式数据通路，使终端能够在稳定周期内完成采样、换算、封包与上传。

模拟式水质传感器虽然接线方便，但如果缺少标定，其结果并不可靠。传感器输出会受到供电、电缆长度、探头状态以及水体特性的共同影响。因此，固件并没有把 ADC 计数值直接上传到平台，而是先把采样值换算为电压，再使用经验标定公式将其转换为业务指标。程序同时加入了简单的阈值裁剪，用来抑制明显不合理的结果。这样的处理不能替代正式标定实验，但在前期验证阶段能够提升上传数据的可读性，避免平台被大量无意义异常值干扰。

终端固件采用时间片主循环而不是实时操作系统。这一决策降低了结构复杂度，也符合当前任务规模。程序在软件中维护了多个定时器，分别用于心跳指示、周期采样、Wi-Fi 重试以及上行上传。这样的组织方式便于从源代码中直接观察执行顺序，也方便通过串口日志定位问题。在教学和原型环境下，一个执行周期明确、逻辑顺序清晰的主循环，往往比层次过深的抽象框架更有价值。

无线传输部分由 ESP-01 模块和 AT 指令链路完成。固件首先检查模块是否能响应基础指令，然后配置为 STA 模式、关闭回显并连接本地无线网络。当网络可用后，终端建立 TCP 连接，通过 HTTP GET 请求上传温度、pH、浊度和 TDS 等数据。选择这一方案的原因在于链路短、状态清晰、易于验证。HTTP 当然不是唯一选择，也不是长期架构的最终目标，但对于证明终端和平台已经能够通过统一接口交换真实测量数据而言，它非常合适。

当前通信设计也为后续扩展预留了空间。虽然实际运行路径仍然是 HTTP，但代码中已经保留了 MQTT 模式的切换常量。这体现了一个务实的工程取舍。对于毕业设计来说，优先完成一条端到端可跑通、可调试、可展示的链路，通常比一次性引入全部高级功能更重要。当基础路径稳定以后，再逐步加入本地缓存、失败补发、消息确认和主题订阅机制，系统演进就会更平滑。

在服务端，首先需要解决的问题是字段归一化。数据既可能来自 STM32 终端，也可能来自国家公开平台、云侧接口或本地数据库本身。不同来源使用的字段名不同，可选属性不同，对同一语义的表达方式也不完全一致。因此，DataTransformer 并不是一个可有可无的辅助层，而是整个后端处理中的关键组件。它把异构输入转换为稳定的内部模型。如果没有这一层，几乎每一个接口和页面都要面对来源差异，系统维护成本会迅速增加。

设备入库网关就是这一思路的直接体现。终端上传的参数较为紧凑，例如 temp、ph、turb 和 tds。后端在接收后，会把这些参数组织成统一负载，其中包含站点身份、地理线索、规范化指标名、时间戳和来源类型。在当前模型中，TDS 通道暂时映射到统一结构中的 conductivity 字段，以便下游存储和展示逻辑继续共享同一套模型。这是一种面向集成的折中方案，关键在于平台收到的设备数据能够与其他监测来源在概念上保持一致。

数据库设计同样是平台实现中的关键部分。当前项目把历史表和最新快照表分开管理。历史表保留全部有效记录，用于趋势分析；快照表只保留每个站点的最新状态，用于首页查询。这种拆分方式虽然简单，却十分有效。如果实时页面总是扫描长时间积累的历史表，随着数据量增长，响应速度必然下降；如果系统只保留最新一条记录，那么历史变化又无法回溯。通过双表协同，平台既保留了可追溯性，又保持了当前状态查询的效率。

断面级展示不仅依赖指标存储，还依赖空间定位。现实数据源中的断面经纬度往往并不完整，甚至同一断面名称的写法也可能不统一。平台为此增加了独立的坐标缓存表和地理编码流程。程序首先复用已有缓存，只有在缓存不存在时才调用外部地理编码接口；新的坐标结果会回写缓存表，并在需要时同步到最新快照表中。这样的设计减少了对外部服务的重复请求，也使地图展示随着系统运行逐步变得完整。

实时展示并不只是依赖 WebSocket 推送，还取决于如何处理“数据未变化”的轮询场景。当前项目会根据站点编号、时间戳、水质类别、水温、pH 和溶解氧等关键字段计算版本摘要。当客户端携带的版本号与服务端当前版本一致时，后端只返回一个很小的响应，表明当前数据未变化，而不再重复发送完整快照列表。这样做可以减少无效传输，避免重复渲染，也更符合监测类页面的大多数刷新周期其实并没有变化的实际情况。

前端展示层围绕五个页面组织：监测概览、综合分析、断面地图、AI 助手和系统设置。概览页面强调当前快照和快速筛选；分析页面突出趋势曲线、均值指标与风险排序；地图页面把断面记录转换为可点击标记，并把空间分布和历史详情结合起来；设置页面负责切换数据源模式；AI 助手页面则利用最新快照生成文字化解释和建议。虽然页面功能不同，但它们都建立在同一套归一化数据模型之上，因此语义保持一致。

当前平台对风险解释采取较为保守的策略。系统使用规则化方法对低溶解氧、异常 pH 和较差水质类别进行排序，而 AI 模块被明确定位为辅助工具，而不是自动决策者。这种分工在工程实践中很有必要。规则透明、可重复、可审计；AI 输出在解释、总结和建议方面确实有帮助，但一旦上下文不完整或外部模型服务不稳定，其结果就可能偏移。记录每次问答的问题、回答、模型名称和耗时，可以保证这一辅助模块在使用过程中保持可追溯。

另一个值得强调的点是可维护性。水质监测终端并不是一个封闭的数学对象，而是部署在变化环境中的物理节点。探头会老化，连接器会松动，供电质量会波动，无线信号也会随时间变化。因此，可维护性应当被视为体系结构的一部分，而不是上线之后才考虑的问题。当前方案保留了调试串口输出、明确的 Wi-Fi 重试逻辑和可读性较强的配置常量。这些细节看上去不如传感器或数据库“显眼”，但它们直接决定系统在现场异常发生后能否被迅速定位和恢复。

服务端的筛选与查询组织同样具有实际意义。监测使用者不会在每次操作中都以同一种方式获取数据。有时他们只关心一个站点，有时关注一个省份，有时则关注某个流域或者较短时间窗口。因此，项目把 realtime、history 和 all_history 三类接口设计为不同粒度的查询入口。realtime 面向当前快照，history 面向一定时间范围内的变化轨迹，all_history 则面向站点全量时序。这样的划分使每类接口的语义更清晰，也避免把过多职责塞进一个统一入口，进而增加前后端两侧的复杂度。

地图页面还反映出一个面向断面监测系统的重要设计经验。地图的价值并不只是“有坐标”，而在于位置、时间和指标数值能否被同时呈现。当前实现不只是简单打点。它会先过滤无坐标数据，再对完全重叠的标记做偏移处理，保证密集站点仍然可以点击，并把每个弹窗与短周期历史曲线关联起来。这样一来，地图就不再是装饰性界面，而成为一个紧凑的空间诊断入口，用户可以从区域分布直接过渡到站点详情，而不需要重新切换页面或再次手动搜索。

AI 助手同样需要在有限而明确的范围内理解。它并不会自行生成监测数据，而是使用由快照表组织好的结构化上下文，对风险点进行解释、对系统状态进行总结，或者回答与运维相关的问题。这种限制非常必要。只要模型在没有受控上下文的情况下自由回答，输出就可能变得流畅却脱离真实记录。正因为如此，平台才会记录请求负载、模型名称、响应状态和耗时，以便让 AI 层保持可观测，并始终围绕真实数据工作，而不是成为一个无法回溯的黑盒。

最后，这套系统的意义并不限于当前监测任务本身，它还具有较强的教学价值。它展示了本科项目如何从分散的技术练习过渡到完整工程流程。嵌入式部分体现采样与通信，后端部分体现归一化和持久化，前端部分体现交互和决策支持。当这些环节真正连接起来以后，学生就能直观理解：工程质量不仅取决于每个模块能否独立运行，更取决于模块边界是否清晰、数据语义是否一致。从这个角度看，本项目既是一套水质监测平台，也是一个面向物联网系统集成的完整实践案例。

从现有项目运行情况看，系统设计已经获得了一定验证。同步日志记录了大量周期任务，而只有部分任务产生新增记录，这说明重复入库已得到控制；地理编码批处理也表现出较高成功率，说明“坐标缓存 + 批量补齐”的策略是有效的。这些现象并不能说明系统已经完全成熟，但至少表明当前架构不只是纸面设计，而是真正支撑了多次数据同步、增量存储、断面定位和前端展示。

从工程角度评价，这套系统最有价值的地方并不是某一个单独的算法或界面，而是它把嵌入式采集、服务端归一化、结构化存储、断面地图展示和交互式分析组织为一条连贯工作流。现有实现已经完整展示了从物理采样到可视化决策支持的全过程，而这正是综合型本科毕业设计最需要体现的内容。
""".strip().split("\n\n")

APPENDIX_BLOCKS = [
    ("chapter_no_number", "附录Ⅰ 英文资料翻译"),
    ("appendix_subtitle", "英文原文"),
    *[("appendix_en", paragraph) for paragraph in APPENDIX_I_EN],
    ("appendix_subtitle", "中文翻译"),
    *[("appendix_cn", paragraph) for paragraph in APPENDIX_I_CN],
    ("chapter_no_number", "附录Ⅱ 程序代码"),
    (
        "appendix_cn",
        "附录Ⅱ保留与本课题直接相关的代表性程序代码，分别对应终端侧采样与上传、服务端设备入库以及前端实时缓存控制。"
        "代码及注释均采用五号 Times New Roman、单倍行距排版。"
    ),
    (
        "code",
        {
            "caption": "附录Ⅱ-1 终端主循环核心代码（程序/USER/main.c）",
            "lines": [
                "while (1)",
                "{",
                "    if ((app_ms - last_heartbeat_ms) >= HEARTBEAT_MS)",
                "    {",
                "        last_heartbeat_ms = app_ms;",
                "        Heartbeat_Toggle();",
                "    }",
                "",
                "    if ((app_ms - last_sensor_ms) >= SENSOR_SCAN_MS)",
                "    {",
                "        last_sensor_ms = app_ms;",
                "        Sensors_Update();",
                "        Debug_PrintSensors();",
                "    }",
                "",
                "    if ((app_ms - last_retry_ms) >= ESP_RETRY_MS)",
                "    {",
                "        last_retry_ms = app_ms;",
                "        if (esp_ready == 0U) esp_ready = ESP01_BasicSetup();",
                "        else if (wifi_ready == 0U) wifi_ready = ESP01_ConnectWiFi();",
                "    }",
                "",
                "    if ((esp_ready != 0U) && (wifi_ready != 0U) && ((app_ms - last_upload_ms) >= UPLOAD_PERIOD_MS))",
                "    {",
                "        last_upload_ms = app_ms;",
                "        if (ESP01_ReportSensorsUpstream(PH_temp, turbidity_temp, temperature_temp, TDS_DAT) == 0U)",
                "            wifi_ready = 0U;",
                "    }",
                "",
                "    DelayMs(LOOP_DELAY_MS);",
                "    app_ms += LOOP_DELAY_MS;",
                "}",
            ],
        },
    ),
    (
        "code",
        {
            "caption": "附录Ⅱ-2 HTTP 上报函数关键代码（程序/USER/esp01_at.c）",
            "lines": [
                "snprintf(",
                "    request,",
                "    sizeof(request),",
                "    \"GET %s?ph=%ld.%01ld&tds=%lu&turb=%lu&temp=%ld.%01ld HTTP/1.1\\r\\n\"",
                "    \"Host: %s\\r\\n\"",
                "    \"Connection: close\\r\\n\"",
                "    \"\\r\\n\",",
                "    ESP01_HTTP_PATH,",
                "    ph_x10 / 10,",
                "    ph_x10 >= 0 ? (ph_x10 % 10) : -(ph_x10 % 10),",
                "    (unsigned long)tds_u32,",
                "    (unsigned long)turb_u32,",
                "    temp_x10 / 10,",
                "    temp_x10 >= 0 ? (temp_x10 % 10) : -(temp_x10 % 10),",
                "    ESP01_HTTP_HOST",
                ");",
                "",
                "request_len = (uint16_t)strlen(request);",
                "snprintf(cmd, sizeof(cmd), \"AT+CIPSEND=%u\", (unsigned int)request_len);",
                "if (!ESP01_SendCommand(cmd, \">\", 5000U))",
                "    return 0U;",
                "",
                "USART2_ClearRxBuffer();",
                "USART2_SendBuffer((const uint8_t *)request, request_len);",
                "if (!ESP01_WaitReply(\"SEND OK\", 8000U))",
                "    return 0U;",
            ],
        },
    ),
    (
        "code",
        {
            "caption": "附录Ⅱ-3 设备入库网关关键代码（backend/apps/sensors/views.py）",
            "lines": [
                "@api_view([\"GET\", \"POST\"])",
                "def device_ingest_gateway(request):",
                "    payload = request.query_params if request.method == \"GET\" else (request.data or {})",
                "    stored_payload = _store_device_sensor_payload(payload)",
                "",
                "    broadcast_realtime_event(",
                "        {",
                "            \"reason\": \"device_ingest\",",
                "            \"source\": stored_payload.get(\"data_source\"),",
                "            \"station_id\": stored_payload.get(\"station_id\"),",
                "        }",
                "    )",
                "    return Response(",
                "        {",
                "            \"code\": 200,",
                "            \"message\": \"success\",",
                "            \"data\": {",
                "                \"station_id\": stored_payload.get(\"station_id\"),",
                "                \"station_name\": stored_payload.get(\"station_name\"),",
                "                \"recorded_at\": stored_payload.get(\"recorded_at\").isoformat(),",
                "                \"data_source\": stored_payload.get(\"data_source\"),",
                "            },",
                "        }",
                "    )",
            ],
        },
    ),
    (
        "code",
        {
            "caption": "附录Ⅱ-4 实时数据缓存控制关键代码（frontend/src/stores/sensorStore.js）",
            "lines": [
                "const result = await fetchFn(cache.dataVersion.value)",
                "if (result?.code === 200) {",
                "  const changed = result.data?.changed !== false",
                "  const dataVersion = result.data?.data_version || cache.dataVersion.value",
                "  if (!changed && hasCacheData) {",
                "    cache.timestamp.value = Date.now()",
                "    cache.dataVersion.value = dataVersion",
                "    return {",
                "      code: 200,",
                "      data: {",
                "        sensors: cache.sensors.value,",
                "        total: cache.total.value,",
                "        timestamp: cache.timestamp.value,",
                "        fromCache: true,",
                "        data_version: cache.dataVersion.value,",
                "        changed: false",
                "      }",
                "    }",
                "  }",
                "}",
            ],
        },
    ),
    ("chapter_no_number", "附录Ⅲ 正文内过长的公式推导、重复性的数据和过大的图表、及说明等"),
    (
        "appendix_cn",
        "附录Ⅲ给出正文中不宜在主体部分展开列出的补充数据与说明，主要包括运行日志统计补充、关键数据表字段说明以及终端与平台之间的字段映射关系。"
    ),
    (
        "table",
        {
            "caption": "附表Ⅲ-1 运行日志统计补充说明",
            "headers": ["统计项", "数值", "补充说明"],
            "rows": [
                ["同步任务总次数", "465", "统计区间为 2026-03-05 至 2026-03-11"],
                ["有新增写入的任务数", "257", "created > 0"],
                ["无新增写入的任务数", "208", "created = 0"],
                ["有新增任务平均写入量", "349.24", "仅统计 created > 0 的任务"],
                ["有新增任务中位数", "174", "反映常规任务规模"],
                ["单次新增上限", "1000", "与同步任务默认 count 上限一致"],
                ["代表性地理编码批次", "297/300、296/300、296/300", "平均成功率约 98.78%"],
            ],
        },
    ),
    (
        "table",
        {
            "caption": "附表Ⅲ-2 sensor_data_latest 关键字段补充说明",
            "headers": ["字段名", "类型或含义", "说明"],
            "rows": [
                ["station_id", "站点唯一标识", "快照表按该字段唯一更新"],
                ["station_name", "站点/断面名称", "用于列表、地图和历史查询展示"],
                ["province、city、river_basin", "区域字段", "用于筛选和统计"],
                ["longitude、latitude", "经纬度", "支撑地图标点展示"],
                ["temperature、ph、dissolved_oxygen", "基础水质指标", "用于实时监测与风险分析"],
                ["conductivity、turbidity", "扩展指标", "兼容设备与公开数据源"],
                ["water_quality", "水质类别", "用于颜色映射和风险排序"],
                ["recorded_at", "监测时间", "判断在线状态与数据新鲜度"],
                ["updated_at", "快照更新时间", "反映数据库最近更新时间"],
            ],
        },
    ),
    (
        "table",
        {
            "caption": "附表Ⅲ-3 终端字段与系统统一字段映射关系",
            "headers": ["终端上传字段", "系统统一字段", "说明"],
            "rows": [
                ["temp", "temperature", "温度值，单位为摄氏度"],
                ["ph", "ph", "酸碱度指标"],
                ["turb", "turbidity", "浊度指标"],
                ["tds", "conductivity", "当前项目中为兼容统一模型而暂时映射"],
                ["station_id/device_id", "station_id", "站点或设备标识"],
                ["station_name/device_name", "station_name", "站点或设备名称"],
                ["timestamp/recorded_at", "recorded_at", "设备采样时间"],
            ],
        },
    ),
]


def set_rfonts(run, zh: str, latin: str) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(
    run,
    *,
    zh: str = "宋体",
    latin: str = "Times New Roman",
    size: float = 12,
    bold: bool = False,
    italic: bool = False,
) -> None:
    set_rfonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def clear_paragraph(paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def add_field(paragraph, instruction: str, default_text: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    if default_text:
        text_run = paragraph.add_run(default_text)
        style_run(text_run, zh="宋体", latin="Times New Roman", size=10.5)
    run_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def set_page_number_format(section, fmt: str, start: int | None = None) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    pg_num_type = OxmlElement("w:pgNumType")
    pg_num_type.set(qn("w:fmt"), fmt)
    if start is not None:
        pg_num_type.set(qn("w:start"), str(start))
    sect_pr.append(pg_num_type)


def set_header_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")


def configure_section(section, *, with_header: bool, page_fmt: str | None = None, start: int | None = None) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(36)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(20)
    section.header_distance = Mm(25)
    section.footer_distance = Mm(18)

    if page_fmt:
        set_page_number_format(section, fmt=page_fmt, start=start)

    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    if with_header:
        header_para = section.header.paragraphs[0]
        clear_paragraph(header_para)
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_para.paragraph_format.space_before = Pt(0)
        header_para.paragraph_format.space_after = Pt(0)
        set_header_border(header_para)
        run = header_para.add_run(HEADER_TEXT)
        style_run(run, zh="宋体", latin="Times New Roman", size=10.5)

        footer_para = section.footer.paragraphs[0]
        clear_paragraph(footer_para)
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.paragraph_format.space_before = Pt(0)
        footer_para.paragraph_format.space_after = Pt(0)
        add_field(footer_para, "PAGE")
        for run in footer_para.runs:
            style_run(run, zh="宋体", latin="Times New Roman", size=10.5)
    else:
        clear_paragraph(section.header.paragraphs[0])
        clear_paragraph(section.footer.paragraphs[0])


def add_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)


def set_heading_style(paragraph, level: int, text: str) -> None:
    paragraph.style = f"Heading {level}"
    clear_paragraph(paragraph)
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = paragraph.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(12)
        pf.first_line_indent = Pt(0)
        run = paragraph.add_run(text)
        style_run(run, zh="黑体", latin="Times New Roman", size=18, bold=True)
    elif level == 2:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = paragraph.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        pf.first_line_indent = Pt(0)
        run = paragraph.add_run(text)
        style_run(run, zh="黑体", latin="Times New Roman", size=16, bold=True)
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = paragraph.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(6)
        pf.space_after = Pt(3)
        pf.first_line_indent = Pt(0)
        run = paragraph.add_run(text)
        style_run(run, zh="黑体", latin="Times New Roman", size=12, bold=True)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(text)
    style_run(run, zh="宋体", latin="Times New Roman", size=12)


def add_appendix_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    run = p.add_run(text)
    style_run(run, zh="黑体", latin="Times New Roman", size=12, bold=True)


def add_appendix_en_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(text)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=10.5)


def add_appendix_cn_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(text)
    style_run(run, zh="宋体", latin="Times New Roman", size=12)


def add_reference_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    pf.hanging_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(text)
    style_run(run, zh="宋体", latin="Times New Roman", size=12)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    run = p.add_run(text)
    style_run(run, zh="宋体", latin="Times New Roman", size=10.5)


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    run = p.add_run(text)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=10.5)


def add_code_block(doc: Document, caption: str, lines: list[str]) -> None:
    add_caption(doc, caption)
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        run = p.add_run(line)
        style_run(run, zh="Times New Roman", latin="Times New Roman", size=10.5)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]]) -> None:
    add_caption(doc, caption)
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row_idx, row_data in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.first_line_indent = Pt(0)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                for run in p.runs:
                    style_run(run, zh="宋体", latin="Times New Roman", size=10.5, bold=(row_idx == 0))


def build_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("学士学位论文")
    style_run(run, zh="黑体", latin="Times New Roman", size=22, bold=True)

    for _ in range(5):
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(TITLE_CN)
    style_run(run, zh="黑体", latin="Times New Roman", size=22, bold=True)

    for _ in range(5):
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(0)

    for label, value in [
        ("设计人", STUDENT),
        ("学号", STUDENT_ID),
        ("指导教师", ADVISOR),
        ("所属系部", COLLEGE),
        ("专业班级", MAJOR_CLASS),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(f"{label}：{value}")
        style_run(run, zh="宋体", latin="Times New Roman", size=14)

    for _ in range(3):
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DATE_TEXT)
    style_run(run, zh="宋体", latin="Times New Roman", size=14)


def build_task_book(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(TASK_BOOK_TITLE)
    style_run(run, zh="黑体", latin="Times New Roman", size=16, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"学院：{COLLEGE}")
    style_run(run, zh="宋体", latin="Times New Roman", size=12)

    table = doc.add_table(rows=len(TASK_BOOK_META) + 1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row_data in enumerate(TASK_BOOK_META):
        for col_idx, value in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    title_row = len(TASK_BOOK_META)
    table.cell(title_row, 0).text = "设计（论文）题目"
    table.cell(title_row, 1).text = TASK_BOOK_TOPIC
    table.cell(title_row, 1).merge(table.cell(title_row, 3))

    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = para.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.first_line_indent = Pt(0)
                for run in para.runs:
                    style_run(run, zh="宋体", latin="Times New Roman", size=10.5)

    def add_task_label(text: str) -> None:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run(text)
        style_run(run, zh="黑体", latin="Times New Roman", size=12, bold=True)

    add_task_label("设计目的要求")
    add_body_paragraph(doc, TASK_BOOK_PURPOSE)

    add_task_label("设计主要内容")
    for idx, text in enumerate(TASK_BOOK_CONTENT, start=1):
        add_body_paragraph(doc, f"{idx}. {text}")

    add_task_label("设计提交资料")
    for idx, text in enumerate(TASK_BOOK_SUBMIT, start=1):
        add_body_paragraph(doc, f"{idx}. {text}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run("学生签名：                    指导教师签名：                    系主任签名：                    主管院长签名：")
    style_run(run, zh="宋体", latin="Times New Roman", size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run("说明：一式两份，一份装订入学生毕业设计（论文）内，一份交学院。")
    style_run(run, zh="宋体", latin="Times New Roman", size=10.5)


def add_cn_abstract(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("中文摘要")
    style_run(run, zh="黑体", latin="Times New Roman", size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(CN_ABSTRACT)
    style_run(run, zh="宋体", latin="Times New Roman", size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run1 = p.add_run("关键词：")
    run2 = p.add_run(KEYWORDS_CN)
    style_run(run1, zh="黑体", latin="Times New Roman", size=12, bold=True)
    style_run(run2, zh="宋体", latin="Times New Roman", size=12)


def add_en_abstract(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE_EN)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("author: Hu Zhuofan")
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("tutor: Li Junji")
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Abstract")
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(EN_ABSTRACT)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run1 = p.add_run("Keywords: ")
    run2 = p.add_run(KEYWORDS_EN)
    style_run(run1, zh="Times New Roman", latin="Times New Roman", size=12, bold=True)
    style_run(run2, zh="Times New Roman", latin="Times New Roman", size=12)


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目录")
    style_run(run, zh="黑体", latin="Times New Roman", size=16, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "打开 Word 或 WPS 后更新目录")
    for run in p.runs:
        style_run(run, zh="宋体", latin="Times New Roman", size=10.5)


def render_blocks(doc: Document, blocks: list[tuple[str, object]]) -> None:
    first_numbered_chapter = True
    for block_type, payload in blocks:
        if block_type == "chapter":
            if not first_numbered_chapter:
                doc.add_page_break()
            first_numbered_chapter = False
            p = doc.add_paragraph()
            set_heading_style(p, 1, str(payload))
        elif block_type == "chapter_no_number":
            doc.add_page_break()
            p = doc.add_paragraph()
            set_heading_style(p, 1, str(payload))
        elif block_type == "section":
            p = doc.add_paragraph()
            set_heading_style(p, 2, str(payload))
        elif block_type == "subsection":
            p = doc.add_paragraph()
            set_heading_style(p, 3, str(payload))
        elif block_type == "paragraph":
            add_body_paragraph(doc, str(payload))
        elif block_type == "appendix_subtitle":
            add_appendix_subtitle(doc, str(payload))
        elif block_type == "appendix_en":
            add_appendix_en_paragraph(doc, str(payload))
        elif block_type == "appendix_cn":
            add_appendix_cn_paragraph(doc, str(payload))
        elif block_type == "equation":
            add_equation(doc, str(payload))
        elif block_type == "table":
            add_table(doc, payload["caption"], payload["headers"], payload["rows"])
        elif block_type == "code":
            add_code_block(doc, payload["caption"], payload["lines"])
        else:
            raise ValueError(f"Unsupported block type: {block_type}")


def main() -> None:
    doc = Document()
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    doc.core_properties.title = TITLE_CN
    doc.core_properties.author = STUDENT
    doc.core_properties.subject = "本科毕业论文"

    build_cover(doc)
    configure_section(doc.sections[0], with_header=False)
    doc.add_page_break()
    build_task_book(doc)

    prelim_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(prelim_section, with_header=True, page_fmt="upperRoman", start=1)
    add_cn_abstract(doc)
    doc.add_page_break()
    add_en_abstract(doc)
    doc.add_page_break()
    add_toc(doc)

    main_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main_section, with_header=True, page_fmt="decimal", start=1)

    render_blocks(doc, BLOCKS)
    for ref in REFERENCES:
        add_reference_paragraph(doc, ref)
    render_blocks(doc, APPENDIX_BLOCKS)

    add_update_fields_on_open(doc)
    doc.save(str(OUT_DOCX))


if __name__ == "__main__":
    main()
