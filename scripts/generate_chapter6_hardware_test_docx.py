from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
OUT = ROOT / "第六章_系统测试与分析_硬件测试补充稿.docx"


def set_fonts(run, zh: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    set_fonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line: bool = True) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_line else Pt(0)


def add_title(doc: Document, text: str, *, size: float = 18) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=size, bold=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(8 if level == 1 else 5)
    pf.space_after = Pt(5)
    pf.first_line_indent = Pt(0)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=16 if level == 1 else 14 if level == 2 else 12, bold=True)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p)
    run = p.add_run(text)
    style_run(run, size=12)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    run = p.add_run(text)
    style_run(run, zh="楷体", size=11)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    run = p.add_run(f"• {text}")
    style_run(run, size=12)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    style_run(run, size=10.5)


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for idx, title in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = p.add_run(title)
        style_run(run, zh="黑体", size=10.5, bold=True)

    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            style_run(run, size=10.5)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)

    add_title(doc, "第六章 系统测试与分析（含硬件测试补充稿）")
    add_note(doc, "说明：本稿用于补充第六章中的硬件测试内容，并将 6.2、6.3、6.4 调整为可直接替换的论文正文。")

    add_heading(doc, "一、建议修改位置", level=1)
    add_bullet(doc, "在 6.1.2 测试覆盖范围之后新增 6.1.3 硬件测试内容与方法。")
    add_bullet(doc, "用新的表6-1替换原功能测试表，把硬件测试和平台测试统一纳入。")
    add_bullet(doc, "在 6.3 前部补入 6.3.1 终端硬件联调结果，原 6.3.1 到 6.3.5 顺延。")
    add_bullet(doc, "用新的 6.4 测试结果分析替换原总结，使硬件、平台和页面三部分一起落地。")

    add_heading(doc, "二、可直接替换正文", level=1)

    add_heading(doc, "第六章 系统测试与分析", level=2)
    add_heading(doc, "6.1 测试环境与方法", level=3)
    add_body(doc, "系统测试围绕终端采样、设备入库、前端展示、硬件联调和运行日志几个层面展开。测试方法以功能联调和结果核对为主，通过观察串口输出、调用平台接口、访问前端页面以及统计运行日志，判断系统是否达到设计要求。由于本课题同时包含 STM32 终端和 Web 平台，因此测试既要覆盖软件链路，也要覆盖终端硬件模块的实际运行情况。")

    add_heading(doc, "6.1.1 测试判据", level=3)
    add_body(doc, "本文采用的测试判据不只停留在页面是否能够打开或接口是否返回 200。对于本系统而言，更重要的是终端采样值能否进入统一数据结构，历史记录与快照数据是否保持一致，概览页面和地图页面是否引用相同来源的数据，缓存与广播逻辑是否能够减少重复刷新，同时终端硬件能否稳定完成采样、串口通信和无线上传。")

    add_heading(doc, "6.1.2 测试覆盖范围", level=3)
    add_body(doc, "基于上述判据，测试覆盖范围包括终端采样与串口输出、/sensor 网关入库、实时与历史查询接口、概览与地图页面、数据源设置页面、AI 辅助分析页面、运行日志统计以及终端硬件联调。通过覆盖这些关键路径，可以对系统链路完整性形成较充分的验证。")

    add_heading(doc, "6.1.3 硬件测试内容与方法", level=3)
    add_body(doc, "考虑到本课题包含 STM32 终端硬件设计，系统测试不仅需要验证平台接口和页面展示，还需要对终端硬件模块本身进行联调验证。硬件测试主要围绕最小系统供电、温度采集、模拟量采集、无线通信以及状态指示几个方面展开，测试方法以实物上电、串口日志观察、采样结果比对和连续运行记录为主。")
    add_body(doc, "在最小系统测试中，重点检查 STM32F103C8T6 是否能够正常上电、复位和下载程序，验证 5 V 转 3.3 V 供电链路是否稳定，确认主控在持续运行过程中不会出现异常复位。温度采集测试主要观察 DS18B20 是否能够被正确识别，并通过串口输出核对温度读数是否连续变化。模拟量采集测试则围绕浊度、pH 和 TDS 三路 ADC 输入展开，检查采样值、电压换算结果和最终业务指标是否能够稳定输出。无线通信测试主要验证 ESP-01 与 USART2 的连接是否正常，WiFi 是否能够建立连接，以及 HTTP 上报是否能够持续完成。状态指示测试则通过调试串口输出和 PC13 心跳灯翻转情况，观察主循环运行是否稳定。")
    add_body(doc, "通过上述硬件测试，可以对终端从上电启动、传感器采集到无线传输的基本运行状态形成较完整的验证，为后续平台入库和页面展示测试提供基础。")

    add_heading(doc, "6.2 功能测试", level=3)
    add_table(
        doc,
        ["测试项", "测试方法", "预期结果", "实际结果"],
        [
            ("最小系统上电与下载", "给终端供电并烧录固件", "主控能够正常启动并运行程序", "符合预期"),
            ("DS18B20 温度采集", "观察串口温度输出", "能够稳定输出水温数据", "符合预期"),
            ("模拟量采集", "观察串口输出 ADC 值和换算值", "能够周期输出 pH、浊度和 TDS", "符合预期"),
            ("ESP-01 联网", "观察串口联网日志", "能够完成 WiFi 连接并进入上传状态", "符合预期"),
            ("心跳指示灯", "观察 PC13 对应 LED 状态", "主循环运行时心跳灯周期翻转", "符合预期"),
            ("设备入库", "终端调用 /sensor 接口", "生成历史记录并更新快照", "符合预期"),
            ("多源同步", "触发 sync_realtime 接口", "完成统一转换并写入数据库", "符合预期"),
            ("实时概览", "访问 Dashboard 页面", "卡片和列表正常刷新", "符合预期"),
            ("历史分析", "选择断面加载历史曲线", "按时间顺序展示趋势图", "符合预期"),
            ("断面地图", "打开 WaterMap 页面并点击标记", "展示断面详情和历史趋势", "符合预期"),
            ("实时广播", "执行设备入库或同步任务", "前端接收到更新事件", "符合预期"),
            ("AI 辅助分析", "提交推荐问题或自定义问题", "返回分析结果并记录日志", "符合预期"),
        ],
    )
    add_caption(doc, "表6-1 主要功能测试结果")
    add_body(doc, "从表6-1可以看出，硬件终端和平台功能均通过了基本测试。终端侧已经能够完成上电启动、温度读取、模拟量采样、无线联网和状态指示；平台侧能够完成数据入库、快照更新、页面展示和广播联动。这说明系统不仅在软件层面具备完整流程，在硬件层面也已经具备连续运行的基本条件。")

    add_heading(doc, "6.3 联调结果与运行统计", level=3)
    add_heading(doc, "6.3.1 终端硬件联调结果", level=3)
    add_body(doc, "终端硬件联调主要验证 STM32 最小系统、DS18B20、浊度/pH/TDS 采样接口、ESP-01 通信链路以及心跳指示模块的协同工作情况。在实物上电后，主控能够正常启动并输出调试信息，DS18B20 可以被识别并返回温度数据，ADC1 与 DMA 能够持续更新三路模拟量采样结果，ESP-01 在完成串口握手后能够建立 WiFi 连接并执行 HTTP 上报，PC13 对应心跳灯能够按设定周期翻转。联调结果表明，终端硬件各模块之间连接关系正确，能够支撑后续平台入库和页面展示功能。")

    add_heading(doc, "6.3.2 采集到查询的链路联调", level=3)
    add_body(doc, "在终端上电并完成网络连接后，串口能够持续输出采样结果，平台在收到 HTTP 请求后写入历史表和快照表，前端页面刷新后能够显示新增站点的最新状态。这一过程验证了从物理采样到页面查询的全链路连通性。")

    add_heading(doc, "6.3.3 运行日志统计结果", level=3)
    add_table(
        doc,
        ["统计指标", "结果", "说明"],
        [
            ("同步任务总次数", "465 次", "统计区间为 2026-03-05 至 2026-03-11"),
            ("产生新增写入的任务数", "257 次", "created 大于 0 的任务数"),
            ("无新增写入的任务数", "208 次", "说明重复同步不会持续膨胀历史表"),
            ("单次新增记录上限", "1000 条", "与同步接口默认 count 上限一致"),
            ("有新增任务平均写入量", "349.24 条", "仅统计 created 大于 0 的任务"),
            ("有新增任务中位数", "174 条", "反映常规同步任务规模"),
            ("代表性地理编码批次", "297/300、296/300、296/300", "平均成功率约 98.78%"),
        ],
    )
    add_caption(doc, "表6-2 运行日志统计结果")
    add_body(doc, "从表6-2可以看出，平台能够在多轮同步过程中保持较稳定的写入节奏。created=0 的任务数量较多，说明平台在识别重复记录方面起到了实际作用，历史表不会因为重复同步而无序膨胀。")

    add_heading(doc, "6.3.4 地图定位与断面展示结果", level=3)
    add_body(doc, "三次代表性地理编码批处理结果分别达到 297/300、296/300 和 296/300，平均成功率为 98.78%。这表明“坐标缓存 + 批量补齐”的方案能够满足当前地图展示需求，绝大多数断面都可以在页面中获得稳定的空间位置。")

    add_heading(doc, "6.3.5 广播与缓存控制结果", level=3)
    add_body(doc, "realtime 接口的数据版本机制与 sensorStore 缓存逻辑配合后，能够在无变化情况下直接返回 changed=false。页面因此不需要在每次轮询时重新渲染整批数据，既降低了前端刷新负担，也保持了监测页面的稳定性。")

    add_heading(doc, "6.3.6 页面级联与配置验证", level=3)
    add_body(doc, "在页面级联验证中，Settings 页面切换 auto/manual 模式后，概览与实时页面会依据新的数据源策略重新获取结果，页面顶部可用性标记也会同步更新。这一结果说明系统配置并非孤立页面，而是能够影响后续数据展示行为。")
    add_table(
        doc,
        ["验证项", "验证方式", "结果说明"],
        [
            ("模式切换生效", "在 Settings 页面保存 auto/manual 模式", "后台配置更新后页面可重新取到对应数据"),
            ("AI 摘要联动", "刷新 AI 页面并查看概览摘要", "监测断面数与告警数能与 overview 结果保持一致"),
            ("地图详情联动", "在 WaterMap 页面点击标记后查看趋势", "详情弹窗与历史查询接口返回一致"),
            ("缓存复用", "在短时间内重复请求实时数据", "无变化时直接复用已有缓存结果"),
        ],
    )
    add_caption(doc, "表6-3 页面级联验证结果")

    add_heading(doc, "6.4 测试结果分析", level=3)
    add_body(doc, "综合硬件测试、功能测试和运行日志统计结果可以认为，系统已经达到论文设计目标。终端侧能够完成稳定上电、温度采集、模拟量采样、无线传输和状态指示，说明硬件部分具备持续运行基础；平台侧能够完成统一入库、历史记录保存、快照维护和页面展示，说明系统软件链路较为完整。终端与平台联调后，从现场采样到页面查询之间不存在明显断点，系统能够形成较完整的数据闭环。")
    add_body(doc, "从本科毕业设计要求来看，当前系统不仅实现了 STM32 终端采集和 Web 平台展示，也完成了终端硬件、服务端处理和前端页面之间的协同验证。测试结果表明，本课题已经具备较清晰的系统结构、较明确的数据流向以及可核对的运行结果，能够体现水质监测系统在硬件采集、数据处理和页面展示三个层面的综合设计与实现过程。")

    doc.save(str(OUT))


if __name__ == "__main__":
    build_docx()
