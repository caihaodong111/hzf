from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
OUT = ROOT / "第二章_需求分析_重写稿.docx"

FUNCTION_TABLE_ROWS = [
    ("终端采集与上报模块", "传感采样、指标换算、HTTP 上报、串口调试", "负责产生本地监测数据并将结果送入平台"),
    ("平台接入与状态管理模块", "设备入库、多源同步、字段转换、快照维护", "负责把不同来源的数据整理为统一业务状态"),
    ("实时监测与总览模块", "概览卡片、实时列表、状态统计、广播刷新", "负责展示当前断面的最新状态和总体分布"),
    ("历史分析与回溯模块", "历史曲线、时间窗口查询、风险站点排序", "负责呈现连续变化过程并辅助判断趋势"),
    ("地图展示与区域筛选模块", "地理编码、地图标点、省市流域筛选、详情弹窗", "负责呈现监测数据的空间分布与区域特征"),
    ("辅助分析与配置模块", "AI 问答、数据源模式切换、手动同步触发、运行提示", "负责提升系统解释能力并支持运行策略调整"),
]

ROLE_TABLE_ROWS = [
    ("平台管理人员", "保障平台运行、观察全局状态、调整数据源策略", "查看概览、触发同步、切换数据源模式、使用 AI 助手分析重点断面"),
    ("监测业务人员", "掌握断面当前状态、历史趋势和空间分布", "查看实时列表、分析历史曲线、浏览地图分布、定位高风险断面"),
    ("演示观察人员", "理解终端到平台的完整监测链路", "观察 STM32 上报、查看页面联动结果、体验 AI 分析与地图展示"),
]

NON_FUNCTION_TABLE_ROWS = [
    ("实时性", "在合理时间内看到最新监测结果和状态变化", "快照表、REST 接口、WebSocket 广播"),
    ("一致性", "不同来源数据以统一字段和统一语义进入平台", "DataTransformer、统一模型字段、双表结构"),
    ("稳定性", "网络波动、重复同步和页面轮询不导致系统失稳", "WiFi 重试、去重写入、data_version 机制"),
    ("可扩展性", "便于接入新数据源、新指标和新页面", "app 分层、数据源模式、统一接口"),
    ("可维护性", "具备调试、日志和结构化修改基础", "调试串口、运行日志、AI 问答日志、部署说明"),
]


def set_fonts(run, zh: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    set_fonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line: bool = True) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_line else Pt(0)


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=16, bold=True)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(8 if level == 1 else 5)
    pf.space_after = Pt(4)
    pf.first_line_indent = Pt(0)
    run = p.add_run(text)
    size_map = {1: 14, 2: 12, 3: 12}
    style_run(run, zh="黑体", size=size_map.get(level, 12), bold=True)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p)
    run = p.add_run(text)
    style_run(run, size=12)


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]], caption: str) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = p.add_run(header)
        style_run(run, zh="黑体", size=10.5, bold=True)

    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            style_run(run, size=10.5)

    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(6)
    run = cp.add_run(caption)
    style_run(run, size=10.5)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)

    add_title(doc, "第二章 需求分析（重写稿）")

    add_heading(doc, "2.1 可行性分析", 1)

    add_heading(doc, "2.1.1 技术可行性", 2)
    add_body(
        doc,
        "从终端侧来看，当前原型已经具备较清楚的硬件实现路径。终端以 STM32F103C8T6 作为控制核心，"
        "接入 DS18B20 温度传感器以及 pH、浊度、TDS 三类模拟传感器，能够完成温度读取、ADC 多通道采样、"
        "指标换算和本地状态控制；ESP-01 模块通过 USART2 与主控连接，用于完成 WiFi 接入与 HTTP 数据上报，"
        "调试串口和心跳灯则用于运行观测与联调定位。STM32F1 系列资料丰富，标准外设库成熟，能够满足本课题对采样、"
        "通信和周期调度等方面的实现需求，因此终端侧技术路线较为明确。"
    )
    add_body(
        doc,
        "从平台侧来看，系统采用 Django、Vue 3、MySQL、Redis/Channels、Celery、ECharts 和高德地图等较成熟的技术方案。"
        "后端负责终端入库、多源同步、历史保存、快照维护和接口供给，前端负责概览监测、历史分析、地图展示、"
        "数据源设置以及 AI 辅助分析，Redis/Channels 用于实时广播，Celery 用于周期同步任务。与此同时，"
        "系统还支持 STM32 终端上报与国家平台、华为云等外部数据源统一转换，这说明平台方案不仅能够运行，"
        "也保留了后续扩展数据来源和监测指标的空间。"
    )
    add_body(
        doc,
        "从整体链路来看，终端通过 /sensor 网关上报或平台通过同步任务获取的数据进入系统后，已经能够完成统一建模、"
        "历史表写入、快照更新、实时广播、页面展示和辅助分析，形成“采集—入库—展示—分析”的完整闭环。"
        "这说明本课题并不是停留在概念描述层面，而是建立在已有工程实现基础之上的，因此从技术实现角度看，"
        "本系统具备较好的可行性。"
    )

    add_heading(doc, "2.1.2 经济可行性", 2)
    add_body(
        doc,
        "本课题所采用的硬件以 STM32F103C8T6 最小系统板、DS18B20 温度传感器、pH 传感器、浊度传感器、"
        "TDS 传感器以及 ESP-01 WiFi 模块为主，这些器件均属于常见开发型模块，采购渠道较广、单价较低，"
        "不依赖价格较高的工业级在线监测设备或专用采集网关。对于本科毕业设计而言，该配置已经能够完成多参数采集、"
        "无线联网和端到端上报等核心功能验证，整体硬件投入较易控制。"
    )
    add_body(
        doc,
        "在系统联调和答辩演示阶段，平台除真实终端上报外，还可以利用现有的多源同步能力、数据库快照和历史数据接口完成"
        "多断面监测场景的展示，这意味着并不需要部署大量硬件节点就能完成系统验证。这样一来，实验成本、设备维护成本"
        "以及联调成本都能够控制在较低水平。"
    )
    add_body(
        doc,
        "软件部分主要依赖 Django、Vue 3、MySQL、Redis、Celery、ECharts 等开源框架和常见服务，不涉及额外商业授权费用。"
        "服务器既可以在本地实验环境中运行，也可以部署到轻量级云主机进行演示，整体运行成本较低。因此，从经费投入和资源条件来看，"
        "本课题具有较好的经济可行性。"
    )

    add_heading(doc, "2.1.3 运行可行性", 2)
    add_body(
        doc,
        "系统当前已经具备较明确的运行条件。项目提供了较完整的前端工程、后端工程、固件代码、部署说明和容器化编排文件，"
        "能够支持本地联调和分模块部署；开发阶段可分别启动 Django 服务、Vue 前端、Redis、Celery Worker/Beat 以及 STM32 终端，"
        "从而较快完成接口验证、页面调试和任务联调，环境复现难度相对较低。"
    )
    add_body(
        doc,
        "另外，系统运行方式较为灵活。在终端侧，可以通过 STM32 实物节点完成采样、串口观察和 WiFi 上报演示；在平台侧，"
        "可以通过手动上报接口、定时同步任务和数据库快照机制补足历史分析、地图展示和多断面监测场景。对于毕业设计阶段的开发、"
        "测试和答辩而言，这种“真实终端 + 平台多源数据”的运行方式既能体现硬件参与，也能保证系统展示的完整性和稳定性。"
    )
    add_body(
        doc,
        "因此，从系统部署、联调和演示条件来看，本课题具备较好的运行可行性。"
    )

    add_heading(doc, "2.2 功能需求分析", 1)

    add_heading(doc, "2.2.1 场景描述", 2)
    add_body(
        doc,
        "本系统主要面向养殖水体巡检、河湖断面监测以及教学实验演示三类场景。在现场监测场景中，STM32 终端负责周期采集水温、"
        "pH、浊度和 TDS 等指标，并通过 WiFi 将结果上传到平台；在区域断面监测场景中，平台还需要接入国家平台和华为云等外部数据源，"
        "将不同来源的监测数据整理为统一状态后供页面展示与分析使用；在教学和答辩演示场景中，系统还应能够清楚展示“终端采样、平台入库、"
        "页面联动、文字分析”这一完整链路。"
    )
    add_body(
        doc,
        "在这一场景下，系统不能只完成单点采样，也不能只停留在静态页面展示，而需要形成一条完整业务链。终端采集到的数据需要能够进入平台，"
        "平台处理后的结果既要能够以实时列表、历史曲线和地图分布的方式呈现，也要能够转化为风险提示和辅助分析结论；只有这样，系统才真正具备"
        "水质监测平台应有的实际价值。"
    )

    add_heading(doc, "2.2.2 数据采集与终端上报功能需求", 2)
    add_body(
        doc,
        "水质变化具有连续性，如果只依赖偶发读数或人工记录，就难以及时发现异常波动。因此，系统首先需要具备稳定的数据采集与终端上报能力，"
        "而且这部分不能只停留在“能读到值”这一层，还应把指标换算、上报组织和调试观测一起纳入考虑。"
    )
    add_body(
        doc,
        "传感采样子功能。终端应能够周期读取 DS18B20 的温度数据，并通过 ADC 采集 pH、浊度和 TDS 三路模拟量，为后续换算和展示提供基础数据。"
        "数据换算与状态控制子功能。终端应能够把原始采样值转换为可展示的业务指标，并通过心跳灯和主循环调度保持运行状态可观测。"
        "报文封装与上报子功能。终端应将站点编号、站点名称、时间戳及各项指标封装为统一结构，通过 HTTP 接口送入平台，保证后端能够按统一规则接收和处理。"
        "调试观测子功能。终端还应能够通过 USART1 输出关键日志，用于联调阶段观察采样结果、联网状态和上报结果。"
    )
    add_body(
        doc,
        "由此可见，终端功能需求不是单纯的读数功能，而是由采样、换算、上报和调试共同构成的一组前端入口能力。"
    )

    add_heading(doc, "2.2.3 平台接入与状态管理功能需求", 2)
    add_body(
        doc,
        "监测数据进入平台后，如果不能及时完成解析、统一转换和状态更新，后续的实时监控、历史分析和地图展示就都无法成立。因此，"
        "系统还需要具备平台接入与状态管理功能，把原始上报数据和外部同步数据转化为可查询、可维护的业务状态。"
    )
    add_body(
        doc,
        "设备入库与报文解析子功能。平台应能够接收 STM32 终端通过 /sensor 或兼容入口上报的数据，提取站点编号、时间戳和监测指标等字段。"
        "多源同步与统一转换子功能。平台应能够从国家平台、华为云等来源获取实时数据，并将字段名称、时间格式、区域字段和水质类别表达转换为统一结构。"
        "历史与快照双表存储子功能。系统应将连续历史记录保存到历史表中，并为每个站点维护一条最新快照，以同时满足“查变化过程”和“查当前状态”两类需求。"
        "接口供给与广播更新子功能。平台还应向上层页面提供实时数据、历史曲线、总览统计和区域筛选接口，并在快照变化后触发广播，保证页面能够较快获取更新结果。"
    )
    add_body(
        doc,
        "这一部分功能承担的是“数据进入系统后的第一次业务落地”，它决定了平台能否把离散数据整理成稳定的断面状态和区域状态。"
    )

    add_heading(doc, "2.2.4 实时监测与总览功能需求", 2)
    add_body(
        doc,
        "水质监测系统不仅要保存数据，还要让管理人员能够及时看到当前状态和整体分布。因此，平台在完成接入之后，还需要提供面向日常值守的实时监测与总览功能。"
    )
    add_body(
        doc,
        "看板展示子功能。系统应提供概览页面，用于展示断面总数、平均指标、异常告警数量和最新更新时间，使管理人员能够快速掌握整体运行情况。"
        "实时列表子功能。系统应支持按省市、关键字等条件筛选当前快照数据，并在列表中展示站点名称、水质类别、监测时间和核心指标。"
        "状态统计子功能。平台应能够基于快照数据计算总体均值、风险数量和类别分布，避免不同页面各自重复统计。"
        "刷新联动子功能。系统还应在数据变化时通过轮询或 WebSocket 方式更新页面，使实时监测页面能够长期运行而不失去时效性。"
    )
    add_body(
        doc,
        "也就是说，实时监测功能不仅要做到“看最新值”，还要能够“看总体状态”和“看更新节奏”，这样页面才真正具备值守意义。"
    )

    add_heading(doc, "2.2.5 历史分析与地图展示功能需求", 2)
    add_body(
        doc,
        "除了当前快照，用户还需要查看某一断面在一定时间窗口内的连续变化情况，以及监测点在空间上的分布特征。因此，系统还需要具备历史分析与地图展示功能。"
    )
    add_body(
        doc,
        "历史曲线子功能。系统应支持按站点查询最近 24 小时、7 天或 30 天内的历史数据，并以折线图方式展示指标变化过程。"
        "风险排序子功能。系统应能够依据快照中的关键指标计算风险分值，对高风险断面进行排序，便于用户优先关注重点站点。"
        "地图展示子功能。对于已具备经纬度信息的站点，系统应能够在地图中进行标点展示，并结合弹窗输出站点详情和短期趋势。"
        "区域筛选子功能。系统应支持按省份、城市和流域进行筛选，使地图和列表能够体现区域监测特征。"
    )
    add_body(
        doc,
        "可以看出，历史与地图功能不仅服务于结果回看，也服务于风险定位和区域对比，是监测平台中不可缺少的分析入口。"
    )

    add_heading(doc, "2.2.6 风险识别与辅助分析功能需求", 2)
    add_body(
        doc,
        "在监测场景中，平台不能只返回原始数值，还需要对重点信号给出更容易理解的提示。因此，系统还应具备风险识别与辅助分析功能。"
    )
    add_body(
        doc,
        "风险识别子功能。系统应能够根据溶解氧、pH、水质类别及其他关键指标计算风险分值，识别当前需重点关注的断面，并在概览与分析页面中给出风险数量和排序结果。"
        "辅助问答子功能。系统应能够基于当前快照构造上下文，将监测结果转换为文字化结论、依据和建议，帮助用户更快理解数据含义。"
        "问答留痕子功能。为便于审计、回溯和优化，系统还应记录 AI 提问、回答和调用结果，使辅助分析过程具备可追踪性。"
    )
    add_body(
        doc,
        "因此，辅助分析功能本质上不是替代人工判断，而是把当前快照中的风险信号转化为更容易理解和复盘的解释结果。"
    )

    add_heading(doc, "2.2.7 数据源配置与系统联调功能需求", 2)
    add_body(
        doc,
        "由于系统同时支持多源数据接入和本地终端演示，因此平台还需要具备面向运行策略调整的数据源配置与联调功能。"
    )
    add_body(
        doc,
        "数据源模式切换子功能。系统应支持自动模式和手动模式，用于控制国家平台、华为云及本地数据在不同运行场景下的优先级。"
        "可用性提示子功能。页面应能够显示各数据源当前是否可用，减少联调时的信息不对称。"
        "手动同步与刷新子功能。系统应支持人工触发同步或刷新操作，使答辩演示和问题定位过程更可控。"
        "联调观察子功能。平台还应与终端串口日志、页面刷新结果和快照数据保持对应关系，便于验证端到端链路是否正常工作。"
    )
    add_body(
        doc,
        "这部分需求体现的是系统在原型阶段的工程实用性，即不仅要能展示结果，也要能支持联调、切换和问题定位。"
    )

    add_heading(doc, "2.2.8 功能需求总结", 2)
    add_body(
        doc,
        "结合上述分析，系统功能需求可以归纳为终端采集与上报、平台接入与状态管理、实时监测与总览、历史分析与回溯、地图展示与区域筛选、辅助分析与配置六类模块。"
        "各模块及其主要子功能如表 2-1 所示。"
    )
    add_table(
        doc,
        ["功能模块", "主要子功能", "作用说明"],
        FUNCTION_TABLE_ROWS,
        "表2-1 系统功能需求与功能模块对应表",
    )

    add_heading(doc, "2.3 角色分析", 1)
    add_body(
        doc,
        "角色分析用于说明系统中的主要使用者是谁、他们更关注什么问题，以及他们与系统之间会发生哪些典型交互。"
        "尽管当前原型并未实现复杂的权限体系，但从业务使用角度看，仍然可以将系统使用者划分为平台管理人员、监测业务人员和演示观察人员三类。"
    )

    add_heading(doc, "2.3.1 角色划分", 2)
    add_body(
        doc,
        "平台管理人员主要负责系统日常运行和联调控制，关注数据源可用性、快照更新状态、风险数量以及 AI 分析结果，需要使用概览、分析、地图、"
        "数据源设置等页面完成运行观察和模式切换。监测业务人员更关注断面本身的变化过程，主要查看实时列表、历史曲线、风险排序和地图位置分布，"
        "用于判断哪些区域和断面需要优先关注。演示观察人员则更关注系统链路是否完整，通常通过观察 STM32 终端、页面联动和 AI 输出结果来理解系统整体工作过程。"
    )
    add_table(
        doc,
        ["角色", "关注目标", "主要操作"],
        ROLE_TABLE_ROWS,
        "表2-2 角色与主要职责对应表",
    )

    add_heading(doc, "2.3.2 角色用例分析", 2)
    add_body(
        doc,
        "依据上述角色划分，可以将平台管理人员的主要用例概括为查看概览、触发同步、调整数据源模式、查看地图和使用 AI 助手；"
        "监测业务人员的主要用例集中在实时查看、历史分析、风险定位和区域筛选；演示观察人员的主要交互则是观察终端上报与页面展示之间的联动关系。"
    )
    add_body(
        doc,
        "从用例关系上看，三类角色都围绕“看状态”这一核心目标展开，但平台管理人员更偏向运行控制，监测业务人员更偏向分析判断，"
        "演示观察人员更偏向理解链路和展示效果。这种划分方式与项目当前的页面结构和使用方式基本一致，也符合毕业设计阶段的实际应用特点。"
    )

    add_heading(doc, "2.4 非功能需求分析", 1)
    add_body(
        doc,
        "除功能需求外，系统还应满足实时性、一致性、稳定性、可扩展性和可维护性等方面的要求。这些内容虽然不直接表现为某个按钮或某个页面，"
        "但会直接影响系统架构设计、接口组织方式以及后续联调效率，因此需要单独说明。"
    )
    add_table(
        doc,
        ["非功能属性", "需求说明", "当前项目中的设计支撑"],
        NON_FUNCTION_TABLE_ROWS,
        "表2-3 非功能需求与设计支撑关系",
    )

    add_heading(doc, "2.4.1 实时性需求", 2)
    add_body(
        doc,
        "水质监测平台对实时性的要求主要体现在合理时间内能够看到最新状态变化，而不是追求毫秒级控制。终端侧应能够按秒级周期完成采样与上报，"
        "平台在接收到数据后应较快完成历史入库、快照更新和页面刷新，使用户能够在当前页面中及时看到最新监测结果。"
    )
    add_body(
        doc,
        "实时性需求还体现在更新链路的完整性上。数据变化后不仅要写入数据库，还要通过广播或刷新机制让前端及时感知，因此快照维护和 WebSocket 更新通知"
        "都是实时性需求的重要组成部分。"
    )

    add_heading(doc, "2.4.2 一致性需求", 2)
    add_body(
        doc,
        "系统同时接入 STM32 终端数据、国家平台数据和华为云数据，如果不同来源在字段命名、时间格式和水质类别表达上不一致，"
        "前端页面和统计逻辑就无法复用。因此，多源数据应在进入平台后转换为统一模型，并以相同语义写入历史表和快照表。"
    )
    add_body(
        doc,
        "一致性需求还意味着概览页面、历史分析页面、地图页面和 AI 页面应尽量共享同一份当前状态来源，避免因为页面使用不同数据口径而出现结果分裂。"
    )

    add_heading(doc, "2.4.3 稳定性需求", 2)
    add_body(
        doc,
        "系统在持续运行中需要避免单点失稳。终端侧应能够在 WiFi 波动时进行自动重试，并保持主循环持续运行；平台侧应能够在重复同步时避免无意义的重复写入，"
        "在前端频繁轮询时避免整批无变化数据被重复下发。"
    )
    add_body(
        doc,
        "对于监测平台而言，稳定性还体现在页面联动上。地图、概览和分析页面应共享稳定的快照数据来源，确保不同页面在同一时刻反映的是相互一致的当前状态。"
    )

    add_heading(doc, "2.4.4 可扩展性需求", 2)
    add_body(
        doc,
        "水质监测平台后续可能继续接入新的数据源、扩展新的水质指标或增加新的可视化页面，因此系统结构应保留扩展空间。"
        "当前平台通过 app 分层、统一数据模型和数据源模式配置，为新增来源和新增页面保留了接口层和存储层的适配基础。"
    )
    add_body(
        doc,
        "后续扩展时，应尽量通过调整接入层、转换层和页面组件完成，而不破坏现有的历史表、快照表和统一接口结构。"
    )

    add_heading(doc, "2.4.5 可维护性需求", 2)
    add_body(
        doc,
        "毕业设计项目不仅要能够运行，还要便于说明、调试和继续修改。因此，系统在组织结构上应保持清晰的前后端边界、明确的数据表职责以及可读的接口语义，"
        "环境复现过程也应尽量简化。"
    )
    add_body(
        doc,
        "当前项目已经具备较清楚的终端固件、后端服务、前端页面、运行日志和部署文件划分；终端保留调试串口和心跳灯，平台保留同步日志、AI 问答日志和可切换的数据源模式。"
        "这些设计都有利于后续定位问题、补充材料和继续扩展。"
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_docx()
