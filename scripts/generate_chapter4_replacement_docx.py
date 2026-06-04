from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
SRC = ROOT / "第四章_系统详细设计_直接替换稿.md"
OUT = ROOT / "第四章_系统详细设计_直接替换稿.docx"


def set_fonts(run, zh: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    set_fonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line: bool = True, spacing=WD_LINE_SPACING.ONE_POINT_FIVE) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_line else Pt(0)


def clean_inline(text: str) -> str:
    return text.replace("`", "").strip()


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(clean_inline(text))
    style_run(run, zh="黑体", size=18, bold=True)


def add_heading(doc: Document, text: str, level: int) -> None:
    sizes = {1: 16, 2: 15, 3: 14, 4: 13}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(8 if level == 1 else 5)
    pf.space_after = Pt(5)
    pf.first_line_indent = Pt(0)
    run = p.add_run(clean_inline(text))
    style_run(run, zh="黑体", size=sizes.get(level, 12), bold=True)


def add_body(doc: Document, text: str) -> None:
    text = clean_inline(text)
    if not text:
        return
    p = doc.add_paragraph()
    is_formula = re.match(r"^[A-Za-z].*（\d+-\d+）$", text) is not None
    if is_formula:
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, spacing=WD_LINE_SPACING.SINGLE)
    else:
        format_paragraph(p)
    run = p.add_run(text)
    style_run(run, size=12)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    run = p.add_run(f"• {clean_inline(text)}")
    style_run(run, size=12)


def split_row(line: str) -> list[str]:
    parts = [clean_inline(cell) for cell in line.strip().strip("|").split("|")]
    return [part for part in parts]


def is_separator_row(line: str) -> bool:
    stripped = line.strip()
    return bool(stripped) and set(stripped) <= {"|", "-", " ", ":"}


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for col, value in enumerate(rows[0]):
        cell = table.rows[0].cells[col]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = p.add_run(value)
        style_run(run, zh="黑体", size=10.5, bold=True)

    for row_values in rows[1:]:
        row = table.add_row()
        for col in range(cols):
            value = row_values[col] if col < len(row_values) else ""
            cell = row.cells[col]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            style_run(run, size=10.5)


def parse_blocks(lines: list[str]) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue
        if stripped.startswith("|"):
            raw_rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                current = lines[i].rstrip()
                if not is_separator_row(current):
                    raw_rows.append(split_row(current))
                i += 1
            if raw_rows:
                blocks.append(("table", raw_rows))
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            blocks.append(("heading", (level, text)))
            i += 1
            continue
        if stripped.startswith("- "):
            blocks.append(("bullet", stripped[2:].strip()))
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt == "---" or nxt.startswith("#") or nxt.startswith("- ") or nxt.startswith("|"):
                break
            paragraph_lines.append(nxt)
            i += 1
        blocks.append(("paragraph", " ".join(paragraph_lines)))
    return blocks


def build_docx() -> None:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    blocks = parse_blocks(lines)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)

    for kind, payload in blocks:
        if kind == "heading":
            level, text = payload
            if level == 1:
                add_title(doc, text)
            elif level == 2:
                add_heading(doc, text, 1)
            elif level == 3:
                add_heading(doc, text, 2)
            elif level == 4:
                add_heading(doc, text, 3)
            else:
                add_heading(doc, text, 4)
        elif kind == "paragraph":
            add_body(doc, payload)
        elif kind == "bullet":
            add_bullet(doc, payload)
        elif kind == "table":
            add_table(doc, payload)

    doc.save(str(OUT))


if __name__ == "__main__":
    build_docx()
