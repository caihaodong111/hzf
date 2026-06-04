from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
ASSETS_DIR = ROOT / "thesis_assets"
OUT_DOCX = ROOT / "基于STM32的水质监测系统及断面数据展示平台_修改补充稿.docx"
TECH_ROUTE_IMG = ASSETS_DIR / "fig-revision-system-route.png"
HARDWARE_IMG = ASSETS_DIR / "fig-revision-hardware-block.png"

DATE_TEXT = "2026年5月27日"


def pick_font(bold: bool = False) -> str:
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc" if bold else "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for item in candidates:
        path = Path(item)
        if path.exists():
            return str(path)
    return ""


FONT_REGULAR = pick_font(False)
FONT_BOLD = pick_font(True)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = FONT_BOLD if bold else FONT_REGULAR
    if path:
        return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def draw_centered_multiline(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    font: ImageFont.ImageFont,
    fill: str,
    spacing: int = 8,
) -> None:
    lines = text.split("\n")
    heights: list[int] = []
    widths: list[int] = []
    for line in lines:
        left, top, right, bottom = draw.textbbox((0, 0), line, font=font)
        widths.append(right - left)
        heights.append(bottom - top)
    total_height = sum(heights) + spacing * max(0, len(lines) - 1)
    x1, y1, x2, y2 = box
    current_y = y1 + (y2 - y1 - total_height) / 2
    for line, width, height in zip(lines, widths, heights):
        current_x = x1 + (x2 - x1 - width) / 2
        draw.text((current_x, current_y), line, font=font, fill=fill)
        current_y += height + spacing


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    *,
    fill: str,
    outline: str,
    radius: int = 24,
    width: int = 4,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow_line(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    fill: str = "#0f172a",
    width: int = 5,
    head_len: int = 18,
    head_half: int = 8,
) -> None:
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=fill, width=width)
    if x2 == x1 and y2 == y1:
        return
    dx = x2 - x1
    dy = y2 - y1
    if abs(dx) >= abs(dy):
        if dx >= 0:
            head = [(x2, y2), (x2 - head_len, y2 - head_half), (x2 - head_len, y2 + head_half)]
        else:
            head = [(x2, y2), (x2 + head_len, y2 - head_half), (x2 + head_len, y2 + head_half)]
    else:
        if dy >= 0:
            head = [(x2, y2), (x2 - head_half, y2 - head_len), (x2 + head_half, y2 - head_len)]
        else:
            head = [(x2, y2), (x2 - head_half, y2 + head_len), (x2 + head_half, y2 + head_len)]
    draw.polygon(head, fill=fill)


def build_tech_route_image(path: Path) -> None:
    image = Image.new("RGB", (1600, 680), "#f8fbff")
    draw = ImageDraw.Draw(image)

    title_font = load_font(40, bold=True)
    box_title_font = load_font(26, bold=True)
    box_body_font = load_font(22, bold=False)
    note_font = load_font(20, bold=False)

    draw.text((520, 40), "本系统技术路线图", fill="#0f172a", font=title_font)
    draw.text((360, 96), "从现场采集终端到平台展示与验证，按完整业务闭环组织系统设计", fill="#475569", font=note_font)

    boxes = [
        ("研究需求", "面向水体巡检、\n断面监测与教学演示"),
        ("终端采集", "温度、pH、浊度、\nTDS 多参数采样"),
        ("终端处理", "STM32 完成调度、\n换算与报文组织"),
        ("无线传输", "ESP-01 接入 WiFi，\n周期上传监测数据"),
        ("平台接入", "服务器完成接收、\n字段统一与数据入库"),
        ("存储治理", "历史表、快照表与\n坐标缓存协同工作"),
        ("展示验证", "概览、分析、地图、\n辅助研判与联调测试"),
    ]

    start_x = 40
    top_y = 200
    box_w = 205
    box_h = 220
    gap = 18
    fills = ["#dbeafe", "#d1fae5", "#e9d5ff", "#fee2e2", "#fde68a", "#cffafe", "#e2e8f0"]
    outlines = ["#2563eb", "#059669", "#7c3aed", "#dc2626", "#d97706", "#0891b2", "#475569"]

    for idx, (title, body) in enumerate(boxes):
        x1 = start_x + idx * (box_w + gap)
        box = (x1, top_y, x1 + box_w, top_y + box_h)
        rounded_box(draw, box, fill=fills[idx], outline=outlines[idx], radius=28, width=4)
        draw_centered_multiline(draw, (x1 + 10, top_y + 24, x1 + box_w - 10, top_y + 92), title, font=box_title_font, fill="#0f172a")
        draw_centered_multiline(draw, (x1 + 12, top_y + 96, x1 + box_w - 12, top_y + box_h - 18), body, font=box_body_font, fill="#334155")
        if idx < len(boxes) - 1:
            arrow_y = top_y + box_h // 2
            arrow_line(draw, (x1 + box_w, arrow_y), (x1 + box_w + gap - 4, arrow_y), fill="#64748b", width=5)

    footer = "建议插入论文 1.3 节后，图号并入正文时按最终章节位置重新编号。"
    draw.text((410, 610), footer, fill="#64748b", font=note_font)
    image.save(path)


def build_hardware_image(path: Path) -> None:
    image = Image.new("RGB", (1800, 1100), "#fbfdff")
    draw = ImageDraw.Draw(image)

    title_font = load_font(44, bold=True)
    label_font = load_font(28, bold=True)
    body_font = load_font(24, bold=False)
    pin_font = load_font(20, bold=False)

    draw.text((560, 36), "终端硬件结构框图", fill="#0f172a", font=title_font)
    draw.text((455, 96), "基于 STM32F103C8T6 的多参数采样终端，按“采集 + 处理 + 通信 + 调试”组织硬件结构", fill="#475569", font=pin_font)

    power_box = (690, 150, 1110, 240)
    rounded_box(draw, power_box, fill="#fef3c7", outline="#d97706", radius=24, width=4)
    draw_centered_multiline(draw, power_box, "3.3V 供电与稳压\n为主控、传感器与 ESP-01 提供工作电压", font=body_font, fill="#7c2d12")

    mcu_box = (620, 350, 1180, 700)
    rounded_box(draw, mcu_box, fill="#dbeafe", outline="#2563eb", radius=32, width=5)
    draw_centered_multiline(
        draw,
        mcu_box,
        "STM32F103C8T6\n主控单元\n采样调度 / 数据换算 / 报文组织 / 通信控制",
        font=label_font,
        fill="#0f172a",
        spacing=10,
    )

    sensors = [
        ((110, 240, 420, 360), "DS18B20\n温度传感器", "PB1 / 单总线"),
        ((110, 410, 420, 530), "浊度传感器模块", "PA6 / ADC1_CH6"),
        ((110, 580, 420, 700), "pH 传感器模块", "PA7 / ADC1_CH7"),
        ((110, 750, 420, 870), "TDS 传感器模块", "PA5 / ADC1_CH5"),
    ]

    for box, title, pin in sensors:
        rounded_box(draw, box, fill="#ecfeff", outline="#0891b2", radius=26, width=4)
        draw_centered_multiline(draw, box, title, font=label_font, fill="#164e63")
        arrow_line(draw, (box[2], (box[1] + box[3]) // 2), (620, (box[1] + box[3]) // 2), fill="#0891b2", width=4)
        draw.text((450, (box[1] + box[3]) // 2 - 12), pin, fill="#334155", font=pin_font)

    esp_box = (1370, 300, 1700, 440)
    rounded_box(draw, esp_box, fill="#fee2e2", outline="#dc2626", radius=26, width=4)
    draw_centered_multiline(draw, esp_box, "ESP-01 WiFi 模块\n负责联网与 HTTP 上传", font=label_font, fill="#7f1d1d")
    arrow_line(draw, (1180, 380), (1370, 380), fill="#dc2626", width=4)
    draw.text((1220, 340), "PA2 / USART2_TX", fill="#334155", font=pin_font)
    draw.text((1220, 390), "PA3 / USART2_RX", fill="#334155", font=pin_font)

    uart_box = (1370, 540, 1700, 670)
    rounded_box(draw, uart_box, fill="#e2e8f0", outline="#475569", radius=26, width=4)
    draw_centered_multiline(draw, uart_box, "调试串口\n输出采样与联网日志", font=label_font, fill="#1e293b")
    arrow_line(draw, (1180, 605), (1370, 605), fill="#475569", width=4)
    draw.text((1215, 570), "PA9 / USART1_TX", fill="#334155", font=pin_font)
    draw.text((1215, 618), "PA10 / USART1_RX", fill="#334155", font=pin_font)

    led_box = (1370, 760, 1700, 880)
    rounded_box(draw, led_box, fill="#dcfce7", outline="#16a34a", radius=26, width=4)
    draw_centered_multiline(draw, led_box, "心跳指示灯\n显示系统运行状态", font=label_font, fill="#14532d")
    arrow_line(draw, (1180, 820), (1370, 820), fill="#16a34a", width=4)
    draw.text((1260, 788), "PC13", fill="#334155", font=pin_font)

    arrow_line(draw, (900, 240), (900, 350), fill="#d97706", width=4)

    footer = "建议插入论文硬件设计章节，图号可按“图4-x”或插入位置重新编号。"
    draw.text((505, 1024), footer, fill="#64748b", font=pin_font)
    image.save(path)


def set_chinese_font(run, zh: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    set_chinese_font(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold


def style_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line: bool = True) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(24) if first_line else Pt(0)


def add_title(doc: Document, text: str, *, size: float = 20, space_after: float = 12) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=size, bold=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(10 if level == 1 else 6)
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Pt(0)
    run = p.add_run(text)
    size = 16 if level == 1 else 14 if level == 2 else 12
    style_run(run, zh="黑体", size=size, bold=True)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    style_paragraph(p)
    run = p.add_run(text)
    style_run(run, zh="宋体", size=12)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    run = p.add_run(text)
    style_run(run, zh="楷体", size=11)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    run = p.add_run(f"• {text}")
    style_run(run, zh="宋体", size=12)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.space_before = Pt(4)
    fmt.space_after = Pt(6)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    style_run(run, zh="宋体", size=10.5)


def add_image(doc: Document, path: Path, width_cm: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def add_interface_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["模块名称", "接口或引脚", "功能说明"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        style_run(run, zh="黑体", size=10.5, bold=True)

    rows = [
        ("STM32F103C8T6", "最小系统板", "负责采样调度、数据换算、串口通信和状态控制"),
        ("DS18B20 温度传感器", "PB1", "获取水温数字量，减少温度通道的模拟干扰"),
        ("浊度传感器模块", "PA6 / ADC1_CH6", "采集浊度模拟电压信号"),
        ("pH 传感器模块", "PA7 / ADC1_CH7", "采集 pH 模拟电压信号"),
        ("TDS 传感器模块", "PA5 / ADC1_CH5", "采集 TDS 模拟电压信号"),
        ("ESP-01 WiFi 模块", "PA2 / USART2_TX；PA3 / USART2_RX", "实现联网、AT 指令交互与 HTTP 上传"),
        ("调试串口", "PA9 / USART1_TX；PA10 / USART1_RX", "输出采样和联网日志，便于联调"),
        ("心跳指示灯", "PC13", "反映主循环运行状态"),
    ]

    for module, pin, desc in rows:
        row = table.add_row()
        values = [module, pin, desc]
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(value)
            style_run(run, zh="宋体", size=10.5)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)

    add_title(doc, "《基于STM32的水质监测系统及断面数据展示平台》论文修改与补充稿", size=20, space_after=6)
    add_title(doc, "结合当前项目软硬件实现整理，可直接用于替换或补充原论文对应章节", size=12, space_after=4)
    add_title(doc, DATE_TEXT, size=12, space_after=12)

    add_note(doc, "说明：本稿重点解决导师指出的三个问题，即行文口吻技术文档化、研究现状缺少问题归纳、硬件设计内容与图示不足。插入原论文后，请按正文位置重新编号图表。")

    add_heading(doc, "一、整体修改方向", level=1)
    add_bullet(doc, "把“按代码讲实现”改成“按系统讲设计”，正文尽量少出现文件路径、函数名和“依据开题报告、课题申请表”等表述。")
    add_bullet(doc, "在国内外研究现状之后增加“已有研究存在的问题”和“本课题针对哪些问题开展研究”的过渡内容，使章节逻辑更完整。")
    add_bullet(doc, "补充终端硬件设计，包括硬件结构框图、接口分配、采样方式、无线通信链路和终端运行流程。")

    add_heading(doc, "二、建议直接替换的章节内容", level=1)

    add_heading(doc, "（一）中文摘要替换稿", level=2)
    add_body(doc, "水质监测在养殖管理、河湖断面巡检和教学实验中具有重要作用。针对人工取样周期长、连续性差以及监测结果难以统一展示的问题，本文设计并实现了一种基于 STM32 的水质监测系统及断面数据展示平台。该系统由现场采集终端、服务器端数据处理模块和 Web 可视化平台三部分组成。终端以 STM32F103C8T6 为核心，完成水温、pH、浊度和 TDS 等参数的周期采集，并通过 ESP-01 无线模块实现数据上传；服务器端负责设备接入、数据标准化、历史数据存储和最新快照更新；前端实现监测概览、历史趋势、断面地图和辅助分析等功能。系统采用历史数据与最新快照协同存储的方式，提高了实时查询与历史追溯效率，并通过站点坐标缓存增强了断面地图展示的稳定性。联调结果表明，该系统能够完成终端采样、平台入库和页面展示的完整闭环，具备较好的工程实现性与应用价值。")
    add_body(doc, "关键词：STM32；水质监测；断面展示；数据可视化；无线传输")

    add_heading(doc, "（二）1.2.3 现有研究存在的问题及本课题切入点（替换稿）", level=2)
    add_body(doc, "综合国内外研究可以看出，现有水质监测系统已经在多参数感知、无线传输和平台展示等方面取得了较多成果，但仍存在一些不足。第一，部分研究更关注单个传感节点的测量与通信，而对平台端的数据统一建模、历史管理和断面化展示关注不够，导致系统完成采集后难以形成完整业务闭环。第二，部分可视化平台主要依赖公开接口数据，缺少与现场采集终端的直接联动，难以体现真实监测系统从感知端到展示端的一体化设计。第三，当系统同时接入多来源数据时，字段命名、单位表达、时间格式和地理坐标往往不统一，这会直接影响后续的趋势分析、地图展示和风险判断。")
    add_body(doc, "针对上述问题，本文以基于 STM32 的现场采集终端和断面数据展示平台为研究对象，重点完成三方面工作：一是设计多参数水质采集终端，实现温度、pH、浊度和 TDS 等指标的周期采集与无线上传；二是构建统一的数据接入与存储机制，使现场终端数据与平台侧数据能够按照统一结构完成入库、查询和管理；三是面向断面监测场景设计可视化页面，支持概览、趋势、地图定位和辅助分析。本文的研究重点不是单独比较某一传感器的性能，而是解决“采得上来、存得下来、看得清楚、便于分析”的系统性问题。")

    add_heading(doc, "（三）1.3 研究内容与技术路线（替换稿）", level=2)
    add_body(doc, "本文围绕水质监测系统的完整实现过程展开，研究内容包括终端设计、平台设计、前端展示和系统测试四个方面。在终端设计方面，系统以 STM32F103C8T6 为控制核心，完成温度、pH、浊度和 TDS 传感器接口设计，利用数字采样与模拟采样相结合的方式获取水体参数，并通过 WiFi 模块完成数据上传。在平台设计方面，系统构建设备接入、统一字段转换、历史数据存储、最新快照维护与坐标缓存等功能，使不同来源的数据能够按统一格式进行管理。在前端展示方面，系统设计了监测概览、综合分析、断面地图和辅助研判等页面，实现对水质数据的多维可视化展示。在系统验证方面，通过终端联调、接口测试、页面功能测试和运行日志统计，对系统的完整性和可用性进行分析。")
    add_body(doc, "本文的技术路线按照“终端采集、无线传输、平台接入、数据存储、页面展示、联调验证”的顺序展开。首先完成现场采集终端的硬件连接与固件开发，实现多参数采样与周期上报；其次完成服务器端的数据接收、字段统一和存储设计，保证监测数据既可追溯又能实时展示；然后实现面向断面监测场景的页面功能，使用户能够从时间和空间两个维度观察水质变化；最后结合运行日志与联调结果验证系统闭环效果。本系统技术路线如图 1 所示。")
    add_image(doc, TECH_ROUTE_IMG, 16)
    add_caption(doc, "图1 本系统技术路线图")

    add_heading(doc, "（四）6.3.2 运行日志统计结果（改写稿）", level=2)
    add_body(doc, "根据 2026 年 3 月 5 日至 2026 年 3 月 11 日的系统运行日志统计，平台共完成 465 次自动同步任务，其中 257 次产生新增数据写入，208 次未产生新增写入。由此可以看出，系统在重复抓取场景下能够通过去重和快照更新机制避免历史数据无序膨胀，从而保证数据库结构的稳定性。")
    add_body(doc, "同一统计区间内，系统共记录 487 个地理编码批次，平均成功率约为 98.97%。这说明断面坐标补齐机制能够较好地支撑地图展示需求。结合终端采样、服务器入库和前端页面展示结果，可以认为本系统已经形成从现场采集到断面展示的完整链路，基本满足毕业设计对系统性和工程性的要求。")

    add_heading(doc, "三、建议新增或扩写的硬件设计内容", level=1)

    add_heading(doc, "（一）建议插入第三章的“终端硬件总体设计”小节", level=2)
    add_body(doc, "为实现现场水质参数的稳定采集与上传，本系统终端采用“主控单元 + 传感器单元 + 无线通信单元 + 调试与指示单元”的结构组织硬件。主控单元选用 STM32F103C8T6，负责采样调度、数据处理和通信控制；传感器单元由 DS18B20 温度传感器、pH 传感器模块、浊度传感器模块和 TDS 传感器模块组成，用于获取水体多参数信息；无线通信单元采用 ESP-01，实现设备接入 WiFi 后向服务器上传监测数据；调试与指示单元通过串口日志和心跳指示灯反映终端运行状态，便于系统联调与故障定位。终端硬件总体结构如图 2 所示。")
    add_image(doc, HARDWARE_IMG, 16.5)
    add_caption(doc, "图2 终端硬件结构框图")

    add_heading(doc, "（二）4.1.1 硬件组成与接口设计（替换/扩写稿）", level=2)
    add_body(doc, "本系统终端以 STM32F103C8T6 最小系统板为核心，利用其 ADC 通道、串口资源和通用 I/O 接口完成多参数采样与无线通信。温度采样采用 DS18B20 数字温度传感器，通过单总线方式与主控连接；pH、浊度和 TDS 传感器模块输出的模拟量分别接入 STM32 的 ADC1 通道，完成连续采样；无线通信由 ESP-01 模块承担，主控通过 USART2 与其进行 AT 指令交互，实现联网和 HTTP 上传。")
    add_body(doc, "为了便于开发和调试，系统保留了调试串口和运行状态指示灯。调试串口用于输出采样值、联网结果和上传状态，便于快速判断终端工作是否正常；心跳指示灯用于反映主循环是否稳定运行。终端主要接口分配如表 1 所示。")
    add_interface_table(doc)
    add_caption(doc, "表1 终端主要接口分配")

    add_heading(doc, "（三）4.1.2 采集接口与供电设计（新增）", level=2)
    add_body(doc, "在采集方式上，系统同时采用数字采样和模拟采样两种方案。DS18B20 直接输出数字温度数据，能够降低模拟测温受线路噪声影响的问题；pH、浊度和 TDS 三路传感器模块输出模拟电压信号，由 ADC1 采用扫描模式进行采样，并结合 DMA 循环搬运方式减少 CPU 在数据采集过程中的占用。按照当前固件实现，浊度、pH 和 TDS 三路通道分别对应 PA6、PA7 和 PA5。")
    add_body(doc, "考虑到 STM32F103C8T6 与 ESP-01 均工作在 3.3V 电平环境下，终端硬件以 3.3V 作为核心工作电压，模拟采样通道的输入信号需控制在 ADC 可接受范围内。对于原型系统而言，这种设计既满足了多参数采样的基本要求，又降低了硬件实现复杂度。为了提升页面展示结果的稳定性，终端程序还会在数据换算后对异常值进行简单裁剪，例如将 pH 约束在合理区间内，并对明显偏低的浊度和 TDS 值进行零值处理。")

    add_heading(doc, "（四）4.1.3 无线通信与数据上传设计（新增）", level=2)
    add_body(doc, "无线通信部分采用 ESP-01 模块作为 WiFi 接入单元。终端上电后，主控首先通过串口完成 AT 指令握手和基础参数配置，然后执行 WiFi 连接流程；网络连接建立后，再按设定周期向服务器发送包含温度、pH、浊度和 TDS 参数的 HTTP 请求。该方案不追求复杂协议栈，而是优先保证原型系统能够完成可靠上报，符合本科毕业设计对可实现性和演示性的要求。")
    add_body(doc, "在平台侧，服务器接收到终端上报数据后，会补充站点标识、时间戳和数据来源等字段，并将监测结果分别写入历史数据表和最新快照表。这样设计的好处在于：一方面可以完整保留采集过程中的时间序列数据，另一方面又能快速得到当前站点的最新状态，为前端概览、历史趋势和断面地图提供统一数据基础。")

    add_heading(doc, "（五）4.1.4 终端运行流程说明（替换稿）", level=2)
    add_body(doc, "终端程序采用时间片主循环方式组织任务。系统启动后依次完成延时初始化、心跳指示灯初始化、调试串口初始化、传感器初始化和 WiFi 串口初始化。进入主循环后，程序按设定时间差依次触发心跳翻转、传感器数据更新、WiFi 重连和数据上传等任务，其中传感器更新周期为 1 s、WiFi 重试周期为 5 s、数据上传周期为 15 s。与引入实时操作系统相比，这种方式结构更清晰、实现更轻量，也便于在论文中说明系统执行顺序。")
    add_body(doc, "从运行效果看，终端主循环能够完成“采样 - 处理 - 上传 - 等待下一周期”的重复执行过程。其核心价值不在于调度算法的复杂性，而在于为系统搭建起稳定的现场数据入口，使后续的平台接入、存储管理和页面展示有了真实的数据基础。")

    add_heading(doc, "四、不建议继续保留的写法", level=1)
    add_bullet(doc, "删除“依据当前项目代码、硬件原型和开题材料”等表述，改为“本文设计并实现了……”“本系统由……组成……”。")
    add_bullet(doc, "删除正文中直接列出的源码路径、文件名和函数名，可用“设备接入网关”“统一字段转换模块”“历史数据表”“最新快照表”等论文化表述替代。")
    add_bullet(doc, "把“当前项目”尽量改成“本系统”或“本文设计的系统”，减少项目汇报和技术文档的口吻。")
    add_bullet(doc, "研究现状部分不要只写别人做了什么，还要补出“还存在哪些不足”和“本文针对哪几个问题开展研究”。")
    add_bullet(doc, "硬件设计部分不要只写‘接了哪些传感器’，还要说明模块划分、接口连接、采样方式、通信链路以及为什么这样设计。")

    doc.save(str(OUT_DOCX))


def main() -> None:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    build_tech_route_image(TECH_ROUTE_IMG)
    build_hardware_image(HARDWARE_IMG)
    build_docx()


if __name__ == "__main__":
    main()
