from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
OUT = ROOT / "水质监测系统元器件清单表.docx"

BOM_ROWS: list[tuple[str, str, str, str]] = [
    ("STM32F103C8T6最小系统板", "U1", "ARM Cortex-M3，72 MHz，64 KB Flash，20 KB SRAM", "1个"),
    ("DS18B20数字温度传感器", "U2", "1-Wire，-55℃~+125℃，分辨率0.0625℃", "1个"),
    ("pH传感器模块", "U3", "0~3.3 V模拟输出，量程pH 0~14", "1个"),
    ("浊度传感器模块", "U4", "0~3.3 V模拟输出，适配ADC采样", "1个"),
    ("TDS传感器模块", "U5", "0~3.3 V模拟输出，支持mg/L或ppm换算", "1个"),
    ("ESP-01 WiFi模块", "U6", "ESP8266，USART2通信，支持AT指令", "1个"),
    ("ST-Link V2仿真器", "TOOL1", "SWD下载与在线调试", "1个"),
    ("数字万用表", "TOOL2", "电压、电流与连通性测试", "1台"),
]


def set_fonts(run, zh: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    set_fonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line: bool = False) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_line else Pt(0)


def add_title(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    style_run(run, zh="黑体", size=16, bold=True)


def add_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = paragraph.add_run(text)
    style_run(run, zh="楷体", size=11)


def add_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["comment", "designator", "value", "number"]
    widths = [Cm(5.2), Cm(3.0), Cm(8.1), Cm(2.2)]

    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = paragraph.add_run(text)
        style_run(run, zh="黑体", size=10.5, bold=True)

    for row_values in BOM_ROWS:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (1, 3) else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            style_run(run, size=10.5)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)

    add_title(doc, "水质监测系统元器件清单表")
    add_note(doc, "说明：表格根据当前硬件开发工具描述整理，仅纳入实体器件与测试工具；XCOM V2.6 为串口调试软件，未列入元器件清单。")
    add_table(doc)

    doc.save(OUT)


if __name__ == "__main__":
    build_docx()
