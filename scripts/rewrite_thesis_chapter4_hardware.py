from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


DOC_PATH = Path("/Users/caihd/Desktop/hzf/基于STM32的水质监测系统及断面数据展示平台_初稿.docx")


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def insert_paragraph_after(paragraph: Paragraph, text: str, style: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def update_table_4_1(doc: Document) -> None:
    target: Table | None = None
    for table in doc.tables:
        header = " | ".join(cell.text.strip() for cell in table.rows[0].cells)
        if header == "模块 | 连接引脚或接口 | 作用说明":
            target = table
            break

    if target is None:
        raise ValueError("Table 4-1 not found")

    rows = [
        ("模块", "连接引脚或接口", "作用说明"),
        ("STM32F103C8T6", "核心控制器", "完成采样调度、数据换算、串口通信和状态控制"),
        ("DS18B20 温度采集模块", "PB1", "读取水温数字量"),
        ("浊度传感器采集模块", "PA6 / ADC1_CH6", "采集浊度模拟电压"),
        ("pH 传感器采集模块", "PA7 / ADC1_CH7", "采集 pH 模块模拟电压"),
        ("TDS 传感器采集模块", "PA5 / ADC1_CH5", "采集 TDS 模块模拟电压"),
        ("ESP-01 WiFi 模块", "PA2 / USART2_TX；PA3 / USART2_RX", "完成无线联网与 HTTP 上报"),
        ("调试串口", "PA9 / USART1_TX；PA10 / USART1_RX", "输出启动、采样与联网日志"),
        ("心跳指示灯", "PC13", "指示主循环运行状态"),
    ]

    while len(target.rows) < len(rows):
        target.add_row()

    for r_idx, row_values in enumerate(rows):
        row = target.rows[r_idx]
        for c_idx, value in enumerate(row_values):
            row.cells[c_idx].text = value


def main() -> None:
    doc = Document(DOC_PATH)

    p_411 = find_paragraph(doc, "4.1.1 硬件组成与接口关系")
    p_411_intro = find_paragraph(
        doc,
        "从固件工程可以确认，终端主控芯片为 STM32F103C8，温度采样采用 DS18B20，模拟量采样由 ADC1 与 DMA 完成，ESP-01 通过串口与主控连接。调试串口和心跳指示灯在代码中均有独立初始化流程，便于开发阶段观察设备状态。终端硬件接口分配如表4-1所示。",
    )
    p_after_table = find_paragraph(
        doc,
        "该组合兼顾了原型实现的可获得性和系统演示的完整性。STM32F103C8 拥有足够的 GPIO、ADC 和串口资源，DS18B20 具备读取简单、结果稳定的优点，ESP-01 则适合在现有项目条件下快速打通终端与平台之间的网络链路。[10][13][14]",
    )

    p_old_412 = find_paragraph(doc, "4.1.2 采样换算与指标计算")
    p_old_412_intro = find_paragraph(
        doc,
        "终端程序将 pH、浊度和 TDS 模块的 ADC 采样结果先换算为电压，再依据当前实验环境下的经验系数转换为业务指标。该方法并不以实验室级标定为目标，而是服务于系统原型阶段对多参数采样链路的验证。主要换算公式如下。",
    )
    p_old_412_tail = find_paragraph(
        doc,
        "为保证页面显示结果的可读性，程序还对部分换算结果进行了简单裁剪。例如 pH 被限制在 0 至 14 范围内，浊度低于 35 时直接置零，TDS 低于 20 时亦置零。这样的处理能够在原型阶段抑制明显异常值，使平台端更容易得到稳定的展示结果。",
    )

    p_old_413 = find_paragraph(doc, "4.1.3 主循环与重试机制")
    p_old_413_intro = find_paragraph(
        doc,
        "终端采用时间片主循环方式组织任务。main.c 中定义了 SENSOR_SCAN_MS、ESP_RETRY_MS、UPLOAD_PERIOD_MS、HEARTBEAT_MS 和 LOOP_DELAY_MS 等常量，主循环每次延时 20 ms 后累加运行时间，根据时间差依次触发心跳翻转、采样更新、WiFi 重试和数据上报。相较于引入实时操作系统，这种实现更适合当前功能规模，也便于在论文中清晰说明执行顺序。",
    )

    p_old_414 = find_paragraph(doc, "4.1.4 上报参数组织设计")
    p_old_414_intro = find_paragraph(
        doc,
        "从 esp01_at.c 可以看出，当前固件把上报报文控制在较为紧凑的范围内，HTTP GET 请求中只包含 ph、tds、turb 和 temp 四个核心参数。这种设计有利于降低终端构造报文的复杂度，也便于在串口和服务端日志中直接核对请求内容。",
    )
    p_old_414_tail = find_paragraph(
        doc,
        "由于终端报文保持了较小的字段集合，后台系统在 device_ingest_gateway 中负责补充统一字段结构，例如 recorded_at、station_id、station_name 和 data_source 等。这样一来，终端实现可以保持简洁，系统内部的数据结构一致性则由服务端负责维护。",
    )

    p_411.text = "4.1.1 终端硬件总体组成与接口分配"
    p_411_intro.text = (
        "本系统终端硬件以 STM32F103C8T6 为核心，外围接入 DS18B20 温度采集模块、浊度传感器模块、pH 传感器模块、TDS 传感器模块、ESP-01 WiFi 模块、调试串口和心跳指示灯。各模块分别承担温度采集、模拟量采集、无线传输以及运行状态显示等任务，模块之间的连接关系较为清楚。终端主要硬件与接口分配如表4-1所示。"
    )
    p_after_table.text = (
        "从接口划分可以看出，PB1、PA5、PA6 和 PA7 主要负责传感器采集，PA2 和 PA3 负责与 ESP-01 进行串口通信，PA9 和 PA10 用于调试输出，PC13 用于主循环状态指示。这样的硬件分配能够满足本系统的采样、上传和联调需求，也便于后续固件按功能模块组织程序。"
    )

    update_table_4_1(doc)

    new_block = [
        ("4.1.2 STM32F103C8T6 最小系统设计", "Heading 3"),
        (
            "STM32F103C8T6 是终端控制部分的核心器件，最小系统电路主要由电源转换、电源去耦、时钟电路、复位电路、BOOT0 启动配置和 SWD 下载接口组成。终端各传感器采样、数据换算、串口通信和状态控制都建立在这一最小系统之上。",
            "Normal",
        ),
        (
            "根据硬件原理图，系统输入电压先由 5 V 转换为 3.3 V，再为主控和外围模块供电；主控外部使用 8 MHz 晶振提供时钟基准，NRST 引脚连接复位电路，BOOT0 默认下拉，程序下载和调试通过 SWD 接口完成。将这些基础电路单独说明后，终端硬件结构会更完整，也更符合本科毕业论文的写法。",
            "Normal",
        ),
        ("4.1.3 温度采集模块设计", "Heading 3"),
        (
            "温度采集模块采用 DS18B20 数字温度传感器，数据线连接到 PB1。该器件通过单总线方式完成通信，硬件连接简单，读取结果稳定，适合本系统的水温采集需求。终端在初始化阶段完成 DS18B20 检测，运行阶段按周期读取温度值，并将其作为水质上报数据的一部分。",
            "Normal",
        ),
        (
            "在接口设计上，DS18B20 数据线需要上拉到 3.3 V，以保证单总线通信过程中的电平稳定。由于温度数据为数字量，终端无需再对其进行 A/D 转换，程序读取后即可直接参与显示和上传，整个温度采集链路比较清楚。",
            "Normal",
        ),
        ("4.1.4 模拟量采集模块设计", "Heading 3"),
        (
            "浊度、pH 和 TDS 三类传感器模块输出的是模拟电压信号，终端通过 ADC1 完成采样，并由 DMA 将转换结果循环搬运到缓存区。这样处理后，主循环不需要反复读取 ADC 寄存器，三路数据可以持续更新，采样效率也更稳定。",
            "Normal",
        ),
        (
            "在硬件接口上，浊度传感器连接 PA6，对应 ADC1_CH6；pH 传感器连接 PA7，对应 ADC1_CH7；TDS 传感器连接 PA5，对应 ADC1_CH5。三路模拟量统一接入 3.3 V 采样系统后，终端再根据 A/D 结果进行电压换算和指标计算。将三类传感器放在同一模块中说明，能够更清楚地体现终端采集通道的整体设计。",
            "Normal",
        ),
        ("4.1.5 无线通信模块设计", "Heading 3"),
        (
            "无线通信模块采用 ESP-01 WiFi 模块，并通过 USART2 与 STM32F103C8T6 连接。其中 PA2 作为 USART2_TX 接到模块 RXD，PA3 作为 USART2_RX 接到模块 TXD，模块 EN(CH_PD) 保持高电平后进入工作状态。主控上电后先发送基础 AT 指令完成串口检测和工作模式配置，再连接无线网络，并将采集结果以 HTTP GET 的方式上传到平台。",
            "Normal",
        ),
        ("4.1.6 调试与状态指示模块设计", "Heading 3"),
        (
            "为便于硬件联调和运行观察，终端还设置了独立的调试与状态指示模块。调试信息由 USART1 输出，PA9 为发送端，PA10 为接收端，上位机能够据此查看启动日志、传感器采样值以及 WiFi 连接状态；PC13 外接心跳指示灯，主循环按固定周期翻转电平，用来反映系统是否处于正常运行状态。通过这一部分设计，程序运行过程更容易定位问题，系统维护也更方便。",
            "Normal",
        ),
    ]

    anchor = p_after_table
    for text, style in new_block:
        anchor = insert_paragraph_after(anchor, text, style)

    p_old_412.text = "4.1.7 采样换算与指标计算"
    p_old_412_intro.text = (
        "在完成各硬件模块连接后，终端还需要把 ADC 采样结果换算为可以直接展示和上传的业务数据。程序先将 pH、浊度和 TDS 的 A/D 结果换算为电压，再结合传感器调试过程中得到的经验系数计算对应指标。主要换算公式如下。"
    )
    p_old_412_tail.text = (
        "为使展示结果更加稳定，程序对部分换算值做了限幅处理。pH 被约束在 0 至 14 之间，浊度低于 35 时按 0 处理，TDS 低于 20 时同样置零。这样可以减少偶发波动对页面展示和上传结果的影响。"
    )

    p_old_413.text = "4.1.8 主循环与重试机制"
    p_old_413_intro.text = (
        "终端程序采用时间片主循环方式组织各项任务。main.c 中定义了 SENSOR_SCAN_MS、ESP_RETRY_MS、UPLOAD_PERIOD_MS、HEARTBEAT_MS 和 LOOP_DELAY_MS 等周期参数，主循环每次延时 20 ms 后累计系统运行时间，再根据时间差依次触发心跳翻转、采样更新、WiFi 重试和数据上报。"
    )
    insert_paragraph_after(
        p_old_413_intro,
        "程序在进入主循环前，会依次完成 ADC1 与 DMA 初始化、DS18B20 检测、调试串口初始化、心跳灯初始化和 USART2 初始化。这样处理后，采样、联网和上报各自按周期执行，程序结构较为清楚，也便于结合图4.1说明终端从启动到周期上传的执行过程。",
        "Normal",
    )

    p_old_414.text = "4.1.9 上报参数组织设计"
    p_old_414_intro.text = (
        "上报参数组织尽量保持简洁，终端 HTTP GET 请求中只包含 ph、tds、turb 和 temp 4 个核心字段。字段数量较少时，终端构造报文更直接，串口日志和服务端日志也更容易核对。"
    )
    p_old_414_tail.text = (
        "终端仅负责上传采样结果，recorded_at、station_id、station_name 和 data_source 等统一字段由后端在 device_ingest_gateway 中补充。这样分工后，终端侧实现不会过于复杂，平台侧的数据结构也能够保持统一。"
    )

    doc.save(DOC_PATH)


if __name__ == "__main__":
    main()
