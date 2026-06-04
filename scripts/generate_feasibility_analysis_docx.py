from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
OUT = ROOT / "3.2_可行性分析_重写稿.docx"


SECTIONS: list[tuple[str, list[str]]] = [
    (
        "3.2.1 技术可行性",
        [
            "从终端侧来看，当前原型已经具备较清楚的硬件实现路径。终端以 STM32F103C8T6 作为控制核心，接入 DS18B20 温度传感器以及 pH、浊度、TDS 三类模拟传感器，能够完成温度读取、ADC 多通道采样、数据换算和本地状态控制；ESP-01 模块通过 USART2 与主控连接，用于完成 WiFi 接入与 HTTP 数据上报，调试串口和心跳灯则用于运行观测与联调定位。STM32F1 系列资料丰富，标准外设库成熟，能够满足本课题对采样、通信和周期任务调度等方面的实现需求，因此终端侧技术路线较为明确。",
            "从平台侧来看，系统采用 Django、Vue 3、MySQL、Redis/Channels、Celery、ECharts 以及高德地图等较成熟的技术方案。后端负责设备接入、历史入库、快照维护、定时同步和接口供给，前端负责概览监测、历史分析、地图展示和 AI 辅助分析，Redis/Channels 用于实时广播，Celery 用于周期同步任务，ECharts 与高德地图用于构建图表和空间可视化。与此同时，系统还支持本地 STM32 终端接入与国家、华为等多源水质数据统一转换，这说明平台方案不仅能够运行，也保留了后续扩展数据来源和监测指标的空间。",
            "从整体链路来看，终端上报或同步任务产生的数据进入平台后，已经能够完成统一建模、历史表写入、快照更新、实时广播、页面展示和辅助分析，形成“采集—入库—展示—分析”的完整闭环。这说明本课题并不是停留在概念描述层面，而是建立在已有工程实现基础之上的。因此，从技术实现角度看，本系统具备较好的可行性。",
        ],
    ),
    (
        "3.2.2 经济可行性",
        [
            "本课题所采用的硬件以 STM32F103C8T6 最小系统板、DS18B20 温度传感器、pH 传感器、浊度传感器、TDS 传感器以及 ESP-01 WiFi 模块为主，这些器件均属于常见开发型模块，采购渠道较广、单价较低，不依赖成本较高的工业级在线监测设备或专用采集网关。对于本科毕业设计而言，该配置已经能够完成多参数采集、无线联网和端到端上报等核心功能验证，整体硬件投入较易控制。",
            "在系统联调和答辩演示阶段，平台除真实终端上报外，还可以利用现有的多源同步能力、数据库快照和历史数据接口完成页面联调与功能验证，这意味着并不需要部署大量硬件节点就能完成系统展示和论文验证。这样一来，实验成本、设备维护成本以及联调成本都能够控制在较低水平。",
            "软件部分主要依赖 Django、Vue 3、MySQL、Redis、Celery、ECharts 等开源框架和常见服务，不涉及额外商业授权费用。服务器部署既可以在本地实验环境中完成，也可以使用轻量级云主机进行演示，整体运行成本较低。因此，从经费投入和资源条件来看，本课题具有较好的经济可行性。",
        ],
    ),
    (
        "3.2.3 运行可行性",
        [
            "系统当前已经具备较明确的运行条件。项目提供了较完整的前后端工程、固件代码、部署说明和容器化编排文件，能够支持本地联调和分模块部署；开发阶段可分别启动 Django 服务、Vue 前端、Redis、Celery Worker/Beat 以及 STM32 终端，从而较快完成接口验证、页面调试和任务联调，环境复现难度相对较低。",
            "另外，系统运行方式较为灵活。在终端侧，可以通过 STM32 实物节点完成采样、串口观察和 WiFi 上报演示；在平台侧，可以通过手动上报接口、定时同步任务和数据库快照机制补足历史分析、地图展示和多断面监测场景。对于毕业设计阶段的开发、测试和答辩而言，这种“真实终端 + 平台多源数据”的运行方式既能体现硬件参与，也能保证系统演示的完整性和稳定性。",
            "因此，从系统部署、联调和展示条件来看，本课题具备较好的运行可行性。",
        ],
    ),
]


def set_fonts(run, zh: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    set_fonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line: bool = True) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_line else Pt(0)


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=16, bold=True)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(8)
    pf.space_after = Pt(4)
    pf.first_line_indent = Pt(0)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=12, bold=True)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p)
    run = p.add_run(text)
    style_run(run, size=12)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)

    add_title(doc, "3.2 可行性分析（重写稿）")

    for heading, paragraphs in SECTIONS:
        add_heading(doc, heading)
        for text in paragraphs:
            add_body(doc, text)

    doc.save(OUT)


if __name__ == "__main__":
    build_docx()
