from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/caihd/Desktop/hzf")
SRC_DOCX = ROOT / "基于STM32的水质监测系统及断面数据展示平台_本科论文初稿.docx"
OUT_DOCX = ROOT / "基于STM32的水质监测系统及断面数据展示平台_本科论文终稿版.docx"

TITLE_CN = "基于STM32的水质监测系统及断面数据展示平台的设计与实现"
TITLE_EN = "Design and Implementation of a Water Quality Monitoring System and Section Data Display Platform Based on STM32"
STUDENT = "胡卓凡"
STUDENT_ID = "202220050409"
ADVISOR = "李俊吉"
COLLEGE = "计算机科学与技术学院"
MAJOR_CLASS = "计算机科学与技术"
DATE_TEXT = "2026年5月"
HEADER_TEXT = "太原科技大学学士学位论文"


def iter_block_items(parent):
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise TypeError(f"Unsupported parent type: {type(parent)!r}")

    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def set_rfonts(run, zh: str, latin: str) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False, italic: bool = False) -> None:
    set_rfonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def add_field(paragraph: Paragraph, instruction: str, default_text: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)

    if default_text:
        text_run = paragraph.add_run(default_text)
        style_run(text_run, zh="宋体", latin="Times New Roman", size=10.5)

    run_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def set_page_number_format(section, fmt: str, start: int | None = None) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    pg_num_type = OxmlElement("w:pgNumType")
    pg_num_type.set(qn("w:fmt"), fmt)
    if start is not None:
        pg_num_type.set(qn("w:start"), str(start))
    sect_pr.append(pg_num_type)


def set_header_border(paragraph: Paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")


def configure_section(section, *, with_header: bool, page_fmt: str | None = None, start: int | None = None) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(36)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(20)
    section.header_distance = Mm(25)
    section.footer_distance = Mm(18)

    if page_fmt:
        set_page_number_format(section, fmt=page_fmt, start=start)

    if with_header:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        header_para = section.header.paragraphs[0]
        clear_paragraph(header_para)
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_para.paragraph_format.space_before = Pt(0)
        header_para.paragraph_format.space_after = Pt(0)
        set_header_border(header_para)
        run = header_para.add_run(HEADER_TEXT)
        style_run(run, zh="宋体", latin="Times New Roman", size=10.5)

        footer_para = section.footer.paragraphs[0]
        clear_paragraph(footer_para)
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.paragraph_format.space_before = Pt(0)
        footer_para.paragraph_format.space_after = Pt(0)
        add_field(footer_para, "PAGE")
        for run in footer_para.runs:
            style_run(run, zh="宋体", latin="Times New Roman", size=10.5)
    else:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        clear_paragraph(section.header.paragraphs[0])
        clear_paragraph(section.footer.paragraphs[0])


def add_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)


def add_paragraph(doc: Document, text: str = "", *, align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> Paragraph:
    p = doc.add_paragraph()
    p.alignment = align
    if text:
        p.add_run(text)
    return p


def format_body_paragraph(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, zh="宋体", latin="Times New Roman", size=12)


def format_reference_paragraph(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    pf.hanging_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, zh="宋体", latin="Times New Roman", size=12)


def format_chapter_heading(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(12)
    run = paragraph.add_run(text)
    style_run(run, zh="黑体", latin="Times New Roman", size=18, bold=True)


def format_section_heading(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    run = paragraph.add_run(text)
    style_run(run, zh="黑体", latin="Times New Roman", size=16, bold=True)


def format_subsection_heading(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    run = paragraph.add_run(text)
    style_run(run, zh="黑体", latin="Times New Roman", size=12, bold=True)


def format_caption(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    run = paragraph.add_run(text)
    style_run(run, zh="宋体", latin="Times New Roman", size=10.5)


def format_equation(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    run = paragraph.add_run(text)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=10.5)


def format_code(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    run = paragraph.add_run(text)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=10.5)


def format_cn_abstract_title(paragraph: Paragraph) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    run = paragraph.add_run("中文摘要")
    style_run(run, zh="黑体", latin="Times New Roman", size=18, bold=True)


def format_en_title(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=12, bold=True)


def format_en_center_line(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=12)


def format_en_abstract_title(paragraph: Paragraph) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    run = paragraph.add_run("Abstract")
    style_run(run, zh="Times New Roman", latin="Times New Roman", size=18, bold=True)


def format_abstract_body(paragraph: Paragraph, text: str, *, english: bool) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = paragraph.add_run(text)
    if english:
        style_run(run, zh="Times New Roman", latin="Times New Roman", size=12)
    else:
        style_run(run, zh="宋体", latin="Times New Roman", size=12)


def format_keywords(paragraph: Paragraph, label: str, content: str, *, english: bool) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run1 = paragraph.add_run(label)
    run2 = paragraph.add_run(content)
    if english:
        style_run(run1, zh="Times New Roman", latin="Times New Roman", size=12, bold=True)
        style_run(run2, zh="Times New Roman", latin="Times New Roman", size=12)
    else:
        style_run(run1, zh="黑体", latin="Times New Roman", size=12, bold=True)
        style_run(run2, zh="宋体", latin="Times New Roman", size=12)


def split_reference_items(text: str) -> list[str]:
    items = re.split(r"(?=\[\d+\])", text.strip())
    return [item.strip() for item in items if item.strip()]


def split_body_chunks(text: str) -> list[str]:
    chunks = re.split(r"(?:\n|\s{2,})+", text.strip())
    chunks = [chunk.strip() for chunk in chunks if chunk.strip()]
    return chunks or [text.strip()]


def copy_table(doc: Document, src_table: Table) -> None:
    dst_table = doc.add_table(rows=len(src_table.rows), cols=len(src_table.columns))
    dst_table.style = "Table Grid"
    dst_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row in enumerate(src_table.rows):
        for j, cell in enumerate(row.cells):
            dst_cell = dst_table.cell(i, j)
            dst_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            dst_cell.text = cell.text.strip()
            for p in dst_cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.first_line_indent = Pt(0)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                for run in p.runs:
                    style_run(
                        run,
                        zh="宋体",
                        latin="Times New Roman",
                        size=10.5,
                        bold=(i == 0),
                    )


def rewrite_text(text: str) -> str:
    stripped = text.strip()

    exact_map = {
        "设一次采样周期内读取到的原始序列为 ，则滤波结果可表示为：": "设一次采样周期内读取到的原始序列为 x1, x2, ..., xn，则滤波结果可写为：",
        "[ {x}= ]": "x_avg = (sum(x_i) - max(x) - min(x)) / (n - 2)",
        "[ EC_{25}= ]": "EC25 = ECt / (1 + a(t - 25))",
        "其中， 表示当前温度下测得的电导率， 表示折算到 25℃ 时的参考值， 为经验补偿系数。若系统需要同时显示 TDS，可再通过经验比例系数将补偿后的电导率换算为 TDS 值。pH 和浊度则分别经过标定曲线或经验关系换算后再上传。": "其中，ECt 表示当前温度下测得的电导率，EC25 表示折算到 25℃ 时的参考值，a 为经验补偿系数。若系统还需要显示 TDS，可再通过经验比例系数把补偿后的电导率换算成 TDS 值。pH 和浊度同样需要经过标定曲线或经验关系换算后再上传。",
        "[ R = 3I(DO<5)+2I(pH<6.5 pH>8.5)+WQ ]": "R = 3I(DO<5) + 2I(pH<6.5 or pH>8.5) + WQ",
        "其中， 为指示函数；当水质等级为Ⅲ类时，；为Ⅳ类时，；为Ⅴ类或劣Ⅴ类时，。这样设计的考虑是：溶解氧过低通常对水体生态影响更直接，因此权重较高；pH 偏离正常范围会影响生物代谢，权重居中；综合水质等级则综合反映整体状况，用于刻画长期风险。": "其中，I(.) 为指示函数；当水质等级为Ⅲ类时，WQ 取 2；为Ⅳ类时，WQ 取 3；为Ⅴ类或劣Ⅴ类时，WQ 取 4。这样设置的原因是，溶解氧过低往往会更直接地影响养殖环境，所以权重更高；pH 偏离正常范围会影响生物代谢，权重次之；综合水质等级主要反映整体风险水平。",
    }
    if stripped in exact_map:
        return exact_map[stripped]

    startswith_map = [
        (
            "水质参数波动具有连续性、时变性和空间差异性",
            "水质监测结果会直接影响养殖管理和现场处置效率。单靠人工取样虽然能够完成阶段性检测，但很难覆盖连续监测场景。结合本次毕业设计项目，本文围绕 STM32 感知终端和 Web 展示平台，完成了一套水质监测系统及断面数据展示平台的设计与实现。系统把设备采样、数据上传、平台入库和可视分析连成一条链路：设备端负责采集温度、浊度、pH 值和 TDS/电导率等指标，平台端负责统一建模、历史存储、快照更新和告警整理，前端则负责把结果展示为列表、图表和地图。针对多源数据字段不统一、断面坐标不完整、页面重复拉取开销较大等实际问题，本文在实现中采用了统一字段映射、历史表与快照表分离、地理编码缓存和数据版本摘要校验等方法。项目运行日志表明，系统可以完成千级断面数据的批量同步，单次同步写入量稳定在 1350 到 1558 条之间；在断面坐标补齐任务中，批处理成功率达到 98.67% 以上；当实时数据未发生变化时，接口响应体可由约 169 KB 降到 173 B。整体来看，系统已经能够满足水质监测、趋势分析、风险识别和断面展示的基本需求，也为后续接入真实硬件终端和消息总线留下了扩展空间。"
        ),
        (
            "Water quality monitoring is strongly affected by continuous fluctuation",
            "Water quality data has a direct impact on aquaculture management and on-site response. Manual sampling can satisfy occasional inspection, but it is not suitable for continuous monitoring. Based on the graduation project, this thesis designs and implements a water quality monitoring system and section data display platform centered on an STM32 terminal and a Web platform. The system connects device sampling, data upload, database storage, and visual analysis into one workflow. The terminal is responsible for collecting temperature, turbidity, pH, and TDS or conductivity data. The platform is responsible for unified modeling, history storage, snapshot updating, and alert processing. The front end presents the results through real-time lists, charts, and maps. To deal with practical issues such as heterogeneous field structures, incomplete section coordinates, and repeated front-end polling overhead, the implementation adopts unified field mapping, history-snapshot separation, geocoding cache, and data-version checksum strategies. Runtime logs show that the platform can handle batch synchronization for more than one thousand sections, with 1350 to 1558 new records written in a single task. During coordinate completion, the success rate stayed above 98.67%. When no real-time data changed, the response body dropped from about 169 KB to 173 B. These results indicate that the system can support routine monitoring, trend analysis, risk identification, and section display, while also leaving room for further expansion toward real hardware integration and message-bus based transmission."
        ),
        (
            "本课题围绕“感知终端、数据链路、展示平台”三个层面展开。",
            "这次毕业设计的核心并不是单独做一个采集板，或者单独做一个网页，而是把感知终端、数据链路和展示平台放到同一条流程里考虑。设备端以 STM32 为核心控制器，连接温度、浊度、pH 和 TDS/电导率等传感器模块，对采样值做基础滤波、标定和封包；通信层预留 MQTT 上行方式，用来支撑终端周期上传；平台端基于 Django 建立统一数据接入、历史存储和快照展示机制，前端再通过 Vue 3、ECharts 和地图组件把结果展示出来。"
        ),
        (
            "软件部分由后端服务与前端页面两部分构成。",
            "软件部分的实现更偏向工程化组织。后端主要负责数据接入、标准化处理、入库和统计分析，前端主要负责把结果展示给用户。现有项目中，后端基于 Django 搭建，核心应用包括 sensors、alerts 和 dashboard；前端基于 Vue 3 与 Vite 实现，页面分成概览、分析、地图、AI 辅助和系统设置几个部分。这样的拆分方式比较符合项目实际开发过程，也方便后续继续加功能。"
        ),
        (
            "平台的数据同步由 sync_realtime_data 负责。",
            "在实际开发中，数据同步是平台最关键的一段逻辑。现有项目把这一部分集中在 sync_realtime_data 中处理。系统会先读取当前数据源模式，再决定优先抓取哪一路数据；拉取完成后，交给 DataTransformer 做字段归一化，再补齐省市信息、时间格式和必要的坐标信息，最后分别写入历史表和快照表。这样的好处是，同步入口保持一致，后面无论接哪一种来源，都不需要把整套存储逻辑再写一遍。"
        ),
        (
            "综合测试结果可以得出以下判断。",
            "结合功能联调结果和运行日志，可以比较直观地看出这个系统已经具备了基本可用性。"
        ),
        (
            "本文围绕“基于 STM32 的水质监测系统及断面数据展示平台”完成了系统设计、平台实现与运行验证工作。",
            "通过这次毕业设计，我围绕“基于 STM32 的水质监测系统及断面数据展示平台”完成了从方案设计、平台实现到运行验证的一整套工作。整个项目并不是把硬件和网页分开处理，而是尽量沿着同一条数据链路去实现：设备端负责采样，平台端负责接收和存储，前端负责展示和分析。这样做之后，系统的结构会更完整，问题也更容易定位。"
        ),
        (
            "从已有运行结果看，系统能够稳定处理千级断面数据",
            "从已有运行结果看，系统已经能够比较稳定地处理千级断面数据，区域筛选、趋势查询、地图展示和风险识别等功能也都能正常工作。对我来说，这部分结果最有价值的地方不只是“页面能打开”，而是说明历史表与快照表分离、地理编码缓存和数据版本摘要这些设计在真实项目里确实是有效的。"
        ),
        (
            "当然，本文工作仍然存在不足。",
            "当然，这次实现还有不少可以继续往下做的地方。设备端与平台端的 MQTT 全链路接入还没有完全打通，WebSocket 实时推送目前还停留在预留阶段，传感器长期标定和现场环境适应性也需要更多实验数据支撑。后续如果能够继续接入自动控制设备、完善消息总线，并结合预测模型做异常预警，系统的实用性还会更强。"
        ),
        (
            "毕业设计从选题、实现到论文撰写，经历了一个不断修改、不断校正的过程。",
            "这次毕业设计从选题到实现，再到最后整理论文，基本上一直处在不断修改和不断回看的状态。指导教师李俊吉老师在选题方向、系统结构和论文写法上给了我很多帮助。每当我把问题想得过于简单时，老师都会提醒我回到项目本身，先把实现依据和设计原因讲清楚。这种要求也让我慢慢意识到，做毕业设计不能只追求把系统跑起来，还要说明白它为什么这样设计。"
        ),
        (
            "同时，也感谢在课程学习和项目实践中给予我帮助的老师与同学。",
            "同时，也感谢在课程学习和项目实践中帮助过我的老师和同学。设备调试、前端联调、后端接口整理以及论文修改的过程中，我都得到过很多具体的建议和提醒。感谢家人对我完成毕业设计的理解和支持，让我能够把更多时间放在系统实现和论文整理上。"
        ),
        (
            "最后，感谢大学阶段的学习经历。",
            "最后，感谢大学阶段一次次课程设计、实验和项目训练带来的积累。本文还有不少不足，但它比较完整地记录了我在本科阶段完成一次工程项目的过程，也让我对嵌入式开发和平台开发之间的关系有了更直接的理解。"
        ),
    ]

    for prefix, replacement in startswith_map:
        if stripped.startswith(prefix):
            return replacement

    stripped = stripped.replace("围绕", "针对", 1) if stripped.startswith("围绕") else stripped
    return stripped


def build_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("学士学位论文")
    style_run(run, zh="黑体", latin="Times New Roman", size=22, bold=True)

    for _ in range(5):
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(TITLE_CN)
    style_run(run, zh="黑体", latin="Times New Roman", size=22, bold=True)

    for _ in range(5):
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(0)

    for label, value in [
        ("设计人", STUDENT),
        ("学号", STUDENT_ID),
        ("指导教师", ADVISOR),
        ("所属系部", COLLEGE),
        ("专业班级", MAJOR_CLASS),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(f"{label}：{value}")
        style_run(run, zh="宋体", latin="Times New Roman", size=14)

    for _ in range(3):
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DATE_TEXT)
    style_run(run, zh="宋体", latin="Times New Roman", size=14)


def extract_key_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def main() -> None:
    src = Document(str(SRC_DOCX))
    doc = Document()

    # clear default body
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    doc.core_properties.title = TITLE_CN
    doc.core_properties.author = STUDENT

    build_cover(doc)
    cover_section = doc.sections[0]
    configure_section(cover_section, with_header=False)

    prelim_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(prelim_section, with_header=True, page_fmt="upperRoman", start=1)

    # Chinese abstract
    p = doc.add_paragraph()
    format_cn_abstract_title(p)
    p = doc.add_paragraph()
    format_abstract_body(p, rewrite_text(src.paragraphs[2].text.strip()), english=False)
    p = doc.add_paragraph()
    format_keywords(p, "关键词：", "STM32；水质监测；断面展示；Django；Vue 3", english=False)

    doc.add_page_break()

    # English abstract
    p = doc.add_paragraph()
    format_en_title(p, TITLE_EN)
    p = doc.add_paragraph()
    format_en_center_line(p, "author: Hu Zhuofan")
    p = doc.add_paragraph()
    format_en_center_line(p, "tutor: Li Junji")
    p = doc.add_paragraph()
    format_en_abstract_title(p)
    p = doc.add_paragraph()
    format_abstract_body(p, rewrite_text(src.paragraphs[7].text.strip()), english=True)
    p = doc.add_paragraph()
    format_keywords(
        p,
        "Keywords: ",
        "STM32; water quality monitoring; section display; Django; Vue 3",
        english=True,
    )

    doc.add_page_break()

    # TOC
    p = doc.add_paragraph()
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目录")
    style_run(run, zh="黑体", latin="Times New Roman", size=16, bold=True)

    toc_para = doc.add_paragraph()
    toc_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc_pf = toc_para.paragraph_format
    toc_pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    toc_pf.first_line_indent = Pt(0)
    add_field(toc_para, 'TOC \\o "1-3" \\h \\z \\u', "打开 Word 后右键更新目录")
    for run in toc_para.runs:
        style_run(run, zh="宋体", latin="Times New Roman", size=10.5)

    main_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main_section, with_header=True, page_fmt="decimal", start=1)

    blocks = list(iter_block_items(src))
    start_idx = None
    for idx, block in enumerate(blocks):
        if isinstance(block, Paragraph) and block.text.strip() == "第一章 系统概述":
            start_idx = idx
            break
    if start_idx is None:
        raise RuntimeError("Main text start not found.")

    pending_figure_caption: str | None = None
    first_chapter = True

    for i in range(start_idx, len(blocks)):
        block = blocks[i]
        next_block = blocks[i + 1] if i + 1 < len(blocks) else None

        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue

            if text.startswith("图") and isinstance(next_block, Paragraph) and next_block.style.name == "Source Code":
                pending_figure_caption = text
                continue

            if block.style.name == "Heading 1":
                if not first_chapter:
                    doc.add_page_break()
                first_chapter = False
                p = doc.add_paragraph()
                format_chapter_heading(p, text)
                continue

            if block.style.name == "Heading 2":
                p = doc.add_paragraph()
                format_section_heading(p, text)
                continue

            if block.style.name == "Heading 3":
                p = doc.add_paragraph()
                format_subsection_heading(p, text)
                continue

            if block.style.name == "Source Code":
                p = doc.add_paragraph()
                format_code(p, rewrite_text(text))
                if pending_figure_caption:
                    p = doc.add_paragraph()
                    format_caption(p, pending_figure_caption)
                    pending_figure_caption = None
                continue

            if text.startswith("表"):
                p = doc.add_paragraph()
                format_caption(p, text)
                continue

            if text == "参考文献":
                p = doc.add_paragraph()
                format_chapter_heading(p, text)
                continue

            if text.startswith("[1]") and "[2]" in text:
                for item in split_reference_items(block.text.strip()):
                    p = doc.add_paragraph()
                    format_reference_paragraph(p, item)
                continue

            if (
                "{x}" in text
                or "EC_{25}" in text
                or "3I(DO<5)" in text
            ):
                p = doc.add_paragraph()
                format_equation(p, rewrite_text(text))
                continue

            if text == "中文摘要" or text == "摘要" or text == "Abstract" or text == "目录":
                continue

            rewritten = rewrite_text(text)
            if text.startswith("x_avg =") or text.startswith("EC25 =") or text.startswith("R ="):
                p = doc.add_paragraph()
                format_equation(p, rewritten)
                continue

            chunks = split_body_chunks(rewritten)
            for chunk in chunks:
                if not chunk:
                    continue
                p = doc.add_paragraph()
                format_body_paragraph(p, chunk)

        else:
            copy_table(doc, block)

    add_update_fields_on_open(doc)
    doc.save(str(OUT_DOCX))


if __name__ == "__main__":
    main()
