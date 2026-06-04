from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/Users/caihd/Desktop/hzf")
OUT_DOCX = ROOT / "基于STM32的水质监测系统及断面数据展示平台_修改补充稿.docx"
PDF_PATH = ROOT / "document.pdf"
PDF_PAGE_DIR = ROOT / "output" / "document_pdf_pages_appkit"
ASSETS_DIR = ROOT / "thesis_assets"
CURRENT_HW_IMG = ASSETS_DIR / "fig-current-hardware-interface.png"
MIN_SYSTEM_IMG = PDF_PAGE_DIR / "page_2.png"
DATE_TEXT = "2026年5月27日"


def pick_font(bold: bool = False) -> str:
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc" if bold else "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for item in candidates:
        path = Path(item)
        if path.exists():
            return str(path)
    return ""


FONT_REGULAR = pick_font(False)
FONT_BOLD = pick_font(True)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = FONT_BOLD if bold else FONT_REGULAR
    if path:
        return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    *,
    fill: str,
    outline: str,
    radius: int = 28,
    width: int = 4,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow_line(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    fill: str,
    width: int = 4,
    head_len: int = 18,
    head_half: int = 8,
) -> None:
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=fill, width=width)
    if abs(x2 - x1) >= abs(y2 - y1):
        if x2 >= x1:
            head = [(x2, y2), (x2 - head_len, y2 - head_half), (x2 - head_len, y2 + head_half)]
        else:
            head = [(x2, y2), (x2 + head_len, y2 - head_half), (x2 + head_len, y2 + head_half)]
    else:
        if y2 >= y1:
            head = [(x2, y2), (x2 - head_half, y2 - head_len), (x2 + head_half, y2 - head_len)]
        else:
            head = [(x2, y2), (x2 - head_half, y2 + head_len), (x2 + head_half, y2 + head_len)]
    draw.polygon(head, fill=fill)


def draw_centered_multiline(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    font: ImageFont.ImageFont,
    fill: str,
    spacing: int = 8,
) -> None:
    lines = text.split("\n")
    widths: list[int] = []
    heights: list[int] = []
    for line in lines:
        left, top, right, bottom = draw.textbbox((0, 0), line, font=font)
        widths.append(right - left)
        heights.append(bottom - top)
    total_h = sum(heights) + spacing * max(0, len(lines) - 1)
    x1, y1, x2, y2 = box
    cur_y = y1 + (y2 - y1 - total_h) / 2
    for line, width, height in zip(lines, widths, heights):
        cur_x = x1 + (x2 - x1 - width) / 2
        draw.text((cur_x, cur_y), line, font=font, fill=fill)
        cur_y += height + spacing


def render_pdf_page(pdf_path: Path, page_index: int, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if out_path.exists():
        return

    from AppKit import NSBitmapImageRep, NSColor, NSGraphicsContext, NSPDFImageRep, NSPNGFileType, NSRectFill
    from Foundation import NSData, NSMakeRect

    data = NSData.dataWithContentsOfFile_(str(pdf_path))
    rep = NSPDFImageRep.imageRepWithData_(data)
    if rep is None:
        raise RuntimeError(f"Failed to render {pdf_path}")
    rep.setCurrentPage_(page_index)
    size = rep.size()
    scale = 2
    width = int(size.width * scale)
    height = int(size.height * scale)
    bitmap = NSBitmapImageRep.alloc().initWithBitmapDataPlanes_pixelsWide_pixelsHigh_bitsPerSample_samplesPerPixel_hasAlpha_isPlanar_colorSpaceName_bitmapFormat_bytesPerRow_bitsPerPixel_(
        None,
        width,
        height,
        8,
        4,
        True,
        False,
        "NSCalibratedRGBColorSpace",
        0,
        0,
        0,
    )
    context = NSGraphicsContext.graphicsContextWithBitmapImageRep_(bitmap)
    NSGraphicsContext.saveGraphicsState()
    NSGraphicsContext.setCurrentContext_(context)
    NSColor.whiteColor().set()
    NSRectFill(NSMakeRect(0, 0, width, height))
    rep.drawInRect_(NSMakeRect(0, 0, width, height))
    NSGraphicsContext.restoreGraphicsState()
    png = bitmap.representationUsingType_properties_(NSPNGFileType, None)
    png.writeToFile_atomically_(str(out_path), True)


def build_current_hardware_image(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 1120), "#fbfdff")
    draw = ImageDraw.Draw(image)

    title_font = load_font(44, bold=True)
    label_font = load_font(28, bold=True)
    body_font = load_font(23, bold=False)
    pin_font = load_font(20, bold=False)

    draw.text((420, 34), "按系统实际接口整理的终端硬件接口图", fill="#0f172a", font=title_font)
    draw.text((250, 94), "该图用于论文硬件章节，接口关系与本系统软硬件设计一致，可替代参考原理图中不一致的外设页", fill="#64748b", font=pin_font)

    power_box = (650, 150, 1150, 250)
    rounded_box(draw, power_box, fill="#fef3c7", outline="#d97706")
    draw_centered_multiline(draw, power_box, "USB 5V / 外部 5V 输入\nLDO 转 3.3V 后供给 STM32 与外围模块", font=body_font, fill="#7c2d12")

    mcu_box = (600, 360, 1200, 760)
    rounded_box(draw, mcu_box, fill="#dbeafe", outline="#2563eb", radius=36, width=5)
    draw_centered_multiline(
        draw,
        mcu_box,
        "STM32F103C8T6\n主控单元\nADC采样 / 温度读取 / AT通信\n主循环调度 / 数据换算 / 上报控制",
        font=label_font,
        fill="#0f172a",
        spacing=10,
    )

    left_nodes = [
        ((80, 230, 430, 350), "DS18B20 温度传感器", "PB1 / 1-Wire"),
        ((80, 400, 430, 520), "浊度传感器模块", "PA6 / ADC1_CH6"),
        ((80, 570, 430, 690), "pH 传感器模块", "PA7 / ADC1_CH7"),
        ((80, 740, 430, 860), "TDS 传感器模块", "PA5 / ADC1_CH5"),
    ]
    for box, title, pin in left_nodes:
        rounded_box(draw, box, fill="#ecfeff", outline="#0891b2")
        draw_centered_multiline(draw, box, title, font=label_font, fill="#164e63")
        mid_y = (box[1] + box[3]) // 2
        arrow_line(draw, (box[2], mid_y), (600, mid_y), fill="#0891b2")
        draw.text((455, mid_y - 12), pin, fill="#334155", font=pin_font)

    right_nodes = [
        ((1360, 260, 1710, 400), "ESP-01 WiFi 模块", "PA2 / USART2_TX\nPA3 / USART2_RX"),
        ((1360, 500, 1710, 640), "调试串口", "PA9 / USART1_TX\nPA10 / USART1_RX"),
        ((1360, 740, 1710, 860), "心跳指示灯", "PC13"),
    ]
    colors = [
        ("#fee2e2", "#dc2626", "#7f1d1d"),
        ("#e2e8f0", "#475569", "#1e293b"),
        ("#dcfce7", "#16a34a", "#14532d"),
    ]
    for (box, title, pin), (fill, outline, text_color) in zip(right_nodes, colors):
        rounded_box(draw, box, fill=fill, outline=outline)
        draw_centered_multiline(draw, (box[0], box[1] + 6, box[2], box[3] - 36), title, font=label_font, fill=text_color)
        draw_centered_multiline(draw, (box[0], box[1] + 60, box[2], box[3]), pin, font=body_font, fill="#334155")
        mid_y = (box[1] + box[3]) // 2
        arrow_line(draw, (1200, mid_y), (1360, mid_y), fill=outline)

    arrow_line(draw, (900, 250), (900, 360), fill="#d97706")

    bottom_box = (530, 920, 1270, 1020)
    rounded_box(draw, bottom_box, fill="#f8fafc", outline="#94a3b8")
    draw_centered_multiline(
        draw,
        bottom_box,
        "说明：原理图参考页中的 pH/TDS 引脚与系统实际接口不一致，论文中应以本图和系统接口分配为准；\nSTM32 最小系统图可继续使用 document.pdf 第 2 页。",
        font=pin_font,
        fill="#475569",
        spacing=8,
    )
    image.save(path)


def set_fonts(run, zh: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def style_run(run, *, zh: str = "宋体", latin: str = "Times New Roman", size: float = 12, bold: bool = False) -> None:
    set_fonts(run, zh=zh, latin=latin)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line: bool = True) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if first_line else Pt(0)


def add_title(doc: Document, text: str, *, size: float, bold: bool = True, after: float = 10) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=size, bold=bold)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(8 if level == 1 else 5)
    pf.space_after = Pt(5)
    pf.first_line_indent = Pt(0)
    run = p.add_run(text)
    style_run(run, zh="黑体", size=16 if level == 1 else 14 if level == 2 else 12, bold=True)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p)
    run = p.add_run(text)
    style_run(run, size=12)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    run = p.add_run(text)
    style_run(run, zh="楷体", size=11)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    style_run(run, size=10.5)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    run = p.add_run(f"• {text}")
    style_run(run, size=12)


def add_image(doc: Document, path: Path, width_cm: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def style_table_text(paragraph, text: str, *, bold: bool = False, center: bool = False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, size=10.5, bold=bold, zh="黑体" if bold else "宋体")


def add_revision_location_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["原论文位置", "处理方式", "需要补充的内容", "建议结果"]
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        style_table_text(cell.paragraphs[0], title, bold=True, center=True)

    rows = [
        ("3.3.1 终端模块划分", "改写", "把终端拆分为主控最小系统、温度采集、模拟量采集、无线通信、调试与状态指示五个模块", "让第三章先交代硬件结构，不再只从软件模块角度写"),
        ("4.1 节标题后", "新增导语", "说明硬件章节依据 document.pdf 第 2 页和系统实际接口整理", "明确“原理图基础 + 系统实际设计”的关系"),
        ("4.1.1 硬件组成与接口关系", "整体替换", "补终端总体结构图、代码一致版接口分配表，并说明最终引脚以程序为准", "把抽象描述改成能落地的硬件章节"),
        ("4.1.1 后", "新增 4.1.2", "插入 STM32 最小系统设计，使用 document.pdf 第 2 页", "补齐电源、晶振、复位、BOOT0、SWD"),
        ("原 4.1.2 之前", "新增 4.1.3", "补 DS18B20、pH、浊度、TDS 的采集接口设计", "把“只有公式没有电路”的问题补齐"),
        ("原 4.1.3 之前", "新增 4.1.4", "补 ESP-01 联网接口、USART1 调试串口、PC13 心跳灯", "把通信与调试硬件写完整"),
        ("5.2.1 采样与状态控制实现", "小改", "把代码实现和第四章的最终硬件接口一一对应", "实现章节与硬件章节保持一致"),
    ]

    for values in rows:
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            style_table_text(cell.paragraphs[0], value, center=(idx == 1))


def add_compare_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["项目", "document.pdf 情况", "系统实际设计", "论文写法建议"]
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        style_table_text(cell.paragraphs[0], title, bold=True, center=True)

    rows = [
        ("STM32 最小系统", "第 2 页完整绘制", "与本系统设计一致", "可直接作为“STM32 最小系统设计图”插图"),
        ("pH / TDS 接口", "第 1 页绘制，但引脚落在 PA1 / PA0", "系统实际采样为 PA7 / PA5", "保留接口思路，正文和接口表必须改为 PA7 / PA5"),
        ("浊度接口", "第 1 页未单独形成最终代码一致版接口", "系统实际采样为 PA6", "需补一张与系统接口一致的接口图"),
        ("DS18B20", "原理图中未体现", "系统实际连接 PB1", "新增温度采集接口说明"),
        ("ESP-01", "原理图中未体现", "系统实际连接 USART2，使用 PA2 / PA3", "新增无线通信接口说明"),
        ("调试串口 / 心跳灯", "原理图外设页未体现", "系统使用 USART1(PA9/PA10) 和 PC13", "新增调试与状态指示说明"),
        ("OLED / 按键 / 蜂鸣器", "第 1 页存在", "不属于本系统论文重点展开的核心硬件", "不能写成“本系统最终已实现的主功能硬件”"),
    ]

    for values in rows:
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = ""
            style_table_text(cell.paragraphs[0], value, center=(idx == 0))


def add_interface_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["功能模块", "实际连接引脚或接口", "说明"]
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        style_table_text(cell.paragraphs[0], title, bold=True, center=True)

    rows = [
        ("DS18B20 温度采集", "PB1", "用于采集温度数据，系统按单总线方式完成初始化和读取"),
        ("浊度采样", "PA6 / ADC1_CH6", "对应 ADCConvertedValue[0]，经换算得到浊度值"),
        ("pH 采样", "PA7 / ADC1_CH7", "对应 ADCConvertedValue[1]，先换算电压后计算 pH"),
        ("TDS 采样", "PA5 / ADC1_CH5", "对应 ADCConvertedValue[2]，按多项式计算 TDS"),
        ("ESP-01 联网", "PA2 / USART2_TX；PA3 / USART2_RX", "通过 AT 指令完成联网和 HTTP 上报"),
        ("调试串口", "PA9 / USART1_TX；PA10 / USART1_RX", "输出采样与联网日志，便于调试"),
        ("心跳灯", "PC13", "用于反映主循环运行状态"),
    ]
    for values in rows:
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = ""
            style_table_text(cell.paragraphs[0], value, center=(idx < 2))


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)

    add_title(doc, "《基于STM32的水质监测系统及断面数据展示平台》", size=20)
    add_title(doc, "硬件部分修改说明与可直接替换正文", size=18, after=4)
    add_title(doc, DATE_TEXT, size=12, after=10)
    add_note(doc, "说明：本稿只处理论文硬件部分，重点回答“原论文哪里改、需要补什么内容、哪些图可以直接用、哪些内容需要按系统实际接口调整”。")

    add_heading(doc, "一、原论文中硬件部分的具体修改位置", level=1)
    add_body(doc, "根据初稿当前结构，硬件部分的主修改位置集中在 3.3.1、4.1 及 5.2.1。其中第三章负责交代终端模块结构，第四章负责完整展开硬件设计，第五章只保留与硬件对应的实现说明。建议修改位置如下表所示。")
    add_revision_location_table(doc)
    add_caption(doc, "表A-1 原论文硬件部分的修改位置与处理方式")

    add_heading(doc, "二、现有原理图与系统实际设计的对应关系", level=1)
    add_body(doc, "现有原理图 document.pdf 可以分开处理。第 2 页展示了 STM32F103C8T6 最小系统，包括 USB 供电、LDO 转 3.3V、晶振、复位、BOOT0 和 SWD 下载口，这一页与本系统硬件设计一致，可直接作为论文中的“STM32 最小系统设计图”。")
    add_body(doc, "第 1 页更适合作为外设接口参考页，不能直接作为“本系统最终传感器接口图”。原因在于，该页中的 pH、TDS 接口落在 PA1、PA0，且包含 OLED、按键和蜂鸣器等外围模块；而本系统实际使用的是 DS18B20、PA6/PA7/PA5 三路 ADC、USART2 连接 ESP-01、USART1 调试串口以及 PC13 心跳灯。写作时可以保留该页提供的接口设计思路，但最终引脚和模块关系应以系统实际接口分配为准。")
    add_image(doc, MIN_SYSTEM_IMG, 15.8)
    add_caption(doc, "图A-1 现有原理图中的 STM32 最小系统页（可直接用于论文硬件章节）")
    add_compare_table(doc)
    add_caption(doc, "表A-2 现有原理图与系统实际设计的对应关系")

    add_heading(doc, "三、建议新增的图表", level=1)
    add_body(doc, "为了让论文硬件部分从“只有模块名和公式”变成完整设计，建议至少补 2 张图和 2 张表。图表编号可并入正文后再统一调整。推荐如下：")
    add_bullet(doc, "图4-x 终端硬件总体结构图：按系统实际接口重画，突出 DS18B20、浊度、pH、TDS、ESP-01、调试串口和心跳灯。")
    add_bullet(doc, "图4-x STM32F103C8T6 最小系统原理图：直接使用 document.pdf 第 2 页。")
    add_bullet(doc, "表4-x 终端硬件接口分配表：列出功能模块、引脚和作用说明。")
    add_bullet(doc, "表4-x 原理图与系统实际设计对照表：说明哪些外设可以直接使用，哪些需要按系统接口调整。")
    add_image(doc, CURRENT_HW_IMG, 16)
    add_caption(doc, "图A-2 按系统实际接口整理的终端硬件接口图")
    add_interface_table(doc)
    add_caption(doc, "表A-3 系统实际接口版终端硬件接口分配表")

    add_heading(doc, "四、可直接替换或新增的正文", level=1)

    add_heading(doc, "（一）替换 3.3.1 终端模块划分", level=2)
    add_body(doc, "本系统终端部分由主控最小系统、温度采集、模拟量采集、无线通信以及调试与状态指示 5 个模块组成。主控最小系统以 STM32F103C8T6 为核心，负责供电、时钟、程序下载和接口调度；温度采集模块采用 DS18B20，用于获取水温参数；模拟量采集模块负责浊度、pH 和 TDS 三路传感器的 ADC 输入；无线通信模块采用 ESP-01，通过串口发送 AT 指令实现联网和 HTTP 上报；调试与状态指示部分由 USART1 调试串口和 PC13 心跳灯构成，用来显示系统运行状态。按这样的方式划分后，终端硬件结构会更清楚，后续各模块的设计说明也更容易展开。")

    add_heading(doc, "（二）新增 4.1 节导语", level=2)
    add_body(doc, "本节围绕终端硬件设计展开说明，内容依据现有硬件原理图和系统实际接口配置进行整理。需要说明的是，原理图中的 STM32 最小系统页可以直接作为硬件设计基础，外设页主要用于接口参考；论文中的传感器连接关系和通信接口分配，应以本系统实际使用的引脚为准。下面将分别对 STM32 最小系统、传感器采集接口以及无线通信与调试接口进行说明。")

    add_heading(doc, "（三）替换 4.1.1 硬件组成与接口关系", level=2)
    add_body(doc, "本系统终端以 STM32F103C8T6 作为核心控制器，硬件结构围绕最小系统、传感器采集、无线通信和调试指示几部分展开。温度采集采用 DS18B20，浊度、pH 和 TDS 三路传感器通过 ADC1 扫描方式接入主控，ESP-01 通过 USART2 与主控建立 AT 指令通信链路，USART1 负责输出运行日志，PC13 连接心跳灯，用于显示系统的运行状态。这样写可以把终端硬件结构、接口说明和后文实现内容对应起来，章节之间也更连贯。")
    add_body(doc, "在本系统中，DS18B20 连接在 PB1，三路 ADC 分别使用 PA6、PA7 和 PA5，对应浊度、pH 和 TDS 三路采样输入；ESP-01 连接在 USART2，对应 PA2 和 PA3；调试串口使用 USART1，对应 PA9 和 PA10；心跳灯位于 PC13。终端硬件总体结构如图4-x所示，接口分配见表4-x。")

    add_heading(doc, "（四）新增 4.1.2 STM32 最小系统设计", level=2)
    add_body(doc, "STM32 最小系统设计是终端硬件能够稳定工作的基础。根据现有原理图的第 2 页，系统采用 USB 5V 作为输入电源，经 LDO 芯片转换为 3.3V，为 STM32 主控及外围模块供电；同时在电源输入与输出两侧布置滤波和去耦电容，以提高供电稳定性。时钟部分采用 8MHz 晶振及配套电容，为主控提供系统时钟基础；复位电路采用上拉电阻和电容组成基本复位网络；BOOT0 通过下拉电阻进行启动方式配置；程序下载和调试则通过 SWD 接口完成。")
    add_body(doc, "这一最小系统设计与本系统硬件方案一致，可以直接作为论文中的硬件设计依据。写作时建议将该原理图页插入 4.1.2 小节，并围绕供电、时钟、复位、启动配置和下载接口依次说明，不宜只简单写成“采用 STM32 最小系统板”。")

    add_heading(doc, "（五）新增 4.1.3 传感器采集接口设计", level=2)
    add_body(doc, "传感器采集部分包括温度采集和模拟量采集两类接口。温度采集采用 DS18B20，系统通过 PB1 完成初始化和温度读取，该模块用于获取水温信息，也为后续水质参数分析提供温度参考。模拟量采集部分由浊度、pH 和 TDS 三路输入组成，系统通过 ADC1 的扫描模式和 DMA 循环搬运机制，对三路模拟信号进行连续采样，其中浊度输入连接 PA6，pH 输入连接 PA7，TDS 输入连接 PA5。")
    add_body(doc, "还需要说明的是，现有原理图外设页中给出的 pH、TDS 接口引脚，与本系统实际接口分配并不一致，因此论文中不宜直接沿用原页中的 PA0、PA1 标注。更合适的写法是，将原理图作为 pH 和 TDS 接口设计的参考，再结合系统实际连接关系，把三路 ADC 接口确定为 PA6、PA7 和 PA5，并补入浊度接口说明，这样既能交代硬件设计来源，也能保证论文内容和系统实现一致。")

    add_heading(doc, "（六）新增 4.1.4 无线通信与调试接口设计", level=2)
    add_body(doc, "无线通信部分采用 ESP-01 WiFi 模块，通过 USART2 与 STM32 主控连接。系统上电后，主控向 ESP-01 发送基础 AT 指令，完成波特率适配、工作模式配置和 WiFi 连接；网络建立后，再以 HTTP GET 方式将温度、pH、浊度和 TDS 等参数上传到平台设备网关。这样的通信链路比较直接，实现起来也较为清晰，能够满足本系统终端数据上报的需求。")
    add_body(doc, "除数据上传链路外，系统还保留了独立的调试与状态指示接口。USART1 用于向上位机输出采样和联网日志，便于观察温度读取、ADC 采样、WiFi 连接和数据上报等过程；PC13 上连接心跳灯，用于反映主循环运行状态。这样处理后，终端硬件在完成数据采集与上传的同时，也具备了较好的调试和状态显示能力。")

    add_heading(doc, "（七）调整原 4.1.2、4.1.3、4.1.4 的写法", level=2)
    add_body(doc, "建议保留你原来 4.1.2 中的公式部分，但在其前面补一句过渡说明：‘在完成传感器接口设计后，系统将三路 ADC 采样值换算为业务指标，以便后续上报和页面展示。’同时建议将原 4.1.2 改名为‘采样换算与指标计算’，把原 4.1.3 改为‘主循环与重试机制’，把原 4.1.4 改为‘上报参数组织设计’。这样章节逻辑将变成“先硬件、后算法、再流程、最后上报”，整体更符合论文写法。")

    add_heading(doc, "（八）修改 5.2.1 采样与状态控制实现", level=2)
    add_body(doc, "在 5.2.1 中，建议把实现说明改写为：‘系统初始化阶段依次完成 ADC1 与 DMA 配置、DS18B20 初始化、调试串口初始化、心跳灯初始化以及 ESP-01 所在 USART2 的初始化；进入运行阶段后，再按固定周期完成温度读取、三路 ADC 数据更新、日志输出、WiFi 连接与数据上传。其对应的硬件接口分别为 PB1、PA6、PA7、PA5、PA2/PA3、PA9/PA10 和 PC13。’这样写后，实现章节会直接呼应第四章中的硬件接口设计。")

    add_heading(doc, "五、论文中不应直接照搬的内容", level=1)
    add_bullet(doc, "不要把 document.pdf 第 1 页原样写成“本系统最终硬件原理图”，因为其中的 pH/TDS 引脚与系统实际接口分配不一致。")
    add_bullet(doc, "不要把 OLED、按键和蜂鸣器写成“本系统当前主功能硬件模块”，除非后文有对应程序实现和测试结果。")
    add_bullet(doc, "不要继续使用 PA0、PA1 作为 pH、TDS 最终引脚写法，本系统实际采用的是 PA7、PA5，浊度为 PA6。")
    add_bullet(doc, "不要只写“STM32 + 传感器 + ESP-01”，而不说明最小系统、电源、晶振、复位、BOOT0 和 SWD。")

    add_heading(doc, "六、可补入参考文献或资料来源的器件文档", level=1)
    add_bullet(doc, "本地资料：document.pdf（STM32 最小系统与参考外设页）。")
    add_bullet(doc, "本地资料：PH 传感器使用手册、TDS 传感器使用手册。")
    add_bullet(doc, "官方资料：STM32F103C8T6 数据手册与 STM32F10xxx 硬件设计应用笔记。")
    add_bullet(doc, "官方资料：DS18B20 数据手册。")
    add_bullet(doc, "官方资料：ESP8266 / ESP-AT 硬件连接与硬件设计指南。")

    doc.save(str(OUT_DOCX))


def main() -> None:
    render_pdf_page(PDF_PATH, 1, MIN_SYSTEM_IMG)
    build_current_hardware_image(CURRENT_HW_IMG)
    build_docx()


if __name__ == "__main__":
    main()
